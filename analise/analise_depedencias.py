import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict

# ======== CAMINHOS ========
input_path = r"C:\Users\GustavoArruda\Mendes Lobo\TI - Documentos\dependencies.xlsx"
output_path = r"C:\Users\GustavoArruda\Mendes Lobo\TI - Documentos\analise_conexoes_indevidas5.xlsx"

df = pd.read_excel(input_path)

# ======== FUNÇÃO PARA PEGAR ABA ========
def get_sheet(cell_ref):
    if '!' in str(cell_ref):
        return str(cell_ref).split('!')[0].strip("'")
    return str(cell_ref)

df['Source_Sheet'] = df['Source'].apply(get_sheet)
df['Target_Sheet'] = df['Target'].apply(get_sheet)

# ======== CAMADAS ========
layer1 = {'DADOS', 'EQUIPE', 'PRODUÇÕES', 'CPU', 'BDI', 'PRATICABILIDADE'}
layer2 = {'ADM', 'EQUIPAMENTOS', 'TOP', 'MÃO DE OBRA', 'QSMS', 'Escavação e Carga', 'LAB', 'MOB'}
layer3 = {'CRONOGRAMA', 'PLANEJAMENTO'}
layer4 = {'PONTO DE EQUILIBRIO', 'RESULTADO'}

layer_map = {}
for l in layer1: layer_map[l] = 1
for l in layer2: layer_map[l] = 2
for l in layer3: layer_map[l] = 3
for l in layer4: layer_map[l] = 4

def get_layer_name(s):
    if s in layer1: return 'Camada 1 - Base'
    if s in layer2: return 'Camada 2 - Cálculo'
    if s in layer3: return 'Camada 3 - Montagem'
    if s in layer4: return 'Camada 4 - Análise Final'
    return 'Outras'

# ======== DETECÇÃO DE PROBLEMAS ========
problems_list = []

for _, row in df.iterrows():
    consumidor = row['Source_Sheet']
    fornecedor = row['Target_Sheet']

    if consumidor == fornecedor:
        continue

    ptype = None

    if fornecedor in layer4 and consumidor not in layer4:
        ptype = 'Camada 4 alimentando camada inferior'
    elif fornecedor in layer3 and consumidor in (layer1 | layer2):
        ptype = 'Camada 3 alimentando camada inferior'
    elif fornecedor in layer2 and consumidor in layer1:
        ptype = 'Camada 2 alimentando Camada 1'

    if ptype:
        problems_list.append({
            'Aba Consumidora': consumidor,
            'Celula que consome': str(row['Source']),
            'Aba Fornecedora Indevida': fornecedor,
            'Celula fornecedora': str(row['Target']),
            'Tipo de Problema': ptype,
            'Camada do Consumidor': get_layer_name(consumidor),
            'Camada do Fornecedor': get_layer_name(fornecedor),
        })

problems_df = pd.DataFrame(problems_list)

# ======== GRAFO (fornecedor -> consumidores) ========
graph = defaultdict(set)

for _, row in df.iterrows():
    fornecedor = row['Target_Sheet']
    consumidor = row['Source_Sheet']
    graph[fornecedor].add(consumidor)

# ======== BUSCAR ANCESTRAIS ========
def get_ancestors(node, graph, max_depth=5):
    ancestors = set()
    stack = [(node, 0)]

    while stack:
        current, depth = stack.pop()

        if depth > max_depth:
            continue

        for parent in graph:
            if current in graph[parent]:
                ancestors.add(parent)
                stack.append((parent, depth + 1))

    return ancestors

# ======== SUGESTÃO INTELIGENTE ========
smart_rebuild = []

for _, row in problems_df.iterrows():
    consumidor = row['Aba Consumidora']
    fornecedor_errado = row['Aba Fornecedora Indevida']

    camada_consumidor = layer_map.get(consumidor, 0)

    ancestors = get_ancestors(fornecedor_errado, graph)

    candidatos_validos = []

    for anc in ancestors:
        camada_anc = layer_map.get(anc, 0)

        if camada_anc <= camada_consumidor:
            candidatos_validos.append((anc, camada_anc))

    if candidatos_validos:
        melhor = sorted(candidatos_validos, key=lambda x: x[1], reverse=True)[0]

        smart_rebuild.append({
            'Consumidor': consumidor,
            'Fornecedor Atual (Errado)': fornecedor_errado,
            'Sugestao Inteligente': f"Usar {melhor[0]} (Camada {melhor[1]}) ao invés de {fornecedor_errado}"
        })

smart_df = pd.DataFrame(smart_rebuild)

# ======== EXCEL ========
wb = Workbook()

BLUE = '1F4E79'

def hdr(ws, r, c, val):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', start_color=BLUE)
    cell.alignment = Alignment(horizontal='center')

def dat(ws, r, c, val):
    ws.cell(row=r, column=c, value=val)

# ======== RESUMO ========
ws1 = wb.active
ws1.title = 'Resumo'

ws1.cell(1,1,'Resumo de Problemas')

summary = problems_df.groupby('Tipo de Problema').size().reset_index(name='Qtd')

for i, row in enumerate(summary.itertuples(), start=3):
    dat(ws1, i, 1, row[1])
    dat(ws1, i, 2, row[2])

# ======== DETALHAMENTO ========
ws2 = wb.create_sheet('Detalhamento')

headers = ['Consumidor','Celula','Fornecedor','Celula Fonte','Problema']
for col, h in enumerate(headers, start=1):
    hdr(ws2, 2, col, h)

for i, row in enumerate(problems_df.itertuples(), start=3):
    dat(ws2, i, 1, row[1])
    dat(ws2, i, 2, row[2])
    dat(ws2, i, 3, row[3])
    dat(ws2, i, 4, row[4])
    dat(ws2, i, 5, row[5])

# ======== SUGESTÕES INTELIGENTES ========
ws3 = wb.create_sheet('Sugestoes Inteligentes')

headers = ['Consumidor','Fornecedor Atual','Sugestao']
for col, h in enumerate(headers, start=1):
    hdr(ws3, 2, col, h)

for i, row in enumerate(smart_df.itertuples(), start=3):
    dat(ws3, i, 1, row[1])
    dat(ws3, i, 2, row[2])
    dat(ws3, i, 3, row[3])

# ======== SALVAR ========
wb.save(output_path)

print(f"\n✔ {len(problems_df)} conexões problemáticas encontradas.\n")
print(problems_df.groupby('Tipo de Problema').size())

print(f"\n🧠 {len(smart_df)} sugestões inteligentes geradas.\n")