import re
import json
from collections import defaultdict
from openpyxl import load_workbook, Workbook

# =========================
# REGEX DE REFERÊNCIAS
# =========================
CELL_REF_REGEX = re.compile(
    r"(?:'([^']+)'|([A-Za-z0-9_]+))?!?\$?[A-Z]{1,3}\$?\d+"
)

# =========================
# NORMALIZA REFERÊNCIA
# =========================
def normalize_ref(sheet_name, match):
    quoted_sheet, plain_sheet = match.groups()
    target_sheet = quoted_sheet or plain_sheet or sheet_name
    return target_sheet

# =========================
# EXTRAI DEPENDÊNCIAS
# =========================
def extract_dependencies_from_formula(formula, current_sheet):
    deps = set()

    for match in CELL_REF_REGEX.finditer(formula):
        full_ref = match.group(0)

        # Se não tiver sheet, assume a atual
        if "!" not in full_ref:
            full_ref = f"{current_sheet}!{full_ref}"

        deps.add(full_ref)

    return deps

# =========================
# MONTA GRAFO
# =========================
def build_dependency_graph(file_path):
    wb = load_workbook(file_path, data_only=False)
    graph = defaultdict(list)

    for sheet in wb.worksheets:
        sheet_name = sheet.title

        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":  # fórmula
                    current = f"{sheet_name}!{cell.coordinate}"
                    formula = cell.value

                    deps = extract_dependencies_from_formula(
                        formula, sheet_name
                    )

                    graph[current].extend(sorted(deps))

    return dict(graph)

# =========================
# EXPORTAR TXT
# =========================
def export_to_txt(graph, filename="dependencies.txt"):
    with open(filename, "w", encoding="utf-8") as f:
        for source, targets in graph.items():
            line = f"{source} -> {', '.join(targets)}\n"
            f.write(line)

# =========================
# EXPORTAR EXCEL (SIMPLES)
# =========================
def export_to_excel(graph, filename="dependencies.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dependencies"

    ws.append(["Source", "Target"])

    for source, targets in graph.items():
        for target in targets:
            ws.append([source, target])

    wb.save(filename)

# =========================
# EXPORTAR EXCEL (ESTRUTURADO)
# =========================
def export_to_excel_structured(graph, filename="dependencies_structured.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dependencies"

    ws.append([
        "Source",
        "Target",
        "Source Sheet",
        "Target Sheet"
    ])

    for source, targets in graph.items():
        source_sheet = source.split("!")[0]

        for target in targets:
            target_sheet = target.split("!")[0]

            ws.append([
                source,
                target,
                source_sheet,
                target_sheet
            ])

    wb.save(filename)

# =========================
# MAIN
# =========================
if __name__ == "__main__":
    file_path = r"C:\Users\GustavoArruda\Mendes Lobo\TI - Documentos\IA\planilha_investimento\2026_03_30_PLANEJAMENTO_LOTE 09 e 10_(CE).xlsm"

    print("Lendo arquivo e construindo grafo...")
    graph = build_dependency_graph(file_path)

    print("Exportando TXT...")
    export_to_txt(graph)

    print("Exportando Excel simples...")
    export_to_excel(graph)

    print("Exportando Excel estruturado...")
    export_to_excel_structured(graph)

    print("Concluído com sucesso 🚀")