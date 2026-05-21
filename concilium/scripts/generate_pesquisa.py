#!/usr/bin/env python3
"""
generate_pesquisa.py — Stage 3 do Pipeline ADR-045

Gera planilha Excel `pesquisa_{UF}_{AAAA-MM}.xlsx` com preços por fonte para uso
como documento de instrução de licitação.

Uso:
    python scripts/generate_pesquisa.py --uf CE --mes 2026-01
    python scripts/generate_pesquisa.py --uf CE --mes 2026-01 --output-dir ./pesquisas/
"""
from __future__ import annotations

import argparse
import io
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]


def _read_sicro(uf: str) -> list[dict]:
    p = ROOT / f"config/referencias/SICRO/BR-{uf}/insumos.json"
    return json.loads(p.read_text(encoding="utf-8")).get("insumos", []) if p.exists() else []


def _read_sinapi(uf: str, mes: str) -> list[dict]:
    p = ROOT / f"config/referencias/SINAPI/BR-{uf}/insumos_{mes}.json"
    return json.loads(p.read_text(encoding="utf-8")).get("insumos", []) if p.exists() else []


def _read_anp(uf: str) -> dict | None:
    anp_dir = ROOT / f"config/referencias/ANP/BR-{uf}"
    if not anp_dir.exists():
        return None
    files = sorted(anp_dir.glob("diesel_*.json"), reverse=True)
    return json.loads(files[0].read_text(encoding="utf-8")) if files else None


def _read_bundle(contexto: str) -> dict:
    p = ROOT / f"config/catalogos/{contexto}/.catalog_bundle.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def _read_coverage(uf: str) -> dict:
    logs_dir = ROOT / "config/referencias/pesquisas/logs"
    files = sorted(logs_dir.glob(f"cobertura_{uf}_*.json"), reverse=True) if logs_dir.exists() else []
    return json.loads(files[0].read_text(encoding="utf-8")) if files else {}


def generate(uf: str, mes: str, contexto: str, output_dir: Path) -> Path:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl necessário: uv add openpyxl", file=sys.stderr)
        sys.exit(2)

    sicro  = _read_sicro(uf)
    sinapi = _read_sinapi(uf, mes)
    anp    = _read_anp(uf)
    bundle = _read_bundle(contexto)
    cov    = _read_coverage(uf)
    insumos_cat = bundle.get("insumos", [])

    wb = openpyxl.Workbook()

    # ── Estilos ────────────────────────────────────────────────────────────────
    HDR_FILL  = PatternFill("solid", fgColor="1F4E79")
    HDR_FONT  = Font(color="FFFFFF", bold=True, size=10)
    TITLE_FONT = Font(bold=True, size=12)
    THIN      = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
    ERR_FILL  = PatternFill("solid", fgColor="FFE0E0")

    def _header_row(ws, cols: list[str], row: int = 1) -> None:
        for ci, h in enumerate(cols, 1):
            c = ws.cell(row=row, column=ci, value=h)
            c.font = HDR_FONT; c.fill = HDR_FILL
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = THIN

    def _auto_width(ws) -> None:
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 55)

    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    # ── Aba 1: Resumo ──────────────────────────────────────────────────────────
    ws_res = wb.active
    ws_res.title = "Resumo"
    ws_res["A1"] = f"Pesquisa de Preços — BR-{uf} — Referência {mes}"
    ws_res["A1"].font = Font(bold=True, size=14)
    ws_res["A3"] = "Estado:"
    ws_res["B3"] = f"BR-{uf}"
    ws_res["A4"] = "Mês Referência:"
    ws_res["B4"] = mes
    ws_res["A5"] = "Gerado em:"
    ws_res["B5"] = now_str
    ws_res["A6"] = "Fontes:"
    ws_res["B6"] = f"SICRO ({len(sicro)} itens) | SINAPI ({len(sinapi)} itens) | ANP (diesel)"
    ws_res["A8"]  = "Cobertura Insumos Catálogo:"
    ws_res["B8"]  = f"{cov.get('resumo', {}).get('atualizados', '?')} de {len(insumos_cat)} insumos atualizados"
    ws_res["A9"]  = "Suspensos (Δ>40%):"
    ws_res["B9"]  = str(cov.get("resumo", {}).get("suspensos", "?"))
    ws_res["A10"] = "Conflitos (Δ>15%):"
    ws_res["B10"] = str(cov.get("resumo", {}).get("conflitos", "?"))
    ws_res["A11"] = "Sem cobertura:"
    ws_res["B11"] = str(cov.get("resumo", {}).get("sem_match", "?"))
    ws_res.column_dimensions["A"].width = 28
    ws_res.column_dimensions["B"].width = 55

    # ── Aba 2: SICRO ────────────────────────────────────────────────────────────
    ws_s = wb.create_sheet("SICRO")
    cols_s = ["Código SICRO", "Descrição", "Tipo", "Unidade", "Custo Produtivo (R$/h)", "Custo Improdutivo", "Preço Publicado"]
    _header_row(ws_s, cols_s)
    for ri, ins in enumerate(sicro, 2):
        vals = [
            ins.get("codigo_referencia"), ins.get("descricao"), ins.get("tipo"),
            ins.get("unidade"), ins.get("custo_produtivo"), ins.get("custo_improdutivo"),
            ins.get("preco_publicado"),
        ]
        for ci, v in enumerate(vals, 1):
            c = ws_s.cell(row=ri, column=ci, value=v)
            c.border = THIN
            if ci in (5, 6, 7) and isinstance(v, (int, float)):
                c.number_format = "R$ #,##0.0000"
    _auto_width(ws_s)

    # ── Aba 3: SINAPI ───────────────────────────────────────────────────────────
    ws_n = wb.create_sheet("SINAPI")
    cols_n = ["Código SINAPI", "Descrição", "Tipo", "Unidade", "Preço Publicado (R$)", "Origem"]
    _header_row(ws_n, cols_n)
    for ri, ins in enumerate(sinapi, 2):
        vals = [ins.get("codigo_referencia"), ins.get("descricao"), ins.get("tipo"),
                ins.get("unidade"), ins.get("preco_publicado"), ins.get("origem")]
        for ci, v in enumerate(vals, 1):
            c = ws_n.cell(row=ri, column=ci, value=v)
            c.border = THIN
            if ci == 5 and isinstance(v, (int, float)):
                c.number_format = "R$ #,##0.0000"
    _auto_width(ws_n)

    # ── Aba 4: ANP ──────────────────────────────────────────────────────────────
    ws_a = wb.create_sheet("ANP")
    cols_a = ["Produto", "Unidade", "Preço Médio (R$/L)", "Data Referência", "Versão", "N° Postos"]
    _header_row(ws_a, cols_a)
    if anp:
        for ri, comb in enumerate(anp.get("combustiveis", []), 2):
            vals = [comb.get("produto"), comb.get("unidade"), comb.get("preco_medio_ponderado"),
                    anp.get("data_vigencia"), anp.get("versao"), comb.get("n_postos")]
            for ci, v in enumerate(vals, 1):
                c = ws_a.cell(row=ri, column=ci, value=v)
                c.border = THIN
                if ci == 3 and isinstance(v, (int, float)):
                    c.number_format = "R$ #,##0.0000"

    # ── Aba 5: Catálogo atualizado ──────────────────────────────────────────────
    ws_cat = wb.create_sheet("Catálogo_Atualizado")
    cols_c = ["Código", "Descrição", "Tipo", "Unidade", "Preço Base (R$)", "Fonte", "Data Ref.", "Preço BR-" + uf]
    _header_row(ws_cat, cols_c)
    for ri, ins in enumerate(insumos_cat, 2):
        regz = ins.get("regionalizacao") or {}
        if isinstance(regz, list):
            regz = {r.get("jurisdicao_id", "?"): r for r in regz}
        preco_uf = (regz.get(f"BR-{uf}") or {}).get("preco")
        vals = [
            ins.get("codigo"), ins.get("descricao"), ins.get("tipo"),
            ins.get("unidade"), ins.get("preco_unitario"), ins.get("fonte_preco"),
            ins.get("data_referencia_preco"), preco_uf,
        ]
        for ci, v in enumerate(vals, 1):
            c = ws_cat.cell(row=ri, column=ci, value=v)
            c.border = THIN
            if ci in (5, 8) and isinstance(v, (int, float)):
                c.number_format = "R$ #,##0.0000"
            # Destaca STALE (data antiga)
            if ci == 7 and v and str(v) <= "2025-12-31":
                c.fill = WARN_FILL
    _auto_width(ws_cat)

    # ── Aba 6: Conflitos ────────────────────────────────────────────────────────
    ws_conf = wb.create_sheet("Conflitos")
    cols_cf = ["Código", "Preço Atual", "Preço SICRO", "Δ%", "Score Match", "Ação"]
    _header_row(ws_conf, cols_cf)
    for ri, conf in enumerate(cov.get("conflitos", []), 2):
        vals = [conf.get("concilium_id"), conf.get("preco_atual"), conf.get("preco_sicro"),
                f"{conf.get('variacao_pct', 0):+.1f}%", conf.get("score"), "Revisar"]
        for ci, v in enumerate(vals, 1):
            c = ws_conf.cell(row=ri, column=ci, value=v)
            c.border = THIN; c.fill = WARN_FILL

    # ── Aba 7: Sem Cobertura ────────────────────────────────────────────────────
    ws_nc = wb.create_sheet("Sem_Cobertura")
    cols_nc = ["Código", "Descrição", "Tipo", "Melhor Score", "Motivo"]
    _header_row(ws_nc, cols_nc)
    for ri, nc in enumerate(cov.get("sem_match", []), 2):
        vals = [nc.get("concilium_id"), nc.get("descricao"), nc.get("tipo"),
                nc.get("melhor_score", ""), nc.get("motivo", "sem_match_sicro")]
        for ci, v in enumerate(vals, 1):
            c = ws_nc.cell(row=ri, column=ci, value=v)
            c.border = THIN; c.fill = ERR_FILL

    # Salva
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"pesquisa_{uf}_{mes}.xlsx"
    wb.save(out_path)
    print(f"  Pesquisa salva: {out_path}")
    return out_path


def main() -> None:
    p = argparse.ArgumentParser(description="ADR-045 Stage 3: Gera pesquisa_ESTADO.xlsx")
    p.add_argument("--uf",         required=True, help="UF (ex: CE)")
    p.add_argument("--mes",        default="2026-01", help="AAAA-MM")
    p.add_argument("--contexto",   default="mendeslobo")
    p.add_argument("--output-dir", default="config/referencias/pesquisas", dest="output_dir")
    args = p.parse_args()

    out_dir = ROOT / args.output_dir
    print(f"\n  generate_pesquisa — BR-{args.uf} {args.mes}")
    path = generate(args.uf, args.mes, args.contexto, out_dir)
    print(f"  Concluído: {path}")


if __name__ == "__main__":
    main()
