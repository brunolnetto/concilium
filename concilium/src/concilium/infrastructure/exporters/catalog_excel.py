from pathlib import Path
from typing import Any

import pandas as pd

from concilium.domain.services.simulation import SimuladorEngenharia


def _apply_body_row(row, fill, standard_border, font_name_body: str, size_body: int) -> None:
    from openpyxl.styles import Font

    for cell in row:
        if fill:
            cell.fill = fill
        cell.font = Font(name=font_name_body, size=size_body)
        cell.border = standard_border


class CatalogExporterExcel:
    def __init__(self, catalogo: dict[str, Any], jurisdicao_id: str = None):  # type: ignore[assignment]
        self.catalogo = catalogo
        self.jurisdicao_id = jurisdicao_id
        self.simulador = SimuladorEngenharia(catalogo)

    def _get_preco_regional(self, insumo: dict) -> float:
        """Resolve o preço considerando a regionalização da jurisdição."""
        if not self.jurisdicao_id:
            return insumo.get("preco_unitario")  # type: ignore[return-value]

        regional = insumo.get("regionalizacao", {})
        if self.jurisdicao_id in regional:
            return regional[self.jurisdicao_id].get("preco")

        return insumo.get("preco_unitario")  # type: ignore[return-value]

    def _cex_build_insumos_dfs(self, insumos_raw):
        data_map = {"MATERIAL": [], "MAO_OBRA": [], "EQUIPAMENTO": []}
        for i in insumos_raw:
            tipo = i.get("tipo")
            if tipo not in data_map:
                continue
            row = {
                "Codigo": i.get("codigo"),
                "Descricao": i.get("descricao"),
                "Unidade": i.get("unidade"),
                "Preco Unitario": self._get_preco_regional(i),
                "Categoria": i.get("categoria"),
                "Fonte": i.get("fonte_preco"),
                "Referencia": i.get("data_referencia_preco"),
            }
            if tipo == "EQUIPAMENTO":
                attrs = i.get("atributos_especificos", {})
                decomp = attrs.get("decomposicao_equip", {})
                row.update(
                    {
                        "H/Mes": attrs.get("horas_mes", 200),
                        "Capital (R$/h)": decomp.get("custo_capital_h", 0),
                        "Operador (R$/h)": decomp.get("custo_mao_obra_h", 0),
                        "Diesel (R$/h)": decomp.get("custo_diesel_h", 0),
                        "Manut (R$/h)": decomp.get("custo_material_h", 0),
                    }
                )
            data_map[tipo].append(row)
        return (
            pd.DataFrame(data_map["MATERIAL"]),
            pd.DataFrame(data_map["MAO_OBRA"]),
            pd.DataFrame(data_map["EQUIPAMENTO"]),
        )

    def _cex_build_cpu_row(self, cpu):
        res_clt = self.simulador.simular_cpu(cpu["codigo"], 1, "7x0", 1, "clt")
        res_pj = self.simulador.simular_cpu(cpu["codigo"], 1, "7x0", 1, "pj")
        equipe = cpu.get("equipe") or cpu.get("composicao_padrao_equipe")
        if not equipe:
            prod_id = cpu.get("producao_id")
            ficha = next((p for p in self.catalogo.get("producoes", []) if p.get("id") == prod_id), None)
            if ficha:
                equipe = ficha.get("equipe") or ficha.get("composicao_padrao_equipe")
        equipe = equipe or {}
        equipe_mo = ", ".join([f"{f['funcao_id']}:{f['num_pessoas']}" for f in equipe.get("funcoes", [])])
        equipe_eq = ", ".join([f"{e['equip_id']}:{e['num_unidades']}" for e in equipe.get("equipamentos", [])])
        return {
            "Codigo": cpu["codigo"],
            "Descricao": cpu["descricao"],
            "Unidade": cpu["unidade"],
            "Categoria": cpu["categoria"],
            "Produtividade (H)": cpu.get("producao_horaria", 0),
            "Equipe (MO)": equipe_mo,
            "Equipe (EQ)": equipe_eq,
            "Custo Direto CLT": res_clt.total_direto,
            "Custo Direto PJ": res_pj.total_direto,
            "MO (Base)": res_clt.custo_mao_obra_base,
            "MO (Encargos)": res_clt.custo_mao_obra_encargos,
            "Diesel Total": res_clt.custo_diesel_total,
        }

    def _cex_write_sheets(self, writer, df_cpus, df_materiais, df_mo, df_equips, df_prods):
        if not df_cpus.empty:
            df_cpus.to_excel(writer, sheet_name="Resumo_CPUs", index=False)
        if not df_materiais.empty:
            df_materiais.to_excel(writer, sheet_name="Materiais", index=False)
        if not df_mo.empty:
            df_mo.to_excel(writer, sheet_name="Mao_de_Obra", index=False)
        if not df_equips.empty:
            df_equips.to_excel(writer, sheet_name="Equipamentos", index=False)
        if not df_prods.empty:
            df_prods.to_excel(writer, sheet_name="Metodologias_3NF", index=False)
        self._aplicar_estilo(writer)

    def export(self, output_path: Path):
        # 1. Preparação dos DataFrames de Insumos (Por Tipo)
        insumos_raw = self.catalogo.get("insumos", [])
        df_materiais, df_mo, df_equips = self._cex_build_insumos_dfs(insumos_raw)

        cpu_data = []
        for cpu in self.catalogo.get("cpus", []):
            try:
                cpu_data.append(self._cex_build_cpu_row(cpu))
            except Exception:
                cpu_data.append({"Codigo": cpu["codigo"], "Descricao": cpu["descricao"], "Status": "Erro Calculo"})
        df_cpus = pd.DataFrame(cpu_data)

        prods_raw = self.catalogo.get("producoes", [])
        formatted_prods = []
        for p in prods_raw:
            formatted_prods.append(
                {
                    "ID": p.get("id"),
                    "Descricao": p.get("descricao"),
                    "Produtividade (H)": p.get("valor_horario"),
                    "Unidade": p.get("unidade"),
                    "Categoria": p.get("categoria"),
                }
            )
        df_prods = pd.DataFrame(formatted_prods)

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            self._cex_write_sheets(writer, df_cpus, df_materiais, df_mo, df_equips, df_prods)
        return output_path

    def _aplicar_estilo(self, writer):
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        # Carrega Identidade Visual do Perfil
        idv = self.catalogo.get("identidade_visual", {})
        cores = idv.get("cores", {})
        fontes = idv.get("fontes", {})

        # Configuração de Estilo baseada em Tokens
        header_bg = cores.get("cabecalho_fundo", "1F4E78")
        header_fg = cores.get("cabecalho_texto", "FFFFFF")
        zebra_impar = cores.get("zebra_impar", "F2F2F2")
        cores.get("zebra_par", "FFFFFF")

        font_name_header = fontes.get("cabecalho", "Arial")
        font_name_body = fontes.get("corpo", "Calibri")
        size_header = fontes.get("tamanho_cabecalho", 11)
        size_body = fontes.get("tamanho_corpo", 10)

        header_fill = PatternFill(start_color=header_bg, end_color=header_bg, fill_type="solid")
        header_font = Font(name=font_name_header, color=header_fg, bold=True, size=size_header)

        zebra_fill = PatternFill(start_color=zebra_impar, end_color=zebra_impar, fill_type="solid")
        border_side = Side(style="thin", color=cores.get("borda", "D9D9D9"))
        standard_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

        center_align = Alignment(horizontal="center", vertical="center")

        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]

            # 1. Cabeçalho
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = standard_border

            # 2. Zebra e Bordas no Corpo
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                fill = zebra_fill if row_idx % 2 != 0 else None
                _apply_body_row(row, fill, standard_border, font_name_body, size_body)

            # 3. Auto-filtro
            ws.auto_filter.ref = ws.dimensions

            # 4. Ajuste de largura de colunas
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        val = str(cell.value)
                        if len(val) > max_length:
                            max_length = len(val)
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 70)
                ws.column_dimensions[column].width = adjusted_width
