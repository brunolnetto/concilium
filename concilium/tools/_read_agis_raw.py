"""Extrai composições de times indiretos (QSMS/LAB/TOP/ADM) dos XLSM brutos AGIS."""
import io, sys, json
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT  = Path(__file__).resolve().parents[1]
RAW   = ROOT / "benchmarks/data/AGIS/artifacts/raw"

try:
    import openpyxl
except ImportError:
    print("openpyxl necessário", file=sys.stderr); sys.exit(2)

# Usa o arquivo mais recente (abril 2026)
xlsx_files = sorted(RAW.glob("*.xlsm"), reverse=True)
print(f"Arquivos XLSM encontrados: {len(xlsx_files)}")
for f in xlsx_files:
    print(f"  {f.name} ({f.stat().st_size/1024:.1f} KB)")

if not xlsx_files:
    sys.exit(1)

# Lê o mais recente
target = xlsx_files[0]
print(f"\nLendo: {target.name}")
wb = openpyxl.load_workbook(target, read_only=True, data_only=True, keep_vba=False)
print(f"Abas: {wb.sheetnames}")

# Procura abas relacionadas a indiretos, QSMS, LAB, TOP, ADM
KEYWORDS = {"qsms", "lab", "top", "adm", "indir", "equipe", "mob", "segur", "saude"}

for sheet_name in wb.sheetnames:
    if any(k in sheet_name.lower() for k in KEYWORDS):
        print(f"\n{'='*60}")
        print(f"ABA: {sheet_name!r}")
        ws = wb[sheet_name]
        rows_read = 0
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            non_null = [str(c)[:50] for c in row if c is not None]
            if non_null:
                print(f"  {' | '.join(non_null[:8])}")
            rows_read += 1
            if rows_read >= 60:
                print("  ...")
                break

# Tenta também ler todas as abas buscando padrão de função/salário
print(f"\n{'='*60}")
print("BUSCA POR PADRÃO FUNÇÃO/SALÁRIO EM TODAS AS ABAS:")
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    found_salary = False
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = [c for c in row if c is not None]
        if not cells:
            continue
        # Detecta padrão numérico que se parece com salário
        nums = [c for c in cells if isinstance(c, (int, float)) and 500 < float(c) < 50000]
        strs = [str(c) for c in cells if isinstance(c, str) and len(str(c)) > 3]
        if nums and strs:
            rows.append((strs[:3], nums[:3]))
            found_salary = True
    if found_salary and len(rows) > 2:
        print(f"\nABA {sheet_name!r} — {len(rows)} linhas com padrão salário:")
        for s, n in rows[:15]:
            print(f"  desc={s} | valores={[round(x, 2) for x in n]}")

wb.close()
