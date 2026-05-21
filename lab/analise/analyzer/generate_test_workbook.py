#!/usr/bin/env python3
"""
generate_test_workbook.py
Gera workbooks de teste para benchmarks de arquitetura de planilhas.

Modos:
  1. clean : Nenhuma violação.
  2. auto  : Apenas violações simples (G1 - AUTOMATICA).
  3. mixed : Violações simples, complexas e ciclos (G1, G2, G4).
"""

import openpyxl
import sys
import argparse
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

# ── Estilos ───────────────────────────────────────────────────────────────────

def _thin_border():
    s = Side(style="thin", color="BBBBBB")
    return Border(left=s, right=s, top=s, bottom=s)

def style_header(cell, text, color="1F4E79"):
    cell.value = text
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.alignment = Alignment(horizontal="center", vertical="center")

def style_subheader(cell, text):
    cell.value = text
    cell.font = Font(bold=True, color="1F4E79", size=10)
    cell.fill = PatternFill("solid", fgColor="D6E4F0")
    cell.alignment = Alignment(horizontal="left", vertical="center")

def style_label(cell, text):
    cell.value = text
    cell.font = Font(color="404040", size=10)
    cell.alignment = Alignment(horizontal="left", vertical="center")

def style_value(cell, value, fmt="number", bold=False):
    cell.value = value
    cell.font = Font(bold=bold, size=10)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    if fmt == "brl":
        cell.number_format = 'R$ #,##0.00'
    elif fmt == "pct":
        cell.number_format = '0.00%'
    elif fmt == "number":
        cell.number_format = '#,##0'
    cell.border = _thin_border()

# ── Construtores de Abas ──────────────────────────────────────────────────────

def build_vendas_brutas(wb, mode):
    ws = wb["Vendas_Brutas"]
    ws.merge_cells("B1:D1")
    style_header(ws["B1"], "VENDAS BRUTAS")
    style_subheader(ws["B2"], "Produto")
    style_subheader(ws["C2"], "Receita")
    
    rows = [("Produto A", 1000), ("Produto B", 2000)]
    for i, (p, v) in enumerate(rows, start=3):
        style_label(ws[f"B{i}"], p)
        style_value(ws[f"C{i}"], v, "brl")
    
    style_label(ws["B5"], "TOTAL")
    style_value(ws["C5"], "=SUM(C3:C4)", "brl", bold=True)

    if mode == "mixed":
        # VIOLACAO G2 (S-LOCAL mas cone maior por causa do Resumo)
        style_label(ws["B8"], "[VIO-G2] Ref Resultados")
        style_value(ws["C8"], "=Resumo_Executivo!C6", "brl")

def build_custos_base(wb, mode):
    ws = wb["Custos_Base"]
    ws.merge_cells("B1:C1")
    style_header(ws["B1"], "CUSTOS BASE", color="7B3F00")
    
    style_label(ws["B3"], "Custo Fixo")
    style_value(ws["C3"], 500, "brl")

    if mode in ["auto", "mixed"]:
        # VIOLACAO G1 (AUTOMATICA)
        style_label(ws["B8"], "[VIO-G1] Ref Calculos")
        style_value(ws["C8"], "=Margem_Contribuicao!C6", "brl")

def build_margem_contribuicao(wb, mode):
    ws = wb["Margem_Contribuicao"]
    ws.merge_cells("B1:C1")
    style_header(ws["B1"], "MARGEM CONTRIBUIÇÃO", color="1B5E20")
    
    style_label(ws["B3"], "Receita")
    style_value(ws["C3"], "=Vendas_Brutas!C5", "brl")
    style_label(ws["B4"], "Custos")
    style_value(ws["C4"], "=Custos_Base!C3", "brl")
    style_label(ws["B6"], "MARGEM")
    style_value(ws["C6"], "=C3-C4", "brl", bold=True)

    if mode == "mixed":
        # VIOLACAO G4 (CYCLE)
        style_label(ws["B9"], "[VIO-G4] Ciclo")
        style_value(ws["C9"], "=Impostos!C10 * 0.5", "brl")

def build_impostos(wb, mode):
    ws = wb["Impostos"]
    ws.merge_cells("B1:C1")
    style_header(ws["B1"], "IMPOSTOS", color="4A148C")
    
    style_label(ws["B3"], "Base")
    style_value(ws["C3"], "=Vendas_Brutas!C5", "brl")
    style_label(ws["B7"], "TOTAL")
    style_value(ws["C7"], "=C3*0.15", "brl", bold=True)

    if mode in ["auto", "mixed"]:
        # VIOLACAO G1 (AUTOMATICA)
        style_label(ws["B9"], "[VIO-G1] Ref Dashboard")
        style_value(ws["C9"], "=Dashboard_Financeiro!C7 * 0.01", "brl")

    if mode == "mixed":
        # VIOLACAO G4 (CYCLE)
        style_label(ws["B10"], "[VIO-G4] Ciclo Volta")
        style_value(ws["C10"], "=Margem_Contribuicao!C9 * 1.1", "brl")

def build_dashboard_financeiro(wb, mode):
    ws = wb["Dashboard_Financeiro"]
    ws.merge_cells("B1:C1")
    style_header(ws["B1"], "DASHBOARD", color="BF360C")
    style_label(ws["B3"], "Margem")
    style_value(ws["C3"], "=Margem_Contribuicao!C6", "brl")
    style_label(ws["B7"], "KPI")
    style_value(ws["C7"], "=C3*1.1", "brl", bold=True)

def build_resumo_executivo(wb, mode):
    ws = wb["Resumo_Executivo"]
    ws.merge_cells("B1:C1")
    style_header(ws["B1"], "RESUMO", color="006064")
    style_label(ws["B6"], "RESULTADO")
    style_value(ws["C6"], "=Dashboard_Financeiro!C7", "brl", bold=True)

# ── Main ──────────────────────────────────────────────────────────────────────

def generate():
    parser = argparse.ArgumentParser(description="Gera workbooks de teste para benchmark.")
    parser.add_argument("mode", choices=["clean", "auto", "mixed"], help="Modo do benchmark")
    parser.add_argument("--out", default="test_benchmark.xlsx", help="Nome do arquivo de saída")
    args = parser.parse_args()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sheets = ["Vendas_Brutas", "Custos_Base", "Margem_Contribuicao", "Impostos", "Dashboard_Financeiro", "Resumo_Executivo"]
    for s in sheets:
        wb.create_sheet(s)

    build_vendas_brutas(wb, args.mode)
    build_custos_base(wb, args.mode)
    build_margem_contribuicao(wb, args.mode)
    build_impostos(wb, args.mode)
    build_dashboard_financeiro(wb, args.mode)
    build_resumo_executivo(wb, args.mode)

    wb.save(args.out)
    print(f"[OK] Benchmark '{args.mode}' gerado: {args.out}")

if __name__ == "__main__":
    generate()
