#!/usr/bin/env python3
import openpyxl
import re
import os
import sys
import argparse
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor
from functools import partial

# Configuração de Regex para extração de referências (suporta abas com espaços e acentos)
REF_PAT = re.compile(
    r"(?:'([^']+)'|([A-Za-z_\u00C0-\u024F][\w.\u00C0-\u024F]*))?!(\$?[A-Z]{1,3}\$?\d{1,7})|(\$?[A-Z]{1,3}\$?\d{1,7})", 
    re.IGNORECASE
)

def evaluate_smart(expr, values_by_sheet, current_sheet):
    """
    Simula o cálculo do Excel para fórmulas simples (aritmética e referências).
    Retorna o valor calculado ou uma string [COMPLEX] se houver funções complexas.
    """
    clean_expr = str(expr).strip().lstrip('=')
    if not clean_expr: return ""
    
    # Se for apenas uma string entre aspas
    if clean_expr.startswith('"') and clean_expr.endswith('"'):
        return clean_expr.replace('"', '')

    def replace_ref(m):
        sheet_name = m.group(1) or m.group(2) or current_sheet
        addr = (m.group(3) or m.group(4)).replace("$", "")
        sheet_data = values_by_sheet.get(sheet_name, {})
        val = sheet_data.get(addr, 0)
        return str(val) if val is not None else "0"

    # Substitui referências pelos seus valores reais do arquivo original
    eval_str = REF_PAT.sub(replace_ref, clean_expr)
    
    # Se contém funções do Excel, não tentamos avaliar (risco de erro de lógica)
    if "(" in eval_str and any(func in eval_str.upper() for func in ["INT", "ROUND", "SUM", "IF", "CONCAT", "VLOOKUP", "PROCV"]):
        return f"[COMPLEX] {eval_str}"

    try:
        # Filtra apenas caracteres seguros para eval()
        safe_str = re.sub(r"[^0-9.+\-*/() ]", "", eval_str)
        if not safe_str.strip(): return eval_str
        return eval(safe_str)
    except:
        return eval_str

def check_cell_task(cell_info, values_by_sheet, fixed_formulas):
    """Worker process: Valida uma única célula."""
    sheet_name, addr, v_orig = cell_info
    f_new = fixed_formulas.get((sheet_name, addr))
    
    if isinstance(f_new, str) and f_new.startswith("="):
        v_fixed = evaluate_smart(f_new, values_by_sheet, sheet_name)
    else:
        v_fixed = f_new

    match = False
    is_complex = str(v_fixed).startswith("[COMPLEX]")
    
    if is_complex:
        return (sheet_name, addr, v_orig, v_fixed, "ASSISTIDA")

    # Comparação numérica com tolerância para ponto flutuante
    if isinstance(v_orig, (int, float)) and isinstance(v_fixed, (int, float)):
        match = abs(v_orig - v_fixed) < 1e-4
    else:
        match = str(v_orig).strip() == str(v_fixed).strip().replace('"', '')

    status = "OK" if match else "DIVERGENTE"
    return (sheet_name, addr, v_orig, v_fixed, status)

def validate_files(orig_path, fixed_path, report_path):
    print(f"\n  [I] Auditoria: {orig_path.name} vs {fixed_path.name}")
    
    # Carrega arquivos
    print("      Carregando valores originais (data_only)...")
    wb_orig = openpyxl.load_workbook(str(orig_path), data_only=True, read_only=True)
    
    print("      Carregando fórmulas corrigidas...")
    wb_new = openpyxl.load_workbook(str(fixed_path), data_only=False, read_only=True)
    
    values_by_sheet = {}
    fixed_formulas = {}
    task_list = []

    print("      Mapeando células para validação...")
    for sheet_name in wb_new.sheetnames:
        if sheet_name not in wb_orig.sheetnames: continue
        
        ws_orig = wb_orig[sheet_name]
        ws_new = wb_new[sheet_name]
        
        # Cache de valores originais
        values_by_sheet[sheet_name] = {c.coordinate: c.value for r in ws_orig.iter_rows() for c in r if c.value is not None}
        
        # Coleta fórmulas do arquivo novo
        for row in ws_new.iter_rows():
            for cell in row:
                val = cell.value
                if val is not None:
                    addr = cell.coordinate
                    fixed_formulas[(sheet_name, addr)] = val
                    
                    # Se é uma fórmula, entra na lista de auditoria
                    if isinstance(val, str) and val.startswith("="):
                        v_orig = values_by_sheet[sheet_name].get(addr)
                        task_list.append((sheet_name, addr, v_orig))

    print(f"      Total de fórmulas para validar: {len(task_list)}")
    
    worker_func = partial(check_cell_task, values_by_sheet=values_by_sheet, fixed_formulas=fixed_formulas)
    results_summary = {"OK": 0, "DIVERGENTE": 0, "ASSISTIDA": 0}
    
    # Grava relatório
    with open(report_path, "w", encoding="utf-8") as f:
        f.write("RELATÓRIO DE AUDITORIA DE CONSISTÊNCIA NUMÉRICA\n")
        f.write("="*80 + "\n")
        f.write(f"Original : {orig_path}\n")
        f.write(f"Corrigido: {fixed_path}\n")
        f.write("="*80 + "\n\n")
        
        with ProcessPoolExecutor() as executor:
            for s, a, v_o, v_f, status in executor.map(worker_func, task_list):
                results_summary[status] += 1
                if status != "OK":
                    f.write(f"[{status}] {s}!{a}\n")
                    f.write(f"   Valor Original  : {v_o}\n")
                    f.write(f"   Valor Calculado : {v_f}\n\n")
        
        f.write("\n" + "="*80 + "\n")
        f.write(f"RESUMO DA AUDITORIA:\n")
        f.write(f"  ✅ Sucesso (OK)         : {results_summary['OK']}\n")
        f.write(f"  ⚠️ Complexas (Assistida) : {results_summary['ASSISTIDA']}\n")
        f.write(f"  ❌ Falhas (Divergente)   : {results_summary['DIVERGENTE']}\n")
        f.write("="*80 + "\n")

    print(f"\n  [OK] Auditoria concluída!")
    print(f"       ✅ Sucesso: {results_summary['OK']}")
    print(f"       ⚠️ Complexas: {results_summary['ASSISTIDA']} (Requer conferência manual)")
    print(f"       ❌ Falhas: {results_summary['DIVERGENTE']}")
    print(f"\n  [>] Relatório detalhado: {report_path.resolve()}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Auditoria de consistência entre arquivos Excel.")
    parser.add_argument("original", help="Caminho do arquivo original (referência)")
    parser.add_argument("corrigido", help="Caminho do arquivo corrigido")
    parser.add_argument("--report", default="data/auditoria_consistencia.txt", help="Caminho do relatório de saída")
    
    args = parser.parse_args()
    
    # Resolução de caminhos para WSL
    orig = Path(os.path.abspath(os.path.join(os.getcwd(), args.original)))
    fixed = Path(os.path.abspath(os.path.join(os.getcwd(), args.corrigido)))
    report = Path(os.path.abspath(os.path.join(os.getcwd(), args.report)))
    
    if not orig.exists() or not fixed.exists():
        print(f"  [!] ERRO: Arquivos de entrada não encontrados.")
        sys.exit(1)
        
    validate_files(orig, fixed, report)
