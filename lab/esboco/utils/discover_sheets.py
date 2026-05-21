#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Discovery Script: List all sheets in the actual AGIS spreadsheet
"""

import sys
from pathlib import Path
import shutil
import tempfile
from openpyxl import load_workbook

# Spreadsheet path
SPREADSHEET = r"C:\Users\GustavoArruda\Mendes Lobo\TI - Documentos\IA\planilha_investimento\2026_03_30_PLANEJAMENTO_LOTE 09 e 10_(CE).xlsm"

def discover_sheets():
    """List all sheets in the workbook"""
    
    original = Path(SPREADSHEET)
    if not original.exists():
        print(f"ERROR: File not found: {SPREADSHEET}")
        return
    
    # Try to create a copy for safer reading
    try:
        temp_dir = Path(tempfile.gettempdir()) / "agis_skills"
        temp_dir.mkdir(exist_ok=True)
        copy_path = temp_dir / original.name
        
        print(f"[*] Creating temporary copy...")
        shutil.copy2(original, copy_path)
        working_file = copy_path
    except:
        print(f"[!] Could not copy, using original file")
        working_file = original
    
    try:
        print(f"\n[+] Loading workbook: {Path(working_file).name}")
        wb = load_workbook(working_file, data_only=False)
        
        print(f"\n[OK] Total sheets: {len(wb.sheetnames)}\n")
        print("=" * 80)
        print("ALL AVAILABLE SHEETS IN WORKBOOK")
        print("=" * 80)
        
        for i, sheet_name in enumerate(wb.sheetnames, 1):
            ws = wb[sheet_name]
            print(f"\n{i:2d}. {sheet_name:30s} | Rows: {ws.max_row:6d} | Cols: {ws.max_column:3d}")
            
            # Show first row content (headers)
            headers = []
            for col in range(1, min(ws.max_column + 1, 6)):
                cell_value = ws.cell(1, col).value
                if cell_value:
                    headers.append(str(cell_value)[:20])
            
            if headers:
                print(f"    Headers: {', '.join(headers)}...")
        
        print("\n" + "=" * 80)
        print("SHEET NAMES (for import references)")
        print("=" * 80)
        print("\nPython list format:")
        print(f"sheets = {wb.sheetnames}")
        
        print("\nImportable names:")
        for sheet_name in wb.sheetnames:
            sanitized = sheet_name.replace(" ", "_").replace("-", "_").replace("/", "_")
            print(f"  - {sheet_name:30s} -> {sanitized}")
        
        wb.close()
        
    except Exception as e:
        print(f"ERROR reading workbook: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    discover_sheets()
