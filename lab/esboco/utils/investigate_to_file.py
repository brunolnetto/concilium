#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Investigation with file output"""

from pathlib import Path
import tempfile
import shutil
from openpyxl import load_workbook

SPREADSHEET = r"C:\Users\GustavoArruda\Mendes Lobo\TI - Documentos\IA\planilha_investimento\2026_03_30_PLANEJAMENTO_LOTE 09 e 10_(CE).xlsm"

src = Path(SPREADSHEET)
dst = Path(tempfile.gettempdir()) / "agis_skills" / "inv2.xlsm"
dst.parent.mkdir(exist_ok=True)

output_file = Path(SPREADSHEET).parent / "INVESTIGATION_RESULTS.txt"

with open(output_file, "w", encoding="utf-8") as out:
    out.write("[*] Copying file...\n")
    shutil.copy2(src, dst)
    
    out.write("[+] Loading workbook...\n")
    out.flush()
    
    wb = load_workbook(dst, data_only=True)
    
    mystery_sheets = ["CPU", "MATERIAIS", "ENCARGOS", "DD - C. EXAMES E EPI", "Transportes"]
    
    for sheet_name in mystery_sheets:
        if sheet_name not in wb.sheetnames:
            out.write(f"[SKIP] {sheet_name}\n")
            continue
        
        ws = wb[sheet_name]
        out.write(f"\n{'='*80}\n")
        out.write(f"{sheet_name} - {ws.max_row} rows x {ws.max_column} cols\n")
        out.write(f"{'='*80}\n\n")
        
        # Print first 8 rows
        for row_idx in range(1, min(9, ws.max_row + 1)):
            row_vals = []
            for col_idx in range(1, min(8, ws.max_column + 1)):
                val = ws.cell(row_idx, col_idx).value
                row_vals.append(str(val)[:12] if val else "[empty]")
            out.write(" | ".join(f"{v:12s}" for v in row_vals) + "\n")
        out.flush()
    
    # Check PLANEJAMENTO line 52
    out.write(f"\n{'='*80}\n")
    out.write("PLANEJAMENTO - Row 52\n")
    out.write(f"{'='*80}\n\n")
    
    ws = wb["PLANEJAMENTO"]
    for col_idx in range(1, min(16, ws.max_column + 1)):
        val = ws.cell(52, col_idx).value
        col_letter = ws.cell(52, col_idx).column_letter
        out.write(f"{col_letter:3s} | {str(val)[:40] if val else '[empty]'}\n")
    out.flush()
    
    out.write("[OK] Done\n")
    
    wb.close()

print(f"Results written to: {output_file}")
