#!/usr/bin/env python3
"""
Merge entre catálogo Concilium e PLANILHA PADRÃO RFI.xlsm

Uso:
    python scripts/benchmark/merge_rfi.py --contexto mendeslobo          # dry-run (relatório)
    python scripts/benchmark/merge_rfi.py --contexto mendeslobo --apply  # aplica merge

Este script é o adaptador específico do formato "PLANILHA PADRÃO RFI.xlsm" — sabe ler as
abas BD_MÃO_DE_OBRA/BD_EQUIPAMENTOS/BD_MATERIAIS/SERVIÇOS TERCEIRIZADOS/BD_PRODUTIVIDADES
e convertê-las para o formato de entidades do catálogo (`extract_rfi`). O motor de
comparação/merge genérico (agnóstico de origem) vive em
`concilium.application.services.catalog.catalog_compare` (issue #134) — este script apenas
consome esse motor para o caso de uso "catálogo vs planilha RFI". Para comparar/mesclar
dois catálogos entre si (ou catálogo vs planilha), use `concilium catalog compare|merge`.

Operações de comparação (como JOIN):
    LEFT_ONLY   : no catálogo, não na planilha   → mantém (sem alteração)
    RIGHT_ONLY  : na planilha, não no catálogo   → candidato a inclusão
    INNER_HIGH  : match confiante (score ≥ 0.85) → atualiza preço/dados se diferente
    INNER_LOW   : match incerto  (0.50–0.85)     → reporta para revisão manual
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

_SRC_ROOT = Path(__file__).resolve().parents[2] / "src"
if str(_SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(_SRC_ROOT))

from concilium.application.services.catalog.catalog_compare import (  # noqa: E402
    THRESHOLD_HIGH,
    THRESHOLD_LOW,
    RelatorioComparacao as MergeReport,
    _norm_unit,
    aplicar_merge as _apply_merge,
    comparar_entidades as _merge_entities,
)

ROOT = Path(__file__).resolve().parents[2]
XLSM = ROOT / "docs/product/3-reference_files/catalogos/PLANILHA PADRÃO RFI.xlsm"
HEADER_ROW = 3  # linha real do header nas abas BD_* (1-indexed)

# ── Mapeamento de fonte ───────────────────────────────────────────────────────
_FONTE_MAP = {
    "cotação mercado": "cotacao_mercado",
    "cotacao mercado": "cotacao_mercado",
    "sinapi": "tabela_referencia",
    "sicro": "tabela_referencia",
    "tabela referencia": "tabela_referencia",
    "negociado": "negociado",
    "catalogo": "catalogo_proprio",
    "proprio": "catalogo_proprio",
}

def _map_fonte(raw: str | None) -> str:
    if not raw:
        return "catalogo_proprio"
    key = unicodedata.normalize("NFD", str(raw).lower())
    key = "".join(c for c in key if unicodedata.category(c) != "Mn")
    for k, v in _FONTE_MAP.items():
        if k in key:
            return v
    return "catalogo_proprio"

def _fmt_date(val: Any) -> str | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val)
    m = re.match(r"(\d{4}-\d{2}-\d{2})", s)
    return m.group(1) if m else None

# ── Parsing de composições  "P0062×1, P0001×3" ───────────────────────────────
def _parse_mo(raw: str | None) -> list[dict]:
    if not raw or str(raw).strip() in ("—", "-", "", "None"):
        return []
    items = []
    for part in re.split(r"[,;]", str(raw)):
        m = re.match(r"\s*([A-Z0-9]+)\s*[x×]\s*(\d+)", part.strip(), re.IGNORECASE)
        if m:
            items.append({"funcao_id": m.group(1), "num_pessoas": int(m.group(2))})
    return items

def _parse_eq(raw: str | None) -> list[dict]:
    if not raw or str(raw).strip() in ("—", "-", "", "None"):
        return []
    items = []
    for part in re.split(r"[,;]", str(raw)):
        m = re.match(r"\s*([A-Z0-9]+)\s*[x×]\s*(\d+)", part.strip(), re.IGNORECASE)
        if m:
            items.append({"equip_id": f"EQ-{m.group(1)}", "num_unidades": int(m.group(2))})
    return items

def _parse_mat(raw: str | None) -> list[dict]:
    if not raw or str(raw).strip() in ("—", "-", "", "None"):
        return []
    items = []
    for part in re.split(r"[,;]", str(raw)):
        m = re.match(r"\s*([A-Z0-9]+)\s*[x×]\s*([\d.]+)", part.strip(), re.IGNORECASE)
        if m:
            items.append({"insumo_id": f"MA-{m.group(1)}", "consumo_por_hora": float(m.group(2)), "unidade": "un"})
    return items

# ── Extractor ─────────────────────────────────────────────────────────────────

def _read_sheet(wb: Any, name: str, header_row: int = HEADER_ROW) -> list[dict]:
    """Lê uma aba como lista de dicts (ignora linhas all-None)."""
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else f"_col{i}" for i, h in enumerate(rows[0])]
    result = []
    for row in rows[1:]:
        if all(c is None for c in row):
            continue
        d = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        result.append(d)
    return result

def _clean_entity(e: dict) -> dict:
    """Normaliza unidade e remove espaços extras de descrição."""
    if "unidade" in e:
        e["unidade"] = _norm_unit(e["unidade"]) or e["unidade"]
    if "descricao" in e and e["descricao"]:
        e["descricao"] = e["descricao"].strip()
    return e

def _is_valid(e: dict, min_desc_len: int = 3) -> bool:
    """Filtra entidades sem descrição real (linhas em branco/incompletas da planilha)."""
    desc = (e.get("descricao") or "").strip()
    return len(desc) >= min_desc_len and bool(e.get("codigo", "").strip())

def extract_rfi(wb: Any) -> dict:
    """Extrai todas as entidades da planilha RFI no formato do catálogo."""
    insumos: list[dict] = []
    producoes: list[dict] = []

    # ── MAO_OBRA ──────────────────────────────────────────────────────────────
    for row in _read_sheet(wb, "BD_MÃO_DE_OBRA"):
        cod = str(row.get("Código") or "").strip()
        if not cod or cod.lower() == "código":
            continue
        descricao = str(row.get("Função") or "").strip()
        tipo_ct = str(row.get("Tipo\n(CLT/PJ)") or row.get("Tipo") or "CLT").strip().upper()
        sal_mensal = row.get("Salário\nMensal (R$)") or row.get("Salário Mensal (R$)") or 0
        horas_mes = row.get("Horas\n/Mês") or row.get("Horas/Mês") or 220
        encargos = row.get("Encargos\nAplicados") or row.get("Encargos Aplicados") or 0
        data_ref = _fmt_date(row.get("Data Atualização"))
        try:
            sal_mensal = float(sal_mensal)
            horas_mes = float(horas_mes)
            encargos = float(encargos)
        except (TypeError, ValueError):
            sal_mensal = 0; horas_mes = 220; encargos = 0
        insumos.append({
            "codigo": cod,
            "tipo": "MAO_OBRA",
            "descricao": descricao,
            "unidade": "mes",
            "preco_unitario": round(sal_mensal, 4),
            "fonte_preco": "negociado",
            "data_referencia_preco": data_ref,
            "atributos_especificos": {
                "tipo_contrato": "CLT" if "CLT" in tipo_ct else "PJ",
                "salario_mensal": sal_mensal,
                "horas_mes": horas_mes,
                "percentual_encargos_local": encargos,
            },
            "_fonte_rfi": "BD_MÃO_DE_OBRA",
        })

    # ── EQUIPAMENTO ──────────────────────────────────────────────────────────
    for row in _read_sheet(wb, "BD_EQUIPAMENTOS"):
        cod = str(row.get("Código") or "").strip()
        if not cod or cod.lower() == "código":
            continue
        descricao = str(row.get("Descrição") or "").strip()
        tipo_prop_raw = str(row.get("Tipo\n(próprio/loc.)") or row.get("Tipo") or "locação").strip().lower()
        tipo_prop = "locacao" if "loc" in tipo_prop_raw else "proprio"
        custo_prod = row.get("Custo Produtivo\n(R$/h)") or row.get("Custo Produtivo (R$/h)") or 0
        pct_manut  = row.get("% Manutenção") or 0
        custo_mensal = row.get("Custo Mensal\n(R$/mês)") or row.get("Custo Mensal (R$/mês)") or 0
        data_ref = _fmt_date(row.get("Data Atualização"))
        try:
            custo_prod = float(custo_prod); pct_manut = float(pct_manut); custo_mensal = float(custo_mensal)
        except (TypeError, ValueError):
            custo_prod = 0; pct_manut = 0; custo_mensal = 0
        # codigo: E0001 → EQ-E0001
        codigo_cat = f"EQ-{cod}" if not cod.startswith("EQ-") else cod
        insumos.append({
            "codigo": codigo_cat,
            "tipo": "EQUIPAMENTO",
            "descricao": descricao,
            "unidade": "h",
            "preco_unitario": round(custo_prod, 4),
            "fonte_preco": "catalogo_proprio",
            "data_referencia_preco": data_ref,
            "atributos_especificos": {
                "tipo_propriedade": tipo_prop,
                "valor_locacao_mensal": round(custo_mensal, 4) if custo_mensal else None,
                "pct_manutencao": pct_manut if pct_manut else None,
            },
            "_fonte_rfi": "BD_EQUIPAMENTOS",
        })

    # ── MATERIAL ──────────────────────────────────────────────────────────────
    for row in _read_sheet(wb, "BD_MATERIAIS"):
        cod = str(row.get("Código") or "").strip()
        if not cod or cod.lower() == "código":
            continue
        descricao = str(row.get("Descrição") or "").strip()
        unidade   = str(row.get("Unidade") or "un").strip()
        preco     = row.get("Preço Unitário\n(R$)") or row.get("Preço Unitário (R$)") or 0
        fornec    = str(row.get("Fornecedor\nReferência") or row.get("Fornecedor Referência") or "").strip()
        fonte_raw = str(row.get("Fonte/Origem") or "").strip()
        data_ref  = _fmt_date(row.get("Data Atualização"))
        try:
            preco = float(preco)
        except (TypeError, ValueError):
            preco = 0
        codigo_cat = f"MA-{cod}" if not cod.startswith("MA-") else cod
        insumos.append({
            "codigo": codigo_cat,
            "tipo": "MATERIAL",
            "descricao": descricao,
            "unidade": unidade,
            "preco_unitario": round(preco, 4),
            "fonte_preco": _map_fonte(fonte_raw),
            "data_referencia_preco": data_ref,
            "atributos_especificos": {
                "fornecedor_referencia": fornec or None,
            },
            "_fonte_rfi": "BD_MATERIAIS",
        })

    # ── TERCEIRIZADOS ─────────────────────────────────────────────────────────
    for row in _read_sheet(wb, "SERVIÇOS TERCEIRIZADOS", header_row=2):
        cod = str(row.get("CÓDIGO") or "").strip()
        if not cod or cod.lower() in ("código", ""):
            continue
        descricao = str(row.get("DESCRIÇÃO") or "").strip()
        unidade   = str(row.get("UNID.") or "un").strip()
        preco     = row.get("R$ UNITÁRIO") or 0
        try:
            preco = float(preco)
        except (TypeError, ValueError):
            preco = 0
        insumos.append({
            "codigo": cod,
            "tipo": "TERCEIRIZADO",
            "descricao": descricao,
            "unidade": unidade,
            "preco_unitario": round(preco, 4),
            "fonte_preco": "negociado",
            "_fonte_rfi": "SERVIÇOS TERCEIRIZADOS",
        })

    # ── PRODUTIVIDADES ────────────────────────────────────────────────────────
    for row in _read_sheet(wb, "BD_PRODUTIVIDADES"):
        cod = str(row.get("Código") or "").strip()
        if not cod or cod.lower() == "código":
            continue
        descricao = str(row.get("Descrição do Serviço") or "").strip()
        unidade   = str(row.get("Unidade") or "un").strip()
        prod_raw  = row.get("Produtividade\n(un/h)") or row.get("Produtividade (un/h)") or 0
        mo_comp   = str(row.get("Composição MO\n(cód. BD_MO)") or row.get("Composição MO (cód. BD_MO)") or "").strip()
        eq_comp   = str(row.get("Composição Equip.\n(cód. BD_EQ)") or row.get("Composição Equip. (cód. BD_EQ)") or "").strip()
        mat_comp  = str(row.get("Material Principal\n(cód. BD_MAT)") or row.get("Material Principal (cód. BD_MAT)") or "").strip()
        try:
            valor_horario = float(prod_raw)
        except (TypeError, ValueError):
            valor_horario = 0
        funcoes      = _parse_mo(mo_comp)
        equipamentos = _parse_eq(eq_comp)
        materiais    = _parse_mat(mat_comp)
        producoes.append({
            "id": cod,
            "codigo": cod,
            "descricao": descricao,
            "valor_horario": round(valor_horario, 6) if valor_horario else None,
            "unidade": unidade,
            "origem": "PROPRIA",
            "composicao_padrao_equipe": {
                "funcoes": funcoes,
                "equipamentos": equipamentos,
                "materiais": materiais,
            },
            "_fonte_rfi": "BD_PRODUTIVIDADES",
        })

    # Normaliza e filtra entidades incompletas
    insumos   = [_clean_entity(e) for e in insumos   if _is_valid(e)]
    producoes = [_clean_entity(e) for e in producoes if _is_valid(e)]

    return {"insumos": insumos, "producoes": producoes}

# ── Report printer ────────────────────────────────────────────────────────────

def _print_report(report: MergeReport, label: str) -> None:
    print(f"\n{'─'*65}")
    print(f"  {label}")
    print(f"{'─'*65}")
    print(f"  INNER_HIGH  (match confiante ≥{THRESHOLD_HIGH}): {len(report.inner_high):4d}")
    print(f"  INNER_LOW   (match incerto  ≥{THRESHOLD_LOW}): {len(report.inner_low):4d}  ← requer revisão")
    print(f"  RIGHT_ONLY  (novo na planilha)              : {len(report.right_only):4d}  ← candidatos a inclusão")
    print(f"  LEFT_ONLY   (só no catálogo)                : {len(report.left_only):4d}  ← mantidos")

    if report.inner_high:
        with_delta = [i for i in report.inner_high if i["delta"]]
        print(f"\n  INNER_HIGH com diferenças ({len(with_delta)}):")
        for item in with_delta[:20]:
            cat = item["catalog"]
            rfi = item["rfi"]
            print(f"    [{item['score']:.2f}] {cat.get('codigo'):25} {cat.get('descricao','')[:35]}")
            for field, chg in item["delta"].items():
                pct = f"  Δ{chg['delta_pct']:+.1f}%" if "delta_pct" in chg else ""
                print(f"           {field}: {str(chg['catalog'])[:20]} → {str(chg['rfi'])[:20]}{pct}")
        if len(with_delta) > 20:
            print(f"    ... e mais {len(with_delta)-20} com diferenças")

    if report.inner_low:
        print(f"\n  INNER_LOW — revisar manualmente ({min(len(report.inner_low), 15)} de {len(report.inner_low)}):")
        for item in report.inner_low[:15]:
            cat = item["catalog"]
            rfi = item["rfi"]
            print(f"    [{item['score']:.2f}] CATALOGO: {cat.get('codigo'):20} {cat.get('descricao','')[:30]}")
            print(f"           RFI:      {rfi.get('codigo'):20} {rfi.get('descricao','')[:30]}")

    if report.right_only:
        print(f"\n  RIGHT_ONLY — novos na planilha ({min(len(report.right_only), 15)} de {len(report.right_only)}):")
        for item in report.right_only[:15]:
            rfi = item["rfi"]
            print(f"    {rfi.get('codigo'):25} {rfi.get('descricao','')[:40]}  [{rfi.get('unidade')}]")
        if len(report.right_only) > 15:
            print(f"    ... e mais {len(report.right_only)-15}")

# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    # Reconfigura stdout/stderr para UTF-8 (necessário em consoles Windows) e garante
    # openpyxl disponível — só ao rodar como script standalone, nunca como side-effect
    # de import: este módulo também é carregado via importlib por catalog_compare.py's
    # consumidores (CLI `catalog compare/merge`, MCP tools, testes), e reconfigurar
    # sys.stdout/stderr globalmente nesses contextos quebra a captura de output deles.
    global openpyxl
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")
    try:
        import openpyxl
    except ImportError:
        print("openpyxl não encontrado. Execute: uv add openpyxl", file=sys.stderr)
        sys.exit(2)

    parser = argparse.ArgumentParser(description="Merge catálogo Concilium ↔ PLANILHA PADRÃO RFI")
    parser.add_argument("--contexto", required=True, help="ID do contexto (ex: mendeslobo)")
    parser.add_argument("--apply", action="store_true", help="Aplica merge (padrão: dry-run)")
    parser.add_argument("--include-right-only", action="store_true", default=True,
                        help="Inclui RIGHT_ONLY (novos da planilha) ao catálogo")
    parser.add_argument("--export-rfi", help="Caminho para salvar catálogo extraído da RFI (JSON)")
    args = parser.parse_args()

    mode = "APPLY" if args.apply else "DRY-RUN"
    print(f"\n{'='*65}")
    print(f"  Merge Catálogo ↔ RFI [{mode}] — contexto: {args.contexto}")
    print(f"{'='*65}")

    if not XLSM.exists():
        print(f"[ERRO] Planilha não encontrada: {XLSM}", file=sys.stderr)
        sys.exit(1)

    # 1. Carrega bundle existente
    bundle_path = ROOT / f"config/catalogos/{args.contexto}/.catalog_bundle.json"
    if not bundle_path.exists():
        print(f"[ERRO] Bundle não encontrado: {bundle_path}", file=sys.stderr)
        sys.exit(1)
    bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
    cat_insumos  = bundle.get("insumos", [])
    cat_producoes = bundle.get("producoes", []) or bundle.get("produtividades", [])

    print(f"\n  Catálogo atual: {len(cat_insumos)} insumos, {len(cat_producoes)} producoes")

    # 2. Extrai da planilha RFI
    print(f"  Lendo planilha: {XLSM.name} ...")
    wb = openpyxl.load_workbook(XLSM, data_only=True, keep_vba=False)
    rfi = extract_rfi(wb)
    wb.close()
    rfi_insumos   = rfi["insumos"]
    rfi_producoes = rfi["producoes"]
    print(f"  Planilha RFI:   {len(rfi_insumos)} insumos, {len(rfi_producoes)} producoes")

    if args.export_rfi:
        Path(args.export_rfi).write_text(
            json.dumps(rfi, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        print(f"  RFI extraída salva em: {args.export_rfi}")

    # 3. Merge de insumos
    report_ins, _ = _merge_entities(cat_insumos, rfi_insumos)
    _print_report(report_ins, "INSUMOS")

    # 4. Merge de producoes
    report_prod, _ = _merge_entities(cat_producoes, rfi_producoes, id_key="id")
    _print_report(report_prod, "PRODUCOES")

    # 5. Resumo e decisão
    total_high_delta = sum(1 for i in report_ins.inner_high if i["delta"]) + \
                       sum(1 for i in report_prod.inner_high if i["delta"])
    total_new = len(report_ins.right_only) + len(report_prod.right_only)
    total_low = len(report_ins.inner_low) + len(report_prod.inner_low)

    print(f"\n{'='*65}")
    print(f"  RESUMO")
    print(f"{'='*65}")
    print(f"  Atualizações de preço/dados (INNER_HIGH com delta): {total_high_delta}")
    print(f"  Novos candidatos a inclusão (RIGHT_ONLY):           {total_new}")
    print(f"  Matches incertos para revisão (INNER_LOW):          {total_low}")
    print()

    if not args.apply:
        print("  Modo DRY-RUN — nenhuma alteração feita.")
        print("  Para aplicar: python scripts/benchmark/merge_rfi.py --contexto mendeslobo --apply")
        return

    # 6. Aplica merge
    print("  Aplicando merge...")
    new_insumos   = _apply_merge(cat_insumos, rfi_insumos, report_ins,
                                  incluir_novos=args.include_right_only)
    new_producoes = _apply_merge(cat_producoes, rfi_producoes, report_prod,
                                  id_key="id", incluir_novos=args.include_right_only)

    if "producoes" in bundle:
        bundle["producoes"] = new_producoes
    else:
        bundle["produtividades"] = new_producoes
    bundle["insumos"] = new_insumos

    bundle_path.write_text(json.dumps(bundle, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"  Bundle atualizado: {bundle_path}")
    print(f"  Total insumos:   {len(new_insumos)}  (era {len(cat_insumos)})")
    print(f"  Total producoes: {len(new_producoes)}  (era {len(cat_producoes)})")


if __name__ == "__main__":
    main()
