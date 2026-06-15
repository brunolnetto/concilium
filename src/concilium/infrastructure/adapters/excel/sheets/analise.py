from __future__ import annotations
from typing import Any
from openpyxl import Workbook
from concilium.domain.models.canonical import CanonicalModel
from ..styles import _fill, _font, _border, _align, _money, _label, _header_row, _title_row, _section_row, _set_col_widths, _auto_fit_columns, _freeze, _C
from ..helpers import _fmt_mes, _week_to_month
from .operacional import _wk_label


def _compute_equip_min_units(eq: object) -> int:
    """Normalize by minimum equipment unit count to recover composite-front count."""
    if not eq.equipamentos:  # type: ignore[union-attr]
        return 1
    return min(
        (e.num_unidades for e in eq.equipamentos if (e.num_unidades or 0) > 0),  # type: ignore[union-attr]
        default=1,
    )


def _aggregate_equip_demand(act: object, eq: object, all_weeks: list, equip_data: dict) -> None:
    """Accumulate equipment demand from one scheduled activity into equip_data."""
    _min_units = _compute_equip_min_units(eq)
    for equip in eq.equipamentos:  # type: ignore[union-attr]
        if equip.equip_id not in equip_data:
            equip_data[equip.equip_id] = {
                "descricao": getattr(equip, "descricao", equip.equip_id) or equip.equip_id,
                "per_week": {wk: 0 for wk in all_weeks},
            }
        for wk, dist in act.distribuicao_semanal.items():  # type: ignore[union-attr]
            n = getattr(dist, "num_equipes", 0)
            if n > 0 and wk in equip_data[equip.equip_id]["per_week"]:
                equip_data[equip.equip_id]["per_week"][wk] += round((n / _min_units) * equip.num_unidades)


def _write_hist_data_row(
    ws: object, row: int, id_key: str, data: dict, unit_label: str, all_weeks: list, fill: object, week_totals: dict
) -> None:
    _label(ws, row, 1, id_key, fill)
    _label(ws, row, 2, data["descricao"], fill)
    _label(ws, row, 3, unit_label, fill, align="center")
    for j, wk in enumerate(all_weeks):
        v = data["per_week"].get(wk, 0)
        cell = ws.cell(row=row, column=4 + j, value=v if v > 0 else None)  # type: ignore[union-attr]
        cell.fill = fill
        cell.border = _border()
        if v > 0:
            cell.alignment = _align("right")
            cell.number_format = "0"
        week_totals[wk] += v


def _write_hist_totals_row(ws: object, row: int, all_weeks: list, fill_tot: object, week_totals: dict) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)  # type: ignore[union-attr]
    cell = ws.cell(row=row, column=1, value="TOTAL SEMANAL")  # type: ignore[union-attr]
    cell.fill = fill_tot
    cell.font = _font(bold=True, color=_C["white"])
    cell.alignment = _align("right")
    cell.border = _border()
    for j, wk in enumerate(all_weeks):
        v = week_totals[wk]
        cell = ws.cell(row=row, column=4 + j, value=v if v > 0 else None)  # type: ignore[union-attr]
        cell.fill = fill_tot
        cell.border = _border()
        if v > 0:
            cell.font = _font(bold=True, color=_C["white"])
            cell.alignment = _align("right")
            cell.number_format = "0"


def render_histograma_eq(wb: Workbook, canonical: CanonicalModel) -> None:
    """ANALISE.HIST_EQUIP — Demanda semanal de equipamentos (un. × semana ISO)."""
    ws = wb.create_sheet("ANALISE.HIST_EQUIP")
    acts = canonical.cronograma.scheduled_activities if canonical.cronograma else []
    all_weeks = sorted({w for act in acts for w in act.distribuicao_semanal})
    if not all_weeks:
        ws.cell(row=1, column=1).value = "Sem dados de cronograma."
        return
    equipe_by_item = {eq.item_id: eq for eq in canonical.equipes_diretas}
    equip_data: dict[str, dict] = {}
    for act in acts:
        eq = equipe_by_item.get(act.activity_id)
        if eq is None:
            continue
        _aggregate_equip_demand(act, eq, all_weeks, equip_data)
    NCOLS = 3 + len(all_weeks)
    _set_col_widths(ws, [14, 32, 6] + [10] * len(all_weeks))  # type: ignore[arg-type]
    _freeze(ws, "D3")
    row = 1
    _title_row(ws, row, "HISTOGRAMA DE EQUIPAMENTOS — DEMANDA SEMANAL (UN.)", NCOLS)
    row += 1
    _header_row(ws, row, ["Equip. ID", "Descrição", "Un."] + list(map(_wk_label, all_weeks)))
    row += 1
    ws.freeze_panes = f"D{row}"
    fill_colors = [_fill(_C["white"]), _fill(_C["row_alt"])]
    week_totals = dict.fromkeys(all_weeks, 0)
    for i, (eid, data) in enumerate(sorted(equip_data.items())):
        fill = fill_colors[i % 2]
        _write_hist_data_row(ws, row, eid, data, "un", all_weeks, fill, week_totals)
        row += 1
    row += 1
    _write_hist_totals_row(ws, row, all_weeks, _fill(_C["nav_dk"]), week_totals)


def _aggregate_labor_demand(act: object, eq: object, all_weeks: list, func_data: dict) -> None:
    """Accumulate labor demand from one scheduled activity into func_data."""
    _min_units = _compute_equip_min_units(eq)
    for func in eq.funcoes:  # type: ignore[union-attr]
        if func.funcao_id not in func_data:
            func_data[func.funcao_id] = {
                "descricao": func.descricao or func.funcao_id,
                "per_week": {wk: 0 for wk in all_weeks},
            }
        for wk, dist in act.distribuicao_semanal.items():  # type: ignore[union-attr]
            n = getattr(dist, "num_equipes", 0)
            if n > 0 and wk in func_data[func.funcao_id]["per_week"]:
                func_data[func.funcao_id]["per_week"][wk] += round((n / _min_units) * func.num_pessoas)


def render_histograma_mo(wb: Workbook, canonical: CanonicalModel) -> None:
    """ANALISE.HIST_MAO_OBRA — Demanda semanal de mão de obra (pessoas × semana ISO)."""
    ws = wb.create_sheet("ANALISE.HIST_MAO_OBRA")
    acts = canonical.cronograma.scheduled_activities if canonical.cronograma else []
    all_weeks = sorted({w for act in acts for w in act.distribuicao_semanal})
    if not all_weeks:
        ws.cell(row=1, column=1).value = "Sem dados de cronograma."
        return
    equipe_by_item = {eq.item_id: eq for eq in canonical.equipes_diretas}
    func_data: dict[str, dict] = {}
    for act in acts:
        eq = equipe_by_item.get(act.activity_id)
        if eq is None:
            continue
        _aggregate_labor_demand(act, eq, all_weeks, func_data)
    NCOLS = 3 + len(all_weeks)
    _set_col_widths(ws, [20, 32, 6] + [10] * len(all_weeks))  # type: ignore[arg-type]
    _freeze(ws, "D3")
    row = 1
    _title_row(ws, row, "HISTOGRAMA DE MÃO DE OBRA — DEMANDA SEMANAL (PESSOAS)", NCOLS)
    row += 1
    _header_row(ws, row, ["Função ID", "Descrição", "Un."] + list(map(_wk_label, all_weeks)))
    row += 1
    ws.freeze_panes = f"D{row}"
    fill_colors = [_fill(_C["white"]), _fill(_C["row_alt"])]
    week_totals = dict.fromkeys(all_weeks, 0)
    for i, (fid, data) in enumerate(sorted(func_data.items())):
        fill = fill_colors[i % 2]
        _write_hist_data_row(ws, row, fid, data, "pax", all_weeks, fill, week_totals)
        row += 1
    row += 1
    _write_hist_totals_row(ws, row, all_weeks, _fill(_C["nav_dk"]), week_totals)


def render_curva_abc(wb, canonical) -> None:
    """ANALISE.CURVA_ABC — ABC por insumo (funcao_id / equip_id) do projeto (ADR-058)."""
    from concilium.domain.services.analytics.curva_abc_builder import CurvaABCBuilder

    abc = getattr(canonical, "curva_abc_insumo", None)
    if abc is None:
        try:
            abc = CurvaABCBuilder().build(canonical)
        except Exception:
            return

    ws = wb.create_sheet("ANALISE.CURVA_ABC")
    NCOLS = 8
    _set_col_widths(ws, [22, 14, 12, 14, 14, 14, 8, 30])
    row = 1
    _title_row(ws, row, "CURVA ABC POR INSUMO — ANALISE DE PARETO (80/20) — ADR-058", NCOLS)
    row += 1
    ws.cell(
        row=row,
        column=1,
        value=f"Total custo direto: R$ {abc.total_custo_direto:,.2f}  |  "
        f"Classe A: {abc.n_itens_a} itens ({abc.pct_a * 100:.1f}%)  |  "
        f"B: {abc.n_itens_b}  |  C: {abc.n_itens_c}",
    ).font = _font(italic=True, size=9)
    row += 2

    _header_row(
        ws,
        row,
        ["Insumo ID", "Descricao", "Categoria", "Custo Total (R$)", "% Custo", "% Acumulado", "Classe", "Atividades"],
    )
    row += 1
    ws.freeze_panes = f"A{row}"

    fills = [_fill(_C["white"]), _fill(_C["row_alt"])]
    classe_colors = {"A": _C["vermelho"], "B": _C["amarelo"], "C": _C["verde"]}
    for i, item in enumerate(abc.itens):
        fill = fills[i % 2]
        for j, v in enumerate(
            [
                item.insumo_id,
                item.descricao,
                item.categoria,
                item.custo_total,
                item.pct_custo * 100,
                item.pct_custo_acumulado * 100,
                item.classe,
                ", ".join(item.atividades[:5]),
            ],
            1,
        ):
            cell = ws.cell(row=row, column=j, value=v)
            cell.fill = fill
            cell.border = _border()
            if j == 4:
                cell.number_format = "R$ #,##0.00"
            if j in (5, 6):
                cell.number_format = '0.00"%"'
                cell.alignment = _align("right")
            if j == 7:
                cell.font = _font(bold=True, color=classe_colors.get(str(v), _C["preto"]))
                cell.alignment = _align("center")
        row += 1


def render_s_curve(wb, canonical) -> None:
    """ANALISE.S_CURVE — Curva S de custo acumulado semana a semana (ADR-058)."""
    from concilium.domain.services.analytics.s_curve_builder import SCurveBuilder

    sc = getattr(canonical, "s_curve", None)
    if sc is None:
        try:
            sc = SCurveBuilder().build(canonical)
        except Exception:
            return
    if not sc.semanas_iso:
        return

    ws = wb.create_sheet("ANALISE.S_CURVE")
    NCOLS = 8
    _set_col_widths(ws, [14, 14, 14, 14, 14, 14, 14, 18])
    row = 1
    _title_row(ws, row, "CURVA S — DESEMBOLSO ACUMULADO PLANEJADO (ADR-058)", NCOLS)
    row += 1
    ws.cell(
        row=row,
        column=1,
        value=f"Total projeto: R$ {sc.total_projeto:,.2f}  |  "
        f"Direto: R$ {sc.total_custo_direto:,.2f}  |  "
        f"Indireto: R$ {sc.total_custo_indireto:,.2f}",
    ).font = _font(italic=True, size=9)
    row += 2
    _header_row(
        ws,
        row,
        [
            "Semana ISO",
            "Equip. Unicos",
            "Equipes Total",
            "Custo Semana (R$)",
            "Custo Acum. (R$)",
            "% Acumulado",
            "C. Direto",
            "C. Indireto",
        ],
    )
    row += 1
    ws.freeze_panes = f"A{row}"

    fills = [_fill(_C["white"]), _fill(_C["row_alt"])]
    for i, (semana, c_sem, c_acum, pct, agr) in enumerate(
        zip(sc.semanas_iso, sc.custo_semanal, sc.custo_acumulado, sc.pct_acumulado, sc.agregacao_semanal, strict=False)
    ):
        fill = fills[i % 2]
        for j, v in enumerate(
            [
                semana,
                agr.n_equipamentos_unicos,
                agr.n_equipes_totais,
                c_sem,
                c_acum,
                pct * 100,
                agr.custo_direto_semana,
                agr.custo_indireto_semana,
            ],
            1,
        ):
            cell = ws.cell(row=row, column=j, value=v)
            cell.fill = fill
            cell.border = _border()
            if j in (4, 5, 7, 8):
                cell.number_format = "R$ #,##0.00"
            if j == 6:
                cell.number_format = '0.00"%"'
                cell.alignment = _align("right")
        row += 1
