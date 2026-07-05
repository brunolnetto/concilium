from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from concilium.domain.models.canonical import CanonicalModel

from ..cash_flow import _cash_flow
from ..styles import (
    _C,
    MESES_PT,
    _align,
    _auto_fit_columns,
    _border,
    _fill,
    _font,
    _header_row,
    _label,
    _section_row,
    _set_col_widths,
    _title_row,
)


def _rg_week_label(week_key: str) -> str:
    year_str, week_str = week_key.split("-W")
    monday = date.fromisocalendar(int(year_str), int(week_str), 1)
    return f"{monday.day:02d}/{MESES_PT[monday.month]}"


def _rg_write_finance_rows(ws, finance_rows: list) -> None:
    finance_col = 1
    finance_row = 20
    ws.cell(row=finance_row, column=finance_col, value="Fluxo mensal")
    headers = ["Mes", "Receita", "Despesa", "Saldo", "Exposicao"]
    for offset, header in enumerate(headers):
        ws.cell(row=finance_row + 1, column=finance_col + offset, value=header)
    for idx, row in enumerate(finance_rows, start=finance_row + 2):
        ws.cell(row=idx, column=finance_col, value=row["mes"])
        for offset, key in enumerate(("receita", "despesa", "saldo", "exposicao"), start=1):
            cell = ws.cell(row=idx, column=finance_col + offset, value=row[key])
            cell.number_format = "#,##0.00"


def _rg_hist_scale(eq, kind: str) -> float:
    """Compute scale factor for an equipe based on kind."""
    if kind == "equipamentos":
        return sum(float(e.num_unidades or 0.0) for e in eq.equipamentos)
    return sum(float(f.num_pessoas or 0.0) for f in eq.funcoes)


def _rg_hist_accumulate_act(act, lookup, kind: str, totals: dict) -> None:
    eq = lookup.get(act.activity_id)
    if eq is None:
        return
    if eq.equipamentos:
        candidates = [e.num_unidades for e in eq.equipamentos if (e.num_unidades or 0) > 0]
        min_units = min(candidates) if candidates else 1
    else:
        min_units = 1
    scale = _rg_hist_scale(eq, kind)
    for week, dist in act.distribuicao_semanal.items():
        n_eq = float(getattr(dist, "num_equipes", 0) or 0.0)
        if n_eq > 0:
            totals[week] += (n_eq / min_units) * scale


def _rg_histogram_totals(canonical, kind: str) -> tuple:
    acts = canonical.cronograma.scheduled_activities if canonical.cronograma else []
    all_weeks_set = set()
    for act in acts:
        for week in act.distribuicao_semanal:
            all_weeks_set.add(week)
    all_weeks = sorted(all_weeks_set)
    if not all_weeks:
        return [], []
    lookup = {}
    for eq in canonical.equipes_diretas:
        lookup[eq.item_id] = eq
    totals = {week: 0.0 for week in all_weeks}
    for act in acts:
        _rg_hist_accumulate_act(act, lookup, kind, totals)
    labels = [_rg_week_label(week) for week in all_weeks]
    values = [round(totals[week], 2) for week in all_weeks]
    return labels, values


def _rg_cash_flow_rows(canonical) -> list:
    rows = _cash_flow(canonical)
    return [
        {
            "mes": row["mes_label"],
            "receita": float(row["receita"]),
            "despesa": float(row["despesa"]),
            "saldo": float(row["saldo"]),
            "exposicao": float(row["exposicao_acumulada"]),
        }
        for row in rows
    ]


def _rg_write_two_col_block(ws, start_col: int, title: str, labels: list, values: list) -> int:
    ws.cell(row=1, column=start_col, value=title)
    ws.cell(row=2, column=start_col, value="Categoria")
    ws.cell(row=2, column=start_col + 1, value="Valor")
    for idx, (label, value) in enumerate(zip(labels, values, strict=False), start=3):
        ws.cell(row=idx, column=start_col, value=label)
        cell = ws.cell(row=idx, column=start_col + 1, value=value)
        cell.number_format = "#,##0.00"
    return len(labels) + 2


def _rg_build_gantt_map(canonical) -> list:
    """Build Gantt map: list of (activity_id, start_offset_days, duration_days)."""
    acts = canonical.cronograma.scheduled_activities if canonical.cronograma else []
    start_dates = []
    for act in acts:
        weeks = sorted(act.distribuicao_semanal.keys())
        if not weeks:
            continue
        year_str, week_str = weeks[0].split("-W")
        start_dates.append(date.fromisocalendar(int(year_str), int(week_str), 1))
    gantt_origin = min(start_dates) if start_dates else None
    gantt_map = []
    for act in acts:
        weeks = sorted(act.distribuicao_semanal.keys())
        if not weeks:
            continue
        year_str, week_str = weeks[0].split("-W")
        start = date.fromisocalendar(int(year_str), int(week_str), 1)
        start_offset = (start - gantt_origin).days if gantt_origin else 0
        gantt_map.append((act.activity_id, start_offset, len(weeks) * 7))
    return gantt_map


def render_graficos(wb: Workbook, canonical: CanonicalModel) -> None:
    """GRAFICOS — painéis resumidos para histograma, Gantt e fluxo mensal."""
    ws = wb.create_sheet("GRAFICOS")
    _set_col_widths(ws, [18, 14, 18, 14, 18, 14, 18, 14, 18, 14, 18, 14])

    eq_labels, eq_values = _rg_histogram_totals(canonical, "equipamentos")
    mo_labels, mo_values = _rg_histogram_totals(canonical, "mao_obra")
    finance_rows = _rg_cash_flow_rows(canonical)

    _rg_write_two_col_block(ws, 1, "Histograma de Equipamentos", eq_labels, eq_values)
    _rg_write_two_col_block(ws, 4, "Histograma de Mao de Obra", mo_labels, mo_values)

    gantt_col = 7
    ws.cell(row=1, column=gantt_col, value="Cronograma (Gantt)")
    ws.cell(row=2, column=gantt_col, value="Tarefa")
    ws.cell(row=2, column=gantt_col + 1, value="Inicio")
    ws.cell(row=2, column=gantt_col + 2, value="Duracao (dias)")
    gantt_map = _rg_build_gantt_map(canonical)
    for idx, (task, start_offset, duration_days) in enumerate(gantt_map, start=3):
        ws.cell(row=idx, column=gantt_col, value=task)
        ws.cell(row=idx, column=gantt_col + 1, value=start_offset)
        ws.cell(row=idx, column=gantt_col + 2, value=duration_days)

    _rg_write_finance_rows(ws, finance_rows)

    if eq_labels:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Histograma de Equipamentos"
        # openpyxl: TitleDescriptor.__set__ aceita str e converte pra Title em runtime
        chart.y_axis.title = "Unidades"  # ty: ignore[invalid-assignment]
        data = Reference(ws, min_col=2, min_row=2, max_row=1 + len(eq_labels) + 1)
        cats = Reference(ws, min_col=1, min_row=3, max_row=2 + len(eq_labels))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "J2")

    if mo_labels:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Histograma de Mao de Obra"
        chart.y_axis.title = "Pessoas"  # ty: ignore[invalid-assignment]
        data = Reference(ws, min_col=5, min_row=2, max_row=1 + len(mo_labels) + 1)
        cats = Reference(ws, min_col=4, min_row=3, max_row=2 + len(mo_labels))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "J18")

    if gantt_map:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.title = "Cronograma Gantt"
        chart.x_axis.title = "Dias"  # ty: ignore[invalid-assignment]
        chart.y_axis.title = "Atividades"  # ty: ignore[invalid-assignment]
        data = Reference(ws, min_col=gantt_col + 1, min_row=2, max_col=gantt_col + 2, max_row=2 + len(gantt_map))
        cats = Reference(ws, min_col=gantt_col, min_row=3, max_row=2 + len(gantt_map))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        if chart.series:
            chart.series[0].graphicalProperties.noFill = True
            chart.series[0].graphicalProperties.line.noFill = True
        ws.add_chart(chart, "J34")

    if finance_rows:
        finance_row = 20
        finance_col = 1
        finance_end = finance_row + 1 + len(finance_rows)
        finance_metrics = [
            (2, "Receita Mensal"),
            (3, "Despesa Mensal"),
            (4, "Saldo Mensal"),
            (5, "Exposicao Acumulada"),
        ]
        chart_positions = ["A36", "J36", "A52", "J52"]
        for (col_idx, title), anchor in zip(finance_metrics, chart_positions, strict=False):
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = title
            chart.y_axis.title = "R$"  # ty: ignore[invalid-assignment]
            data = Reference(ws, min_col=col_idx, min_row=finance_row + 1, max_row=finance_end)
            cats = Reference(ws, min_col=finance_col, min_row=finance_row + 2, max_row=finance_end)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, anchor)


def _rcc_render_summary_section(ws, row: int, grupos, cubo, fills) -> int:
    for i, (nome, val) in enumerate(grupos):
        fill = fills[i % 2]
        pct = val / cubo.total_preco_venda * 100 if cubo.total_preco_venda > 0 else 0
        _label(ws, row, 1, nome, fill, bold=True, color=_C["nav_dk"])
        cell_val = ws.cell(row=row, column=2, value=val)
        cell_val.fill = fill
        cell_val.border = _border()
        cell_val.number_format = "R$ #,##0.00"
        cell_pct = ws.cell(row=row, column=3, value=round(pct, 2))
        cell_pct.fill = fill
        cell_pct.border = _border()
        cell_pct.number_format = '0.00"%"'
        row += 1
    return row


def _rcc_render_tributaria_section(ws, row: int, trib_map, cubo, fills) -> int:
    for i, (trib, val) in enumerate(sorted(trib_map.items(), key=lambda kv: kv[1], reverse=True)):
        fill = fills[i % 2]
        pct = val / cubo.total_preco_venda * 100 if cubo.total_preco_venda > 0 else 0
        _label(ws, row, 1, trib.replace("tributo_", "").upper(), fill, color=_C["vermelho"])
        ws.cell(row=row, column=2, value=val).fill = fill
        ws.cell(row=row, column=2).border = _border()
        ws.cell(row=row, column=2).number_format = "R$ #,##0.00"
        cell_p = ws.cell(row=row, column=3, value=round(pct, 2))
        cell_p.fill = fill
        cell_p.border = _border()
        cell_p.number_format = '0.00"%"'
        row += 1
    return row


def _rcc_render_margem_section(ws, row: int, margem_ativ, total_margem: float, fills) -> int:
    for i, (ativ, val) in enumerate(margem_ativ.items()):
        fill = fills[i % 2]
        pct = val / total_margem * 100
        _label(ws, row, 1, ativ, fill, color=_C["verde"])
        ws.cell(row=row, column=2, value=val).fill = fill
        ws.cell(row=row, column=2).border = _border()
        ws.cell(row=row, column=2).number_format = "R$ #,##0.00"
        cell_p = ws.cell(row=row, column=3, value=round(pct, 2))
        cell_p.fill = fill
        cell_p.border = _border()
        cell_p.number_format = '0.00"%"'
        row += 1
    return row


def _rcc_render_temporal_row(ws, row: int, semana, comps, fill, total_pv: float, acum: float):
    c_dir = comps.get("custo_direto", 0.0)
    c_oh = sum(comps.get(k, 0.0) for k in ("overhead_ac", "overhead_df", "overhead_r", "overhead_cr"))
    c_marg = comps.get("margem_l", 0.0)
    c_trib = sum(
        comps.get(k, 0.0) for k in ("tributo_cofins", "tributo_iss", "tributo_pis", "tributo_cssl", "tributo_ir")
    )
    total_sem = c_dir + c_oh + c_marg + c_trib
    acum += total_sem
    pct_sem = total_sem / total_pv * 100
    pct_acum = acum / total_pv * 100
    for j, v in enumerate([semana, c_dir, c_oh, c_marg, c_trib, total_sem, pct_acum, pct_sem], 1):
        cell = ws.cell(row=row, column=j, value=round(v, 2) if isinstance(v, float) else v)
        cell.fill = fill
        cell.border = _border()
        if j in (2, 3, 4, 5, 6):
            cell.number_format = "R$ #,##0.00"
        if j in (7, 8):
            cell.number_format = '0.00"%"'
            cell.alignment = _align("right")
        if j == 4:
            cell.font = _font(color=_C["verde"])
        if j == 5:
            cell.font = _font(color=_C["vermelho"])
    return acum, row + 1


def render_cubo_comp(wb, canonical) -> None:
    """CUBO.COMP — Composicao do preco de venda ao longo do tempo (ADR-058 ext.).

    Exibe a decomposicao do preco de venda semanal em 4 grupos:
        Custo Direto | Overhead (AC+DF+R+CR) | Margem L | Tributos (ISS+COFINS+PIS+CSSL+IR)

    Permite identificar semanas de alta exposicao tributaria e pico de margem.
    """
    from concilium.domain.services.analytics.cubo_analitico_builder import CuboAnaliticoBuilder

    cubo = getattr(canonical, "cubo_analitico", None)
    if cubo is None:
        try:
            cubo = CuboAnaliticoBuilder().build(canonical)
        except Exception:
            return
    if not cubo.celulas:
        return

    ws = wb.create_sheet("CUBO.COMP")
    NCOLS = 12
    _set_col_widths(ws, [14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 12])
    row = 1

    _title_row(ws, row, "COMPOSICAO DO PRECO DE VENDA — CUBO 4D (CATEGORIA x ATIVIDADE x TEMPO x BDI)", NCOLS)
    row += 1
    ws.cell(
        row=row,
        column=1,
        value="Relevancia: evidencia composicao de preco por semana para auditoria de tributos/margem.  |  "
        f"Regime BDI: {cubo.regime_bdi}  |  Total PV: R$ {cubo.total_preco_venda:,.2f}  |  "
        f"Custo Direto: {cubo.total_custo_direto / cubo.total_preco_venda * 100:.1f}%  |  "
        f"Margem: {cubo.pct_margem * 100:.1f}%  |  Tributos: {cubo.pct_tributos * 100:.1f}%",
    ).font = _font(italic=True, size=9)
    row += 2

    # Secao 1: Sumario por componente
    _section_row(ws, row, "1. SUMARIO POR COMPONENTE BDI (total do projeto)", NCOLS)
    row += 1
    _header_row(ws, row, ["Componente", "R$ Total", "% PV"])
    row += 1
    grupos = [
        ("Custo Direto", cubo.total_custo_direto),
        ("Overhead Total", cubo.total_overhead),
        ("Margem L", cubo.total_margem),
        ("Tributos Total", cubo.total_tributos),
    ]
    fills = [_fill(_C["white"]), _fill(_C["row_alt"])]
    row = _rcc_render_summary_section(ws, row, grupos, cubo, fills)
    row += 1

    # Secao 2: Exposicao tributaria por tipo
    _section_row(ws, row, "2. EXPOSICAO TRIBUTARIA POR TIPO (projeto inteiro)", NCOLS)
    row += 1
    _header_row(ws, row, ["Tributo", "R$ Total", "% PV"])
    row += 1
    trib_map = cubo.exposicao_tributaria_por_tipo()
    row = _rcc_render_tributaria_section(ws, row, trib_map, cubo, fills)
    row += 1

    # Secao 3: Margem por atividade
    _section_row(ws, row, "3. MARGEM (L) POR ATIVIDADE", NCOLS)
    row += 1
    _header_row(ws, row, ["Atividade", "Margem (R$)", "% Margem Total"])
    row += 1
    margem_ativ = cubo.margem_por_atividade()
    total_margem = cubo.total_margem or 1.0
    row = _rcc_render_margem_section(ws, row, margem_ativ, total_margem, fills)
    row += 1

    # Secao 4: Composicao temporal (custo+overhead+margem+tributos por semana)
    _section_row(ws, row, "4. COMPOSICAO DO PRECO POR SEMANA (custo | overhead | margem | tributos)", NCOLS)
    row += 1
    _header_row(
        ws,
        row,
        [
            "Semana ISO",
            "Custo Direto",
            "Overhead",
            "Margem L",
            "Tributos",
            "Total Semana",
            "% Acumulado",
            "% Semana/Total",
        ],
    )
    row += 1
    ws.freeze_panes = f"A{row}"

    comp_temp = cubo.composicao_preco_temporal()
    total_pv = cubo.total_preco_venda or 1.0
    acum = 0.0
    for i, (semana, comps) in enumerate(comp_temp.items()):
        fill = fills[i % 2]
        acum, row = _rcc_render_temporal_row(ws, row, semana, comps, fill, total_pv, acum)
    _auto_fit_columns(ws)


def render_campanha_frota(wb, canonical) -> None:
    """CAMPANHA.FROTA — Plano de campanhas de permanência de equipamentos (ADR-059)."""
    from concilium.domain.services.campaign.campanha_equipamento_builder import CampanhaEquipamentoBuilder

    analise = getattr(canonical, "analise_campanhas", None)
    if analise is None:
        try:
            analise = CampanhaEquipamentoBuilder().build(canonical)
        except Exception:
            analise = None

    ws = wb.create_sheet("CAMPANHA.FROTA")
    NCOLS = 12
    _set_col_widths(ws, [22, 14, 10, 10, 10, 10, 14, 14, 14, 14, 14, 30])
    row = 1
    _title_row(ws, row, "CAMPANHAS DE PERMANENCIA DE EQUIPAMENTOS (ADR-059)", NCOLS)
    row += 1
    if analise is None:
        ws.cell(row=row, column=1, value="Campanhas de permanencia nao disponiveis para este projeto.").font = _font(
            italic=True, size=9
        )
        row += 2
    else:
        ws.cell(
            row=row,
            column=1,
            value=f"Campanhas: {analise.n_campanhas}  |  "
            f"Economia vs. dismissal: R$ {analise.economia_total:,.2f}  |  "
            f"Standby total: R$ {analise.custo_standby_total:,.2f}  |  "
            f"Economia liquida: R$ {analise.economia_liquida:,.2f}",
        ).font = _font(italic=True, size=9)
        row += 2

        if analise.n_campanhas == 0:
            n_equip_tipos = len(
                {
                    str(e.equip_id)
                    for eq in (canonical.equipes_diretas or [])
                    for e in (eq.equipamentos or [])
                    if getattr(e, "equip_id", "")
                }
            )
            n_ativ_com_equip = sum(1 for eq in (canonical.equipes_diretas or []) if (eq.equipamentos or []))
            ws.cell(
                row=row,
                column=1,
                value=(
                    "Nenhuma campanha detectada. Interpretação: cronograma sem gaps elegíveis para compartilhamento "
                    f"(tipos equip.: {n_equip_tipos}; atividades com equipamento: {n_ativ_com_equip})."
                ),
            ).font = _font(italic=True, size=9)
            row += 2

        _header_row(
            ws,
            row,
            [
                "Instancia",
                "Tipo Equip.",
                "Sem. Entrada",
                "Sem. Saida",
                "Sem. Ativas",
                "Sem. Standby",
                "Mob (R$)",
                "Demob (R$)",
                "Standby Total",
                "Custo Camp.",
                "Economia",
                "Atividades",
            ],
        )
        row += 1
        ws.freeze_panes = f"A{row}"

        fills = [_fill(_C["white"]), _fill(_C["row_alt"])]
        for i, camp in enumerate(analise.campanhas):
            fill = fills[i % 2]
            economia_color = _C["verde"] if camp.economia_vs_dismissal > 0 else _C["preto"]
            for j, v in enumerate(
                [
                    camp.instancia_id,
                    camp.tipo_id,
                    camp.semana_entrada,
                    camp.semana_saida,
                    camp.semanas_ativas,
                    camp.semanas_standby,
                    camp.custo_mob,
                    camp.custo_demob,
                    camp.custo_standby_total,
                    camp.custo_campanha,
                    camp.economia_vs_dismissal,
                    " | ".join(camp.atividades),
                ],
                1,
            ):
                cell = ws.cell(row=row, column=j, value=v)
                cell.fill = fill
                cell.border = _border()
                if j in (7, 8, 9, 10):
                    cell.number_format = "R$ #,##0.00"
                if j == 11:
                    cell.number_format = "R$ #,##0.00"
                    cell.font = _font(bold=True, color=economia_color)
            row += 1

        if analise.alertas_custo_duplicado:
            row += 1
            _section_row(ws, row, "ALERTAS VAL-CAMP-01 — Dismissals potencialmente mais caros que standby", NCOLS)
            row += 1
            for alert in analise.alertas_custo_duplicado:
                ws.cell(row=row, column=1, value=alert).font = _font(color=_C["amarelo"], size=9)
                row += 1

    _auto_fit_columns(ws)
