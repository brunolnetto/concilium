"""Excel adapter — abas MEMORIAL — render_memorial_equipamentos."""

from __future__ import annotations

from typing import Any

from openpyxl import Workbook

from concilium.domain.models.canonical import CanonicalModel

from ...styles import (
    _C,
    _align,
    _auto_fit_columns,
    _border,
    _fill,
    _font,
    _header_row,
    _set_col_widths,
    _title_row,
)


def _xl_memorial_equip_accumulate(agg: dict[str, dict[str, Any]], eq: Any, e: Any) -> None:
    """Acumula um equipamento de uma equipe direta no dict agregador (mutação in-place)."""
    key = str(e.equip_id or "")
    if not key:
        return
    rec = agg.setdefault(
        key,
        {
            "descricao": e.descricao or key,
            "qtd": 0.0,
            "custo": 0.0,
            "comb": 0.0,
        },
    )
    n = float(getattr(e, "num_unidades", 0.0) or 0.0) * float(getattr(eq, "num_equipes_base", 0.0) or 0.0)
    rec["qtd"] += n
    rec["custo"] += float(getattr(e, "custo_mensal", 0.0) or 0.0) * n
    rec["comb"] += float(getattr(e, "custo_combustivel_mensal", 0.0) or 0.0) * n


def _xl_memorial_equip_agg(canonical: CanonicalModel) -> dict[str, dict[str, Any]]:
    """Agrega equipamentos por equip_id a partir das equipes diretas."""
    agg: dict[str, dict[str, Any]] = {}
    for eq in canonical.equipes_diretas or []:
        for e in eq.equipamentos or []:
            _xl_memorial_equip_accumulate(agg, eq, e)
    return agg


def _xl_memorial_equip_row(ws: Any, row: int, equip_id: str, rec: dict[str, Any], fill: Any) -> None:
    total = float(rec["custo"]) + float(rec["comb"])
    vals = [equip_id, rec["descricao"], rec["qtd"], rec["custo"], rec["comb"], total]
    for col, val in enumerate(vals, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.fill = fill
        cell.border = _border()
        if col >= 4:
            cell.number_format = "R$ #,##0.00"
        if col == 3:
            cell.number_format = "0.00"
            cell.alignment = _align("right")


def render_memorial_equipamentos(wb: Workbook, canonical: CanonicalModel) -> None:
    """MEMORIAL.EQUIPAMENTOS — base de equipamentos efetivamente usada no projeto."""
    ws = wb.create_sheet("MEMORIAL.EQUIPAMENTOS")
    _set_col_widths(ws, [22, 32, 12, 16, 16, 16])
    row = 1
    _title_row(ws, row, "MEMORIAL.EQUIPAMENTOS — INSUMOS UTILIZADOS", 6)
    row += 1
    ws.cell(row=row, column=1, value="Fonte: composições de EQUIPE.DIRETO.").font = _font(italic=True, size=9)
    row += 2
    _header_row(ws, row, ["Equip ID", "Descricao", "Qtd Base", "Custo Mensal", "Comb. Mensal", "Total Mensal"])
    row += 1

    agg = _xl_memorial_equip_agg(canonical)

    fills = [_fill(_C["white"]), _fill(_C["row_alt"])]
    if not agg:
        ws.cell(row=row, column=1, value="Sem dados de equipamentos para memorial.").font = _font(italic=True, size=9)
    for i, (equip_id, rec) in enumerate(sorted(agg.items())):
        fill = fills[i % 2]
        _xl_memorial_equip_row(ws, row, equip_id, rec, fill)
        row += 1
    _auto_fit_columns(ws)


