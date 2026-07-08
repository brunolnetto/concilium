"""Excel adapter — abas MEMORIAL — render_memorial_materiais."""

from __future__ import annotations

from typing import Any

from openpyxl import Workbook

from concilium.domain.models.canonical import CanonicalModel

from ...styles import (
    _C,
    _align,
    _auto_fit_columns,
    _border,
    _fill,
    _font,
    _header_row,
    _set_col_widths,
    _title_row,
)


def _xl_memorial_material_row(ws: Any, row: int, p: Any, fill: Any) -> None:
    qtd = float(getattr(p, "quantitativo", 0.0) or 0.0)
    pu_mat = float(getattr(getattr(p, "preco_unitario", None), "material", 0.0) or 0.0)
    sem_bdi = float(getattr(getattr(getattr(p, "subtotal", None), "sem_bdi", None), "material", 0.0) or 0.0)
    com_bdi = float(getattr(getattr(getattr(p, "subtotal", None), "com_bdi", None), "material", 0.0) or 0.0)
    vals = [getattr(p, "item_id", "-"), getattr(p, "cpu_id", "-"), qtd, pu_mat, sem_bdi, com_bdi]
    for col, val in enumerate(vals, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.fill = fill
        cell.border = _border()
        if col >= 4:
            cell.number_format = "R$ #,##0.00"
        if col == 3:
            cell.number_format = "0.00"
            cell.alignment = _align("right")


def render_memorial_materiais(wb: Workbook, canonical: CanonicalModel) -> None:
    """MEMORIAL.MATERIAIS — materiais e valores vindos do planejamento/catálogo."""
    ws = wb.create_sheet("MEMORIAL.MATERIAIS")
    _set_col_widths(ws, [14, 36, 12, 16, 16, 16])
    row = 1
    _title_row(ws, row, "MEMORIAL.MATERIAIS — ITENS COM MATERIAL", 6)
    row += 1
    ws.cell(
        row=row,
        column=1,
        value="Fonte: PLANEJAMENTO (preco_unitario.material e subtotal sem/com BDI).",
    ).font = _font(italic=True, size=9)
    row += 2
    _header_row(ws, row, ["Item ID", "CPU ID", "Qtd", "PU Mat. s/BDI", "Subtotal Mat. s/BDI", "Subtotal Mat. c/BDI"])
    row += 1

    fills = [_fill(_C["white"]), _fill(_C["row_alt"])]
    has_rows = False
    for i, p in enumerate(canonical.planejamento or []):
        pu_mat = float(getattr(getattr(p, "preco_unitario", None), "material", 0.0) or 0.0)
        if pu_mat <= 0.0:
            continue
        fill = fills[i % 2]
        _xl_memorial_material_row(ws, row, p, fill)
        row += 1
        has_rows = True

    if not has_rows:
        ws.cell(row=row, column=1, value="Sem materiais com custo explícito neste planejamento.").font = _font(
            italic=True, size=9
        )

    _auto_fit_columns(ws)


