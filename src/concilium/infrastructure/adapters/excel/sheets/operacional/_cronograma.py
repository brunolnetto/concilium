"""Excel adapter — aba OPERACIONAL — render_cronograma (métricas semanais por atividade)."""

from __future__ import annotations

from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from concilium.domain.models.canonical import CanonicalModel

from ...styles import (
    _C,
    MESES_PT,
    _align,
    _border,
    _fill,
    _font,
    _header_row,
    _label,
    _set_col_widths,
    _title_row,
)

_CRONOGRAMA_METRICS = [
    ("Dias Trabalhados", "dias", "0"),
    ("Nº De Equipes", "n_eq", "0"),
    ("Qtd. Planejada", "qtd_exec", "#,##0.00"),
    ("Qtd Plan. Acumulada", "qtd_acum", "#,##0.00"),
    ("% Semanal", "pct_sem", "0.00%"),
    ("% Acumulado", "pct_acum", "0.00%"),
]


def _wk_label(wk: str) -> str:
    yr, wn = wk.split("-W")
    mon = date.fromisocalendar(int(yr), int(wn), 1)
    return f"{mon.day:02d}/{MESES_PT[mon.month]}"


def _compute_pct(qtd: float, total_item: float) -> Any | None:
    if total_item > 0 and qtd > 0:
        return round(qtd / total_item, 4)
    return None


def _cronograma_qtd_val(key: str, wk_entry: dict, accum_prop: dict, wk: str) -> Any:
    v = wk_entry.get("qtd_exec", 0.0) if key == "qtd_exec" else accum_prop.get(wk, 0.0)
    return round(v, 2) if v else None


def _compute_cell_metric(
    key: str,
    wk_entry: dict,
    accum_prop: dict,
    wk: str,
    total_item: float,
    dias_semana: int,
) -> Any:
    n_eq = wk_entry.get("n_eq", 0)
    if key == "dias":
        return dias_semana if n_eq > 0 else None
    if key == "n_eq":
        return n_eq if n_eq > 0 else None
    if key in ("qtd_exec", "qtd_acum"):
        return _cronograma_qtd_val(key, wk_entry, accum_prop, wk)
    if key == "pct_sem":
        return _compute_pct(wk_entry.get("qtd_exec", 0.0), total_item)
    if key == "pct_acum":
        return _compute_pct(accum_prop.get(wk, 0.0), total_item)
    return None


def _build_act_lookup_and_totals(acts: list) -> tuple[dict, dict, list[str]]:
    act_lookup: dict[str, dict[str, Any]] = {}
    item_totals: dict[str, float] = {}
    all_weeks_set: set[str] = set()
    for act in acts:
        all_weeks_set.update(act.distribuicao_semanal.keys())
        act_lookup[act.activity_id] = {
            wk: {
                "n_eq": getattr(d, "num_equipes", 0),
                "qtd_exec": getattr(d, "qtd_executada", 0.0),
                "qtd_acum": getattr(d, "qtd_acumulada", 0.0),
            }
            for wk, d in act.distribuicao_semanal.items()
        }
        if act.distribuicao_semanal:
            last_week = sorted(act.distribuicao_semanal.keys())[-1]
            item_totals[act.activity_id] = getattr(act.distribuicao_semanal[last_week], "qtd_acumulada", 0.0)
    return act_lookup, item_totals, sorted(all_weeks_set)


def _compute_accum_prop(wk_data: dict, all_weeks: list) -> tuple[dict, float]:
    accum_prop: dict[str, float] = {}
    last_acum = 0.0
    for wk in all_weeks:
        wk_entry = wk_data.get(wk, {})
        acum_wk = float(wk_entry.get("qtd_acum", 0.0) or 0.0)
        if acum_wk > 0:
            last_acum = acum_wk
        accum_prop[wk] = last_acum
    return accum_prop, last_acum


def _get_pratic_map(canonical: CanonicalModel) -> dict:
    if canonical.praticabilidade:
        return canonical.praticabilidade.percentual_produtivo_semanal or {}
    return {}


def _render_praticabilidade_row(ws: Any, row: int, all_weeks: list, pratic_map: dict) -> None:
    pratic_fill = _fill(_C["amar_lt"])
    for col_p, wk_p in enumerate(["", "", "", "", "Praticabilidade"] + list(all_weeks), 1):
        cell_p = ws.cell(row=row, column=col_p)
        cell_p.fill = pratic_fill
        cell_p.border = _border()
        if isinstance(wk_p, str) and wk_p.startswith("20") and "-W" in wk_p:
            v_p = pratic_map.get(wk_p, 1.0)
            cell_p.value = v_p
            cell_p.number_format = "0%"
            cell_p.alignment = _align("center")
            cell_p.font = _font(bold=(v_p < 1.0), color=_C["orange"] if v_p < 1.0 else _C["preto"])
        elif wk_p == "Praticabilidade":
            cell_p.value = wk_p
            cell_p.font = _font(bold=True, color=_C["amarelo"])
            cell_p.alignment = _align("left")


def _render_item_metrics_rows(
    ws: Any,
    row: int,
    item: Any,
    wk_data: dict,
    accum_prop: dict,
    total_item: float,
    all_weeks: list,
    ncols_fixed: int,
    dias_semana: int,
    fill_item_hdr: Any,
    fill_sub: Any,
) -> int:
    for i_met, (label, key, fmt) in enumerate(_CRONOGRAMA_METRICS):
        is_hdr = i_met == 0
        f = fill_item_hdr if is_hdr else fill_sub
        _label(ws, row, 1, item.item_id if is_hdr else None, f, bold=is_hdr)
        _label(ws, row, 2, item.descricao if is_hdr else None, f, bold=is_hdr)
        _label(ws, row, 3, item.unidade if is_hdr else None, f, align="center")
        _label(ws, row, 4, item.quantitativo if is_hdr else None, f, align="right")
        _label(ws, row, 5, label, f, bold=False, color=_C["nav_dk"])
        for j_wk, wk in enumerate(all_weeks):
            wk_entry = wk_data.get(wk, {})
            val = _compute_cell_metric(key, wk_entry, accum_prop, wk, total_item, dias_semana)
            cell = ws.cell(row=row, column=ncols_fixed + 1 + j_wk, value=val)
            cell.fill = f
            cell.border = _border()
            cell.alignment = _align("right")
            if val is not None:
                cell.number_format = fmt
        row += 1
    return row


def render_cronograma(wb: Workbook, canonical: CanonicalModel) -> None:
    """CRONOGRAMA SEMANAL — matriz itens × semanas.

    Colunas: semanas ISO ordenadas cronologicamente.
        Por item EAP, 6 sub-linhas:
      1. Dias Trabalhados   2. Nº De Equipes
            3. Qtd. Planejada     4. Qtd Plan. Acumulada
            5. % Semanal          6. % Acumulado
    Cabeçalho da semana: DD/MMM (data da segunda-feira).
    """
    ws = wb.create_sheet("CRONOGRAMA.DET")

    acts = canonical.cronograma.scheduled_activities if canonical.cronograma else []
    act_lookup, item_totals, all_weeks = _build_act_lookup_and_totals(acts)
    n_weeks = len(all_weeks)

    if n_weeks == 0:
        ws.cell(row=1, column=1).value = "Sem dados de cronograma semanal."
        return

    dias_semana: int = canonical.projeto.regime_trabalho.dias_semana if canonical.projeto else 6

    NCOLS_FIXED = 5  # Item ID | Descrição | Un. | Quantitativo | Métrica
    NCOLS = NCOLS_FIXED + n_weeks
    _set_col_widths(ws, [10, 30, 6, 12, 22] + [13] * n_weeks)

    row = 1
    _title_row(ws, row, "CRONOGRAMA SEMANAL — ACOMPANHAMENTO FÍSICO DE PRODUÇÃO", NCOLS)
    row += 1
    col_headers = ["Item ID", "Descrição", "Un.", "Quantitativo", "Métrica"] + [_wk_label(wk) for wk in all_weeks]
    _header_row(ws, row, col_headers)
    row += 1
    pratic_map = _get_pratic_map(canonical)
    _render_praticabilidade_row(ws, row, all_weeks, pratic_map)
    row += 1
    ws.freeze_panes = f"F{row}"

    FILL_ITEM_EVEN = _fill(_C["orange_lt"])
    FILL_ITEM_ODD = PatternFill(start_color="EAD9C3", end_color="EAD9C3", fill_type="solid")

    for i_item, item in enumerate(canonical.eap):
        wk_data = act_lookup.get(item.item_id, {})
        total_item = item_totals.get(item.item_id, 0.0)
        accum_prop, last_acum = _compute_accum_prop(wk_data, all_weeks)
        if total_item <= 0 and last_acum > 0:
            total_item = last_acum
        fill_item_hdr = FILL_ITEM_EVEN if i_item % 2 == 0 else FILL_ITEM_ODD
        fill_sub = _fill(_C["white"]) if i_item % 2 == 0 else _fill(_C["row_alt"])
        row = _render_item_metrics_rows(
            ws,
            row,
            item,
            wk_data,
            accum_prop,
            total_item,
            all_weeks,
            NCOLS_FIXED,
            dias_semana,
            fill_item_hdr,
            fill_sub,
        )


