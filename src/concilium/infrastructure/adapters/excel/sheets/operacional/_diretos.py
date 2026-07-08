"""Excel adapter — aba OPERACIONAL — render_diretos (equipes diretas: mão de obra, equipamento, material)."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from concilium.domain.models.canonical import CanonicalModel, EquipeDireta

from ...helpers import _find_eap_item
from ...styles import (
    _C,
    _align,
    _auto_fit_columns,
    _border,
    _fill,
    _font,
    _freeze,
    _header_row,
    _label,
    _money,
    _section_row,
    _set_col_widths,
    _title_row,
)


def _xl_diretos_labor(
    ws: Worksheet, row: int, eq: EquipeDireta, fill_base: PatternFill | None, PAPEL_COLOR: dict
) -> int:
    for func in eq.funcoes:
        custo_unit = func.custo_mensal_unitario if func.custo_mensal_unitario else 0.0
        papel = getattr(func, "papel", "direto") or "direto"
        papel_color = PAPEL_COLOR.get(papel, _C["preto"])
        _label(ws, row, 1, "MÃO DE OBRA", fill_base)
        _label(ws, row, 2, func.funcao_id, fill_base)
        _label(ws, row, 3, func.descricao or func.funcao_id, fill_base)
        _label(ws, row, 4, papel.upper(), fill_base, bold=True, color=papel_color, align="center")
        _label(ws, row, 5, func.num_pessoas, fill_base, align="right")
        _money(ws, row, 6, custo_unit, fill_base)
        _money(ws, row, 7, f"=E{row}*F{row}", fill_base)
        _label(ws, row, 8, round(eq.permanencia_meses, 2), fill_base, align="right")
        _money(ws, row, 9, f"=G{row}*H{row}", fill_base)
        row += 1
    return row


def _xl_diretos_equip(ws: Worksheet, row: int, eq: EquipeDireta, fill_base: PatternFill | None) -> int:
    for equip in eq.equipamentos:
        n_un = getattr(equip, "num_unidades", 1) or 1
        custo_equip_u = getattr(equip, "custo_mensal", 0.0) or 0.0
        _label(ws, row, 1, "EQUIPAMENTO", fill_base)
        _label(ws, row, 2, equip.equip_id, fill_base)
        _label(ws, row, 3, getattr(equip, "descricao", equip.equip_id), fill_base)
        _label(ws, row, 4, "—", fill_base, align="center")
        _label(ws, row, 5, n_un, fill_base, align="right")
        _money(ws, row, 6, custo_equip_u, fill_base)
        _money(ws, row, 7, f"=E{row}*F{row}", fill_base)
        _label(ws, row, 8, round(eq.permanencia_meses, 2), fill_base, align="right")
        _money(ws, row, 9, f"=G{row}*H{row}", fill_base)
        row += 1
        custo_comb_u = getattr(equip, "custo_combustivel_mensal", 0.0) or 0.0
        if custo_comb_u > 0:
            consumo_str = f"{equip.consumo_litros_hora:.1f} L/h" if getattr(equip, "consumo_litros_hora", None) else "—"
            _label(ws, row, 1, "COMBUSTÍVEL", fill_base, bold=False, color=_C["orange"])
            _label(ws, row, 2, equip.equip_id, fill_base)
            _label(ws, row, 3, f"Diesel S-10 ({consumo_str})", fill_base)
            _label(ws, row, 4, "—", fill_base, align="center")
            _label(ws, row, 5, n_un, fill_base, align="right")
            _money(ws, row, 6, custo_comb_u, fill_base)
            _money(ws, row, 7, f"=E{row}*F{row}", fill_base)
            _label(ws, row, 8, round(eq.permanencia_meses, 2), fill_base, align="right")
            _money(ws, row, 9, f"=G{row}*H{row}", fill_base)
            row += 1
    return row


def _xl_diretos_material(ws: Worksheet, row: int, eq: EquipeDireta, fill_base: PatternFill | None) -> int:
    custo_mat = getattr(eq, "custo_material_mensal", 0.0) or 0.0
    if custo_mat > 0:
        _label(ws, row, 1, "MATERIAL", fill_base, bold=False, color=_C["nav_md"])
        _label(ws, row, 2, eq.item_id, fill_base)
        _label(ws, row, 3, "Materiais diretos", fill_base)
        _label(ws, row, 4, "—", fill_base, align="center")
        _label(ws, row, 5, 1, fill_base, align="right")
        _money(ws, row, 6, custo_mat, fill_base)
        _money(ws, row, 7, f"=E{row}*F{row}", fill_base)
        _label(ws, row, 8, round(eq.permanencia_meses, 2), fill_base, align="right")
        _money(ws, row, 9, f"=G{row}*H{row}", fill_base)
        row += 1
    return row


def _xl_diretos_subtotal(
    ws: Worksheet, row: int, eq: EquipeDireta, fill_sub: PatternFill, start_row: int, subtotal_rows: list
) -> None:
    _label(ws, row, 1, f"Subtotal {eq.item_id}", fill_sub, bold=True)
    for c in range(2, 9):
        ws.cell(row=row, column=c).fill = fill_sub
        ws.cell(row=row, column=c).border = _border()
    _money(ws, row, 7, f"=SUM(G{start_row}:G{row - 1})", fill_sub, bold=True, color=_C["orange"])
    _label(ws, row, 8, round(eq.permanencia_meses, 2), fill_sub, bold=True, align="right")
    _money(ws, row, 9, f"=SUM(I{start_row}:I{row - 1})", fill_sub, bold=True, color=_C["orange"])
    subtotal_rows.append(row)


def render_diretos(wb: Workbook, canonical: CanonicalModel) -> int:
    """DIRETOS — composição granular por item EAP: mão de obra (com papel), equipamento,
    combustível (linha separada por equipamento) e materiais diretos.

    Retorna o número da linha do TOTAL DIRETO para referência cruzada em RESULTADO.
    """
    ws = wb.create_sheet("DIRETOS.DET")
    NCOLS = 9
    _set_col_widths(ws, [14, 20, 36, 10, 6, 16, 16, 10, 16])
    _freeze(ws, "A3")
    row = 1
    _title_row(ws, row, "DIRETOS — EQUIPES E RECURSOS", NCOLS)
    row += 1
    headers = [
        "Tipo",
        "Código",
        "Descrição",
        "Papel",
        "Qtd",
        "Custo Unit/Mês (R$)",
        "Custo Total/Mês (R$)",
        "Perm. (meses)",
        "Custo Total (R$)",
    ]
    _header_row(ws, row, headers)
    row += 1
    ws.freeze_panes = f"A{row}"
    _PAPEL_COLOR = {"direto": _C["verde"], "indireto": _C["nav_md"], "gestao": _C["orange"]}
    subtotal_rows: list[int] = []
    for i_eq, eq in enumerate(canonical.equipes_diretas):
        eap_item = _find_eap_item(canonical.eap, eq.item_id)
        descricao_item = eap_item.descricao if eap_item else eq.item_id
        _section_row(ws, row, f"{eq.item_id} — {descricao_item}", NCOLS)
        row += 1
        start_row = row
        fill_base = _fill(_C["white"]) if i_eq % 2 == 0 else _fill(_C["row_alt"])
        row = _xl_diretos_labor(ws, row, eq, fill_base, _PAPEL_COLOR)
        row = _xl_diretos_equip(ws, row, eq, fill_base)
        row = _xl_diretos_material(ws, row, eq, fill_base)
        fill_sub = _fill(_C["orange_lt"])
        _xl_diretos_subtotal(ws, row, eq, fill_sub, start_row, subtotal_rows)
        row += 1
    row += 1
    total_row = row
    fill_tot = _fill(_C["nav_dk"])
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    cell = ws.cell(row=row, column=1, value="TOTAL DIRETO")
    cell.fill = fill_tot
    cell.font = _font(bold=True, color=_C["white"], size=11)
    cell.alignment = _align("right")
    cell.border = _border()
    f_tot = f"={'+'.join(f'I{r}' for r in subtotal_rows)}" if subtotal_rows else 0.0
    _money(ws, row, 9, f_tot, fill_tot, bold=True, color=_C["white"])
    wb.defined_names["TOTAL_DIRETO"] = DefinedName(
        "TOTAL_DIRETO",
        attr_text=f"'DIRETOS.DET'!${get_column_letter(9)}${total_row}",
    )
    _auto_fit_columns(ws)
    return total_row


