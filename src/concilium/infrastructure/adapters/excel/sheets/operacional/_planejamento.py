"""Excel adapter — aba OPERACIONAL — render_planejamento_ml (linhas de planejamento)."""

from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from concilium.domain.models.canonical import CanonicalModel

from ...styles import (
    _C,
    _align,
    _auto_fit_columns,
    _border,
    _fill,
    _font,
    _freeze,
    _header_row,
    _money,
    _set_col_widths,
    _title_row,
)


def _xl_plan_sem_vals(sem: Any) -> tuple[float, float, float, float]:
    """s/BDI component values (mo, eq, diesel, mat) — 4 ternaries."""
    return (
        sem.mao_obra if sem else 0.0,
        sem.equipamentos if sem else 0.0,
        sem.diesel if sem else 0.0,
        sem.material if sem else 0.0,
    )


def _xl_plan_row_fill(i: int, classe: str | None, CRIT_HEX: dict) -> Any:
    """Row fill based on criticidade class."""
    crit_hex = CRIT_HEX.get(classe) if classe else None
    return _fill(crit_hex) if crit_hex else (_fill(_C["white"]) if i % 2 == 0 else _fill(_C["row_alt"]))


def _xl_plan_row_vals(
    p: Any,
    eap_item: Any,
    classe: str | None,
    row: int,
    has_crit: bool,
    mo_s: float,
    eq_s: float,
    dsl_s: float,
    mat_s: float,
) -> list[Any]:
    """Construct the values list for one planning row."""
    tot_s = f"=E{row}+G{row}+I{row}+K{row}"
    tot_c = f"=F{row}+H{row}+J{row}+L{row}"
    mo_c = f"=E{row}*(1+'MEMORIAL.BDI'!$M$3)"
    eq_c = f"=G{row}*(1+'MEMORIAL.BDI'!$M$5)"
    dsl_c = f"=I{row}*(1+'MEMORIAL.BDI'!$M$7)"
    mat_c = f"=K{row}*(1+'MEMORIAL.BDI'!$M$4)"
    vals: list[Any] = [
        p.item_id,
        eap_item.descricao if eap_item else p.cpu_id,
        eap_item.unidade if eap_item else "un",
        p.quantitativo,
        mo_s,
        mo_c,
        eq_s,
        eq_c,
        dsl_s,
        dsl_c,
        mat_s,
        mat_c,
        tot_s,
        tot_c,
    ]
    if has_crit:
        vals.append(classe or "—")
    return vals


def _xl_plan_write_data_row(ws: Any, row: int, vals: list[Any], row_fill: Any) -> None:
    """Write one planning row with numeric formatting."""
    for col, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.fill = row_fill
        cell.border = _border()
        if 5 <= col <= 14:
            cell.number_format = "#,##0.00"
            cell.alignment = _align("right")
        elif col == 4:
            cell.number_format = "#,##0.##"
            cell.alignment = _align("right")
        else:
            cell.alignment = _align("left")


def _xl_plan_data_rows(
    ws: Any,
    row: int,
    planejamento_sorted: list[Any],
    canonical_eap: list[Any],
    has_crit: bool,
    criticidade_lookup: dict,
    CRIT_HEX: dict,
) -> int:
    """Write all EAP planning data rows; returns updated row."""
    for i, p in enumerate(planejamento_sorted):
        eap_item = next((e for e in canonical_eap if e.item_id == p.item_id), None)
        sem = p.subtotal.sem_bdi if p.subtotal else None
        classe = criticidade_lookup.get(p.item_id) if has_crit else None
        mo_s, eq_s, dsl_s, mat_s = _xl_plan_sem_vals(sem)
        row_fill = _xl_plan_row_fill(i, classe, CRIT_HEX)
        vals = _xl_plan_row_vals(p, eap_item, classe, row, has_crit, mo_s, eq_s, dsl_s, mat_s)
        _xl_plan_write_data_row(ws, row, vals, row_fill)
        row += 1
    return row


def render_planejamento_ml(wb: Workbook, canonical: CanonicalModel) -> int | None:
    """PLANEJAMENTO — EAP ordenada topologicamente, componentes s/BDI e c/BDI por item.

      Colunas (sem criticidade = 14, com criticidade = 15):
    1  Item ID
    2  Descrição
    3  Un.
    4  Quantitativo
          5  MO s/BDI
          6  MO c/BDI
          7  EQ s/BDI
          8  EQ c/BDI
          9  Combustível s/BDI
          10 Combustível c/BDI
          11 Material s/BDI
          12 Material c/BDI
          13 Total s/BDI
          14 Total c/BDI         ← TOTAL_COM_BDI
    15 Criticidade (se disponível)
    """
    ws = wb.create_sheet("PLANEJAMENTO")

    criticidade_lookup: dict[str, str] = {}
    if canonical.cost_criticality_report is not None:
        for prof in canonical.cost_criticality_report.atividades:
            criticidade_lookup[prof.atividade_id] = prof.classe
    has_crit = bool(criticidade_lookup)

    _CRIT_HEX = {"A+": "FF4444", "A": "FFA500", "B": "FFFF33"}

    # Ordenação topológica pelo campo EapItem.ordem
    eap_ordem: dict[str, int] = {e.item_id: (e.ordem if e.ordem is not None else 9999) for e in canonical.eap}
    planejamento_sorted = sorted(
        canonical.planejamento,
        key=lambda p: eap_ordem.get(p.item_id, 9999),
    )

    NCOLS = 15 if has_crit else 14
    _set_col_widths(ws, [12, 42, 6, 13, 15, 15, 13, 13, 17, 15, 15, 13, 13, 17, 12][:NCOLS])
    _freeze(ws, "C3")

    row = 1
    _title_row(ws, row, "PLANEJAMENTO — ORÇAMENTO POR EAP (COMPONENTES s/BDI e c/BDI)", NCOLS)
    row += 1

    headers = [
        "Item ID",
        "Descrição",
        "Un.",
        "Quantitativo",
        "MO s/BDI (R$)",
        "MO c/BDI (R$)",
        "EQ s/BDI (R$)",
        "EQ c/BDI (R$)",
        "Comb. s/BDI (R$)",
        "Comb. c/BDI (R$)",
        "MAT s/BDI (R$)",
        "MAT c/BDI (R$)",
        "Total s/BDI (R$)",
        "Total c/BDI (R$)",
    ]
    if has_crit:
        headers.append("Criticidade")
    _header_row(ws, row, headers)
    row += 1
    ws.freeze_panes = f"C{row}"

    data_start = row
    row = _xl_plan_data_rows(ws, row, planejamento_sorted, canonical.eap, has_crit, criticidade_lookup, _CRIT_HEX)

    # Linha em branco + total
    row += 1
    total_row = row
    fill_tot = _fill(_C["nav_dk"])
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(row=row, column=1, value="TOTAL")
    cell.fill = fill_tot
    cell.font = _font(bold=True, color=_C["white"], size=11)
    cell.alignment = _align("right")
    cell.border = _border()
    for c in range(5, 15):
        col_letter = get_column_letter(c)
        _money(
            ws, row, c, f"=SUM({col_letter}{data_start}:{col_letter}{row - 1})", fill_tot, bold=True, color=_C["white"]
        )

    # Named ranges — col 13 = sem_bdi, col 14 = com_bdi (para aba RESULTADO)
    wb.defined_names["TOTAL_SEM_BDI"] = DefinedName(
        "TOTAL_SEM_BDI",
        attr_text=f"PLANEJAMENTO!${get_column_letter(13)}${total_row}",
    )
    wb.defined_names["TOTAL_COM_BDI"] = DefinedName(
        "TOTAL_COM_BDI",
        attr_text=f"PLANEJAMENTO!${get_column_letter(14)}${total_row}",
    )
    _auto_fit_columns(ws)
    return total_row


