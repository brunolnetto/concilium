from __future__ import annotations

from openpyxl import Workbook

from concilium.domain.models.canonical import CanonicalModel

from ..styles import (
    _C,
    _align,
    _auto_fit_columns,
    _border,
    _fill,
    _font,
    _freeze,
    _header_row,
    _label,
    _money,
    _section_row,
    _set_col_widths,
    _title_row,
)


def _calendario_clima_totais_por_semana(canonical: CanonicalModel) -> tuple[dict[str, float], dict[str, float]]:
    """Extrai fator_clima e fator_total por semana ISO a partir do cronograma."""
    clima_por_semana: dict[str, float] = {}
    total_por_semana: dict[str, float] = {}
    cron = getattr(canonical, "cronograma", None)
    if cron:
        for sa in cron.scheduled_activities:
            for semana_iso, dist in sa.distribuicao_semanal.items():
                if semana_iso not in clima_por_semana:
                    clima_por_semana[semana_iso] = float(getattr(dist, "fator_clima", getattr(dist, "eficiencia", 1.0)))
                    total_por_semana[semana_iso] = float(getattr(dist, "fator_total", getattr(dist, "eficiencia", 1.0)))
    return clima_por_semana, total_por_semana


def _calendario_row_data(item) -> dict:
    """Normaliza um item do calendário operacional (model ou dict) em um dict simples."""
    if hasattr(item, "model_dump"):
        return item.model_dump()
    if isinstance(item, dict):
        return item
    return {}


def _render_calendario_row(
    ws, row: int, idx: int, item, clima_por_semana: dict[str, float], total_por_semana: dict[str, float]
) -> None:
    """Renderiza uma linha da tabela CALENDARIO.OP."""
    fill = _fill(_C["row_alt"]) if idx % 2 else _fill(_C["white"])
    row_data = _calendario_row_data(item)

    semana = str(row_data.get("semana_iso", ""))
    _label(ws, row, 1, semana, fill)
    _label(ws, row, 2, float(row_data.get("dias_uteis_base", 0.0) or 0.0), fill, align="right")
    _label(ws, row, 3, float(row_data.get("dias_uteis_efetivos", 0.0) or 0.0), fill, align="right")
    _fc_cal = row_data.get("fator_calendario")
    _label(ws, row, 4, float(_fc_cal if _fc_cal is not None else 1.0), fill, align="right")
    _label(ws, row, 5, round(clima_por_semana.get(semana, 1.0), 4), fill, align="right")
    _label(ws, row, 6, round(total_por_semana.get(semana, clima_por_semana.get(semana, 1.0)), 4), fill, align="right")
    eventos = row_data.get("eventos_aplicados", [])
    eventos_txt = ", ".join(eventos) if isinstance(eventos, (list, tuple)) else str(eventos or "")
    _label(ws, row, 7, eventos_txt, fill)


def render_calendario_operacional(wb: Workbook, canonical: CanonicalModel) -> None:
    """ADR-068: renderiza visão semanal de clima x calendário x fator total."""
    ws = wb.create_sheet("CALENDARIO.OP")
    _title_row(ws, 1, "CALENDARIO OPERACIONAL", 7)
    ws.cell(
        row=2,
        column=1,
        value=(
            "Relevancia: evidencia a perda/ganho de produtividade por calendario e clima, "
            "base para validar prazo e permanencia das equipes."
        ),
    ).font = _font(italic=True, size=9)
    _header_row(
        ws,
        4,
        [
            "Semana ISO",
            "Dias úteis base",
            "Dias úteis efetivos",
            "Fator calendário",
            "Fator clima",
            "Fator total",
            "Eventos aplicados",
        ],
    )

    _set_col_widths(ws, [14, 16, 18, 16, 12, 12, 38])
    row = 5

    calendario_rows = list(getattr(canonical, "calendario_operacional", []) or [])
    if not calendario_rows:
        _label(ws, row, 1, "Calendário operacional não disponível.", _fill(_C["row_alt"]))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        _freeze(ws, "A5")
        return

    clima_por_semana, total_por_semana = _calendario_clima_totais_por_semana(canonical)

    for idx, item in enumerate(calendario_rows):
        _render_calendario_row(ws, row, idx, item, clima_por_semana, total_por_semana)
        row += 1

    _freeze(ws, "A5")


def _req_render_equipe_subtotal(ws, row: int, eq, ncols: int, equipe_total_rows: list) -> int:
    """Render subtotal row for one equipe. Returns next row."""
    fill_sub = _fill(_C["orange_lt"])
    _label(ws, row, 1, f"Subtotal {eq.item_id}", fill_sub, bold=True)
    for col in range(2, ncols):
        ws.cell(row=row, column=col).fill = fill_sub
        ws.cell(row=row, column=col).border = _border()
    _money(ws, row, ncols, eq.custo_total, fill_sub, bold=True, color=_C["orange"])
    equipe_total_rows.append(row)
    return row + 1


def _req_render_func_row(ws, row: int, eq, func, fill, papel_color: dict) -> int:
    """Render one funcao row in EQUIPE sheet. Returns next row."""
    papel = getattr(func, "papel", "direto") or "direto"
    cm = func.custo_mensal_unitario * func.num_pessoas
    _label(ws, row, 1, eq.item_id, fill)
    _label(ws, row, 2, eq.cpu_id, fill)
    _label(ws, row, 3, func.descricao or func.funcao_id, fill)
    _label(ws, row, 4, papel.upper(), fill, bold=True, color=papel_color.get(papel, _C["preto"]), align="center")
    _label(ws, row, 5, func.num_pessoas, fill, align="right")
    _money(ws, row, 6, cm, fill)
    _label(ws, row, 7, round(eq.permanencia_meses, 2), fill, align="right")
    _label(ws, row, 8, eq.num_equipes_base, fill, align="right")
    _money(ws, row, 9, cm * eq.permanencia_meses * eq.num_equipes_base, fill)
    return row + 1


def _req_render_equip_rows(ws, row: int, eq, equip, fill) -> int:
    """Render equipment row (+ optional combustivel row) in EQUIPE sheet. Returns next row."""
    cm_eq = getattr(equip, "custo_mensal", 0.0) or 0.0
    desc = getattr(equip, "descricao", equip.equip_id) or equip.equip_id
    _label(ws, row, 1, eq.item_id, fill)
    _label(ws, row, 2, equip.equip_id, fill)
    _label(ws, row, 3, desc + " [EQ]", fill)
    _label(ws, row, 4, "EQ", fill, align="center", color=_C["nav_dk"])
    _label(ws, row, 5, equip.num_unidades, fill, align="right")
    _money(ws, row, 6, cm_eq, fill)
    _label(ws, row, 7, round(eq.permanencia_meses, 2), fill, align="right")
    _label(ws, row, 8, eq.num_equipes_base, fill, align="right")
    _money(ws, row, 9, cm_eq * eq.permanencia_meses * eq.num_equipes_base, fill)
    row += 1
    comb = getattr(equip, "custo_combustivel_mensal", 0.0) or 0.0
    if comb > 0:
        _label(ws, row, 1, eq.item_id, fill)
        _label(ws, row, 2, equip.equip_id, fill)
        _label(ws, row, 3, f"Diesel — {desc}", fill, color=_C["orange"])
        _label(ws, row, 4, "COMB", fill, align="center", color=_C["orange"])
        _label(ws, row, 5, equip.num_unidades, fill, align="right")
        _money(ws, row, 6, comb, fill)
        _label(ws, row, 7, round(eq.permanencia_meses, 2), fill, align="right")
        _label(ws, row, 8, eq.num_equipes_base, fill, align="right")
        _money(ws, row, 9, comb * eq.permanencia_meses * eq.num_equipes_base, fill)
        row += 1
    return row


def render_equipe(wb: Workbook, canonical: CanonicalModel) -> None:
    """EQUIPE.DIRETO — Composição detalhada de funções e equipamentos por equipe direta."""
    ws = wb.create_sheet("EQUIPE.DIRETO")
    NCOLS = 9
    _set_col_widths(ws, [12, 16, 28, 8, 6, 18, 10, 8, 18])
    _freeze(ws, "A3")
    row = 1
    _title_row(ws, row, "EQUIPE.DIRETO — COMPOSIÇÃO DE RECURSOS DIRETOS POR ATIVIDADE", NCOLS)
    row += 1
    _header_row(
        ws,
        row,
        [
            "Item ID",
            "CPU ID",
            "Nome / Tipo",
            "Papel",
            "Qtd",
            "Custo Mensal (R$)",
            "Perm. (m)",
            "Nº Eq.",
            "Custo Total (R$)",
        ],
    )
    row += 1
    ws.freeze_panes = f"A{row}"
    _PAPEL_COLOR = {"direto": _C["verde"], "indireto": _C["nav_md"], "gestao": _C["orange"]}
    equipe_total_rows = []
    for i_eq, eq in enumerate(canonical.equipes_diretas):
        eap_item = next((e for e in canonical.eap if e.item_id == eq.item_id), None)
        item_label = f"{eq.item_id} — {eap_item.descricao if eap_item else eq.item_id}"
        _section_row(ws, row, item_label, NCOLS)
        row += 1
        fill = _fill(_C["white"]) if i_eq % 2 == 0 else _fill(_C["row_alt"])
        for func in eq.funcoes:
            row = _req_render_func_row(ws, row, eq, func, fill, _PAPEL_COLOR)
        for equip in eq.equipamentos:
            row = _req_render_equip_rows(ws, row, eq, equip, fill)
        row = _req_render_equipe_subtotal(ws, row, eq, NCOLS, equipe_total_rows)
    row += 1
    fill_tot = _fill(_C["nav_dk"])
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS - 1)
    cell = ws.cell(row=row, column=1, value="TOTAL EQUIPES DIRETAS")
    cell.fill = fill_tot
    cell.font = _font(bold=True, color=_C["white"], size=11)
    cell.alignment = _align("right")
    cell.border = _border()
    f_tot = ("=" + "+".join(f"I{r}" for r in equipe_total_rows)) if equipe_total_rows else "0"
    _money(ws, row, NCOLS, f_tot, fill_tot, bold=True, color=_C["white"])


def _find_indireto_entry(canonical: CanonicalModel, cat: str):
    """Localiza a entrada de indiretos cuja categoria (uppercased) bate com `cat`."""
    return next((ind for ind in (canonical.indiretos or []) if str(getattr(ind, "categoria", "")).upper() == cat), None)


def _render_indireto_categoria_ausente(ws, row: int, cat: str) -> None:
    """Renderiza o aviso de categoria não encontrada e finaliza a planilha."""
    ws.cell(row=row, column=1, value=f"Categoria {cat} não encontrada em indiretos.").font = _font(italic=True, size=9)
    _auto_fit_columns(ws)


def _render_indireto_resumo_custos(ws, row: int, entry) -> None:
    """Renderiza a linha de resumo de custos (total/MO/EQ/Comb) da categoria."""
    ws.cell(
        row=row,
        column=1,
        value=(
            f"Custo total: R$ {float(getattr(entry, 'custo_total', 0.0) or 0.0):,.2f} | "
            f"MO: R$ {float(getattr(entry, 'custo_mao_obra', 0.0) or 0.0):,.2f} | "
            f"EQ: R$ {float(getattr(entry, 'custo_equipamentos', 0.0) or 0.0):,.2f} | "
            f"Comb: R$ {float(getattr(entry, 'custo_combustivel', 0.0) or 0.0):,.2f}"
        ),
    ).font = _font(italic=True, size=9)


def _write_indireto_detail_row(ws, row: int, fill, vals: list) -> None:
    """Escreve uma linha de detalhe (natureza/id/descrição/qtd/un/custos/observação)."""
    for col, val in enumerate(vals, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.fill = fill
        cell.border = _border()
        if col in (6, 7):
            cell.number_format = "R$ #,##0.00"
        if col == 4:
            cell.number_format = "0.00"
            cell.alignment = _align("right")


def _first_field(d: dict, *keys: str, default: object = 0):
    """Retorna o valor da primeira chave presente e truthy em ``d``, ou ``default``."""
    for k in keys:
        v = d.get(k)
        if v:
            return v
    return default


def _build_indireto_funcao_vals(f: dict) -> list:
    """Monta os valores da linha de detalhe para uma função de mão de obra indireta."""
    qtd = float(_first_field(f, "num_pessoas", "quantidade"))
    unit = float(_first_field(f, "custo_mensal_unitario", "custo_unitario"))
    total = float(_first_field(f, "custo_total", default=qtd * unit))
    return [
        "MÃO DE OBRA",
        _first_field(f, "funcao_id", "id", default="-"),
        _first_field(f, "descricao", default="-"),
        qtd,
        "pax",
        unit,
        total,
        _first_field(f, "origem", default=""),
    ]


def _build_indireto_equipamento_vals(e: dict) -> list:
    """Monta os valores da linha de detalhe para um equipamento indireto."""
    qtd = float(_first_field(e, "num_unidades", "quantidade"))
    unit = float(_first_field(e, "custo_mensal", "custo_unitario"))
    total = float(_first_field(e, "custo_total", default=qtd * unit))
    return [
        "EQUIPAMENTO",
        _first_field(e, "equip_id", "id", default="-"),
        _first_field(e, "descricao", default="-"),
        qtd,
        "un",
        unit,
        total,
        _first_field(e, "origem", default=""),
    ]


def _render_indireto_funcoes(ws, row: int, entry, fills: list) -> tuple[int, bool]:
    """Renderiza as linhas de mão de obra indireta. Retorna (próxima linha, teve linhas)."""
    has_rows = False
    for idx, f in enumerate(getattr(entry, "funcoes", []) or []):
        fill = fills[idx % 2]
        _write_indireto_detail_row(ws, row, fill, _build_indireto_funcao_vals(f))
        row += 1
        has_rows = True
    return row, has_rows


def _render_indireto_equipamentos(ws, row: int, entry, fills: list, start_idx: int) -> tuple[int, bool]:
    """Renderiza as linhas de equipamento indireto. Retorna (próxima linha, teve linhas)."""
    has_rows = False
    for off, e in enumerate(getattr(entry, "equipamentos", []) or []):
        fill = fills[(start_idx + off) % 2]
        _write_indireto_detail_row(ws, row, fill, _build_indireto_equipamento_vals(e))
        row += 1
        has_rows = True
    return row, has_rows


def render_equipe_indireto(wb: Workbook, canonical: CanonicalModel, categoria: str) -> None:
    """EQUIPE.INDIRETO.<CAT> — detalhamento de recursos indiretos por categoria."""
    cat = str(categoria or "").upper()
    ws = wb.create_sheet(f"EQUIPE.INDIRETO.{cat}")
    NCOLS = 8
    _set_col_widths(ws, [18, 18, 30, 8, 10, 16, 16, 28])
    row = 1
    _title_row(ws, row, f"EQUIPE.INDIRETO.{cat} — RECURSOS INDIRETOS", NCOLS)
    row += 1

    entry = _find_indireto_entry(canonical, cat)
    if entry is None:
        _render_indireto_categoria_ausente(ws, row, cat)
        return

    _render_indireto_resumo_custos(ws, row, entry)
    row += 2

    _header_row(ws, row, ["Natureza", "ID", "Descrição", "Qtd", "Un.", "Custo Unit.", "Custo Total", "Observação"])
    row += 1

    fills = [_fill(_C["white"]), _fill(_C["row_alt"])]
    row, has_rows_funcoes = _render_indireto_funcoes(ws, row, entry, fills)

    start_idx = len(getattr(entry, "funcoes", []) or [])
    row, has_rows_equipamentos = _render_indireto_equipamentos(ws, row, entry, fills, start_idx)

    has_rows = has_rows_funcoes or has_rows_equipamentos
    if not has_rows:
        ws.cell(
            row=row, column=1, value=f"Sem equipes explícitas para {cat}; apenas total agregado em INDIRETOS.DET/RES."
        ).font = _font(
            italic=True,
            size=9,
        )

    _auto_fit_columns(ws)
