"""Excel sheets — resultado (parte extraída de resultado.py, grupo _cashflow)."""

from __future__ import annotations

import math

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from concilium.domain.models.canonical import CanonicalModel

from ...cash_flow import _cash_flow
from ...helpers import _project_months
from ...styles import (
    _C,
    _align,
    _border,
    _fill,
    _font,
    _freeze,
    _header_row,
    _label,
    _money,
    _section_row,
    _set_col_widths,
    _title_row,
)


def _xl_cashflow_row_styling(rec: dict) -> tuple:
    """Return (fill_saldo, color_saldo, fill_sa, color_sa, fill_exp, color_exp)."""
    saldo = rec["saldo"]
    fill_saldo = _fill(_C["verde_lt"]) if saldo >= 0 else _fill(_C["verm_lt"])
    color_saldo = _C["verde"] if saldo >= 0 else _C["vermelho"]
    saldo_atual = rec["saldo_atual"]
    fill_sa = _fill(_C["verde_lt"]) if saldo_atual >= 0 else _fill(_C["verm_lt"])
    color_sa = _C["verde"] if saldo_atual >= 0 else _C["vermelho"]
    exp = rec["exposicao_acumulada"]
    fill_exp = _fill(_C["verm_lt"]) if exp > 0 else _fill(_C["verde_lt"])
    color_exp = _C["vermelho"] if exp > 0 else _C["verde"]
    return fill_saldo, color_saldo, fill_sa, color_sa, fill_exp, color_exp


def _xl_cashflow_metrics_block(ws: Worksheet, row: int, metricas: list, NCOLS: int) -> int:
    for label, valor in metricas:
        fill_m = _fill(_C["cinza_lt"])
        _label(ws, row, 1, label, fill_m, bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell_v = ws.cell(row=row, column=3, value=valor)
        cell_v.border = _border()
        cell_v.alignment = _align("right")
        if isinstance(valor, float):
            cell_v.number_format = "R$ #,##0.00"
            is_bad = (label == "Exposição máxima" and valor > 0) or (label == "Saldo final do projeto" and valor < 0)
            cell_v.font = _font(bold=True, color=_C["vermelho"] if is_bad else _C["verde"])
            cell_v.fill = _fill(_C["verm_lt"] if is_bad else _C["verde_lt"])
        else:
            cell_v.font = _font(bold=True)
            cell_v.fill = fill_m
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=NCOLS)
        row += 1
    return row


def _xl_cashflow_render_row(ws: Worksheet, row: int, rec: dict, i: int) -> None:
    fill_base = _fill(_C["white"]) if i % 2 == 0 else _fill(_C["row_alt"])
    fill_saldo, color_saldo, fill_sa, color_sa, fill_exp, color_exp = _xl_cashflow_row_styling(rec)
    _label(ws, row, 1, rec["mes_label"], fill_base, bold=False, align="center")
    _money(ws, row, 2, rec["despesa"], fill_base)
    _money(ws, row, 3, rec["receita"], fill_base)
    _money(ws, row, 4, rec["receita_acumulada"], fill_base)
    _money(ws, row, 5, rec["saldo"], fill_saldo, color=color_saldo)
    _money(ws, row, 6, rec["saldo_atual"], fill_sa, color=color_sa)
    _money(ws, row, 7, rec["exposicao_acumulada"], fill_exp, color=color_exp)


def render_fluxo_caixa(wb: Workbook, canonical: CanonicalModel) -> None:
    """FLUXO DE CAIXA — despesa, receita, saldo e exposição acumulada por mês."""
    ws = wb.create_sheet("ANALISE.FC_CAIXA")
    cf = _cash_flow(canonical)
    n_exec = len(_project_months(canonical))
    NCOLS = 7
    _set_col_widths(ws, [10, 18, 18, 18, 16, 16, 20])
    _freeze(ws, "A3")
    row = 1
    _title_row(ws, row, "FLUXO DE CAIXA — ACOMPANHAMENTO MENSAL", NCOLS)
    row += 1
    headers = [
        "Mês",
        "Despesa (R$)",
        "Receita (R$)",
        "Rec. Acumulada (R$)",
        "Saldo do Mês (R$)",
        "Saldo Atual (R$)",
        "Exposição Acum. (R$)",
    ]
    _header_row(ws, row, headers)
    row += 1
    ws.freeze_panes = f"A{row}"
    total_despesa = 0.0
    total_receita = 0.0
    max_exposicao = 0.0
    mes_max_exp = ""
    for i, rec in enumerate(cf):
        _xl_cashflow_render_row(ws, row, rec, i)
        total_despesa += rec["despesa"]
        total_receita += rec["receita"]
        if rec["exposicao_acumulada"] > max_exposicao:
            max_exposicao = rec["exposicao_acumulada"]
            mes_max_exp = rec["mes_label"]
        row += 1
    row += 1
    fill_tot = _fill(_C["nav_dk"])
    _label(ws, row, 1, "TOTAL", fill_tot, bold=True, color=_C["white"], align="center")
    _money(ws, row, 2, total_despesa, fill_tot, bold=True, color=_C["white"])
    _money(ws, row, 3, total_receita, fill_tot, bold=True, color=_C["white"])
    _money(ws, row, 4, total_receita, fill_tot, bold=True, color=_C["white"])
    saldo_final = total_receita - total_despesa
    _money(ws, row, 5, saldo_final, fill_tot, bold=True, color=_C["verde"] if saldo_final >= 0 else _C["orange"])
    _money(ws, row, 6, saldo_final, fill_tot, bold=True, color=_C["verde"] if saldo_final >= 0 else _C["orange"])
    ws.cell(row=row, column=7).fill = fill_tot
    ws.cell(row=row, column=7).border = _border()
    row += 3
    _section_row(ws, row, "MÉTRICAS DO FLUXO DE CAIXA", NCOLS)
    row += 1
    M_val = canonical.resultado.venda.total_com_bdi if canonical.resultado else 0.0
    metricas = [
        ("Duração da execução (meses)", n_exec),
        ("Parcela mensal (meses 2 a n)", math.floor(M_val / n_exec) if n_exec and canonical.resultado else "—"),
        (
            "Resíduo (mês n+1)",
            M_val - math.floor(M_val / n_exec) * (n_exec - 1) if n_exec and canonical.resultado else "—",
        ),
        ("Exposição máxima", max_exposicao),
        ("Mês da exposição máxima", mes_max_exp),
        ("Saldo final do projeto", saldo_final),
    ]
    row = _xl_cashflow_metrics_block(ws, row, metricas, NCOLS)
    row += 1
    nota = (
        "Mês 1: receita = 0.  Meses 2..n: ⌊M/n⌋ cada.  Mês n+1: resíduo M − ⌊M/n⌋×(n−1).  "
        "Saldo Atual = Rec. Acumulada − Desp. Acumulada.  "
        "Exposição Acum. = −Saldo Atual (positivo = empresa precisa de capital de giro)."
    )
    cell_nota = ws.cell(row=row, column=1, value=nota)
    cell_nota.font = _font(italic=True, size=9, color="555555")
    cell_nota.alignment = _align("left", wrap=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
    ws.row_dimensions[row].height = 36


