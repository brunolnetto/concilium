"""Excel sheets — resultado (parte extraída de resultado.py, grupo _cronograma)."""

from __future__ import annotations

from openpyxl import Workbook

from concilium.domain.models.canonical import CanonicalModel

from ...helpers import (
    _fmt_mes,
    _project_months,
)
from ...styles import (
    _C,
    _align,
    _border,
    _fill,
    _freeze,
    _header_row,
    _label,
    _set_col_widths,
    _title_row,
)


def _rcr_build_qty_by_item(canonical, project_months: list) -> dict:
    """Build dict of {activity_id: {(year,month): qty}} from cronograma."""
    qty_by_item: dict = {}
    if not canonical.cronograma:
        return qty_by_item
    for act in canonical.cronograma.scheduled_activities:
        if act.activity_id not in qty_by_item:
            qty_by_item[act.activity_id] = {ym: 0.0 for ym in project_months}
        for wk, dist in act.distribuicao_semanal.items():
            from ...helpers import _week_to_month

            ym = _week_to_month(wk)
            if ym in qty_by_item[act.activity_id]:
                qty_by_item[act.activity_id][ym] += getattr(dist, "qtd_executada", 0.0)
    return qty_by_item


def _rcr_render_item_row(ws, row: int, item, fill, monthly: dict, project_months: list) -> int:
    """Render one EAP item row in CRONOGRAMA.RES sheet. Returns next row."""
    _label(ws, row, 1, item.item_id, fill)
    _label(ws, row, 2, item.descricao, fill, wrap=True)
    _label(ws, row, 3, item.unidade, fill, align="center")
    cell = ws.cell(row=row, column=4, value=item.quantitativo)
    cell.fill = fill
    cell.border = _border()
    cell.number_format = "#,##0.##"
    cell.alignment = _align("right")
    for j, ym in enumerate(project_months):
        v = monthly.get(ym, 0.0)
        cell = ws.cell(row=row, column=5 + j, value=round(v, 2) if v else None)
        cell.fill = fill
        cell.border = _border()
        if v:
            cell.number_format = "#,##0.##"
            cell.alignment = _align("right")
    return row + 1


def render_cronograma_res(wb: Workbook, canonical: CanonicalModel) -> None:
    """CRONOGRAMA.RES — Resumo mensal de produção física por item EAP."""
    ws = wb.create_sheet("CRONOGRAMA.RES")
    project_months = _project_months(canonical)
    if not project_months:
        ws.cell(row=1, column=1).value = "Sem dados de cronograma."
        return
    qty_by_item = _rcr_build_qty_by_item(canonical, project_months)
    NCOLS = 4 + len(project_months)
    _set_col_widths(ws, [10, 30, 6, 12] + [12] * len(project_months))
    _freeze(ws, "E3")
    row = 1
    _title_row(ws, row, "CRONOGRAMA RESUMIDO — PRODUÇÃO MENSAL POR EAP", NCOLS)
    row += 1
    _header_row(ws, row, ["Item ID", "Descrição", "Un.", "Total Qtd"] + [_fmt_mes(ym) for ym in project_months])
    row += 1
    ws.freeze_panes = f"E{row}"
    fill_colors = [_fill(_C["white"]), _fill(_C["row_alt"])]
    for i, item in enumerate(canonical.eap):
        fill = fill_colors[i % 2]
        monthly = qty_by_item.get(item.item_id, {ym: 0.0 for ym in project_months})
        row = _rcr_render_item_row(ws, row, item, fill, monthly, project_months)


