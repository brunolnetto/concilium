"""Excel sheets — resultado (parte extraída de resultado.py, grupo _diretos)."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from concilium.domain.models.canonical import (
    CanonicalModel,
    EapItem,
    Planejamento,
    SubtotalCategorias,
)

from ...styles import (
    _C,
    _align,
    _border,
    _fill,
    _font,
    _freeze,
    _header_row,
    _money,
    _set_col_widths,
    _title_row,
)


def _xl_dr_sem_vals(sem: SubtotalCategorias | None) -> tuple:
    return (
        sem.mao_obra if sem else 0.0,
        sem.equipamentos if sem else 0.0,
        sem.diesel if sem else 0.0,
        sem.material if sem else 0.0,
        sem.total if sem else 0.0,
    )


def _xl_dr_com_vals(com: SubtotalCategorias | None) -> tuple:
    return (
        com.mao_obra if com else 0.0,
        com.equipamentos if com else 0.0,
        com.diesel if com else 0.0,
        com.material if com else 0.0,
        com.total if com else 0.0,
    )


def _xl_dr_write_cell(ws: Worksheet, row: int, col: int, v: object, fill: PatternFill | None) -> None:
    cell = ws.cell(row=row, column=col, value=v)
    cell.fill = fill
    cell.border = _border()
    if col >= 5 and isinstance(v, float):
        cell.number_format = "#,##0.00"
        cell.alignment = _align("right")
    elif col == 4:
        cell.number_format = "#,##0.##"
        cell.alignment = _align("right")
    else:
        cell.alignment = _align("left")


def _xl_dr_totals_row(ws: Worksheet, row: int, data_start: int, NCOLS: int, fill_tot: PatternFill | None) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(row=row, column=1, value="TOTAL DIRETO")
    cell.fill = fill_tot
    cell.font = _font(bold=True, color=_C["white"], size=11)
    cell.alignment = _align("right")
    cell.border = _border()
    for c in range(5, NCOLS + 1):
        col_letter = get_column_letter(c)
        _money(
            ws, row, c, f"=SUM({col_letter}{data_start}:{col_letter}{row - 1})", fill_tot, bold=True, color=_C["white"]
        )


def _xl_dr_write_data_row(
    ws: Worksheet,
    row: int,
    p: Planejamento,
    eap_item: EapItem | None,
    sem: SubtotalCategorias | None,
    com: SubtotalCategorias | None,
    fill: PatternFill | None,
) -> None:
    vals: list = [
        p.item_id,
        eap_item.descricao if eap_item else p.item_id,
        eap_item.unidade if eap_item else "un",
        p.quantitativo,
        *_xl_dr_sem_vals(sem),
        *_xl_dr_com_vals(com),
    ]
    for col, v in enumerate(vals, 1):
        _xl_dr_write_cell(ws, row, col, v, fill)


def render_diretos_res(wb: Workbook, canonical: CanonicalModel) -> int | None:
    """DIRETOS.RES — Custos diretos resumidos por EAP (MO/EQ/COMB/MAT s/BDI e c/BDI)."""
    ws = wb.create_sheet("DIRETOS.RES")
    eap_ordem = {e.item_id: (e.ordem if e.ordem is not None else 9999) for e in canonical.eap}
    plan_sorted = sorted(canonical.planejamento, key=lambda p: eap_ordem.get(p.item_id, 9999))
    NCOLS = 14
    _set_col_widths(ws, [12, 42, 6, 13, 15, 15, 13, 13, 17, 15, 15, 13, 13, 17])
    _freeze(ws, "C3")
    row = 1
    _title_row(ws, row, "DIRETOS RESUMIDO — COMPONENTES MO/EQ/COMB/MAT s/BDI e c/BDI", NCOLS)
    row += 1
    _header_row(
        ws,
        row,
        [
            "Item ID",
            "Descrição",
            "Un.",
            "Quantitativo",
            "MO s/BDI",
            "EQ s/BDI",
            "Comb. s/BDI",
            "MAT s/BDI",
            "Total s/BDI",
            "MO c/BDI",
            "EQ c/BDI",
            "Comb. c/BDI",
            "MAT c/BDI",
            "Total c/BDI",
        ],
    )
    row += 1
    ws.freeze_panes = f"C{row}"
    data_start = row
    for i, p in enumerate(plan_sorted):
        eap_item = next((e for e in canonical.eap if e.item_id == p.item_id), None)
        sem = p.subtotal.sem_bdi if p.subtotal else None
        com = p.subtotal.com_bdi if p.subtotal else None
        fill = _fill(_C["white"]) if i % 2 == 0 else _fill(_C["row_alt"])
        _xl_dr_write_data_row(ws, row, p, eap_item, sem, com, fill)
        row += 1
    row += 1
    total_row = row
    _xl_dr_totals_row(ws, row, data_start, NCOLS, _fill(_C["nav_dk"]))
    wb.defined_names["DIRETOS_RES_SEM_BDI"] = DefinedName(
        "DIRETOS_RES_SEM_BDI", attr_text=f"'DIRETOS.RES'!${get_column_letter(9)}${total_row}"
    )
    wb.defined_names["DIRETOS_RES_COM_BDI"] = DefinedName(
        "DIRETOS_RES_COM_BDI", attr_text=f"'DIRETOS.RES'!${get_column_letter(14)}${total_row}"
    )
    return total_row


