"""Excel sheets — resultado (parte extraída de resultado.py, grupo _equilibrio)."""

from __future__ import annotations

from openpyxl import Workbook

from concilium.domain.models.canonical import CanonicalModel

from ...styles import (
    _C,
    _align,
    _border,
    _fill,
    _font,
    _freeze,
    _header_row,
    _label,
    _money,
    _section_row,
    _set_col_widths,
    _title_row,
)


def render_equilibrio(wb: Workbook, canonical: CanonicalModel) -> None:
    """EQUILIBRIO — Ponto de equilíbrio por EAP e balanço estrutural."""
    ws = wb.create_sheet("EQUILIBRIO")
    NCOLS = 6
    _set_col_widths(ws, [12, 42, 18, 18, 16, 12])
    _freeze(ws, "A3")
    row = 1
    _title_row(ws, row, "PONTO DE EQUILÍBRIO — CUSTO FIXO vs PREÇO DE VENDA (CPU)", NCOLS)
    row += 1
    _header_row(
        ws, row, ["Item ID", "CPU ID", "Custo Fixo Equipe (R$)", "Preço Venda CPU (R$)", "Delta (R$)", "Status"]
    )
    row += 1
    ws.freeze_panes = f"A{row}"
    _ST_CLR = {"OK": _C["verde"], "AVISO": _C["amarelo"], "ERRO": _C["vermelho"], "INFO": _C["nav_dk"]}
    _ST_BG = {"OK": _C["verde_lt"], "AVISO": _C["amar_lt"], "ERRO": _C["verm_lt"], "INFO": _C["row_alt"]}
    fill_colors = [_fill(_C["white"]), _fill(_C["row_alt"])]
    for i, pe in enumerate(canonical.ponto_equilibrio):
        st = pe.status_validacao
        fill = _fill(_ST_BG.get(st, _C["white"])) if st in ("ERRO", "AVISO") else fill_colors[i % 2]
        eap_item = next((e for e in canonical.eap if e.item_id == pe.item_id), None)
        _label(ws, row, 1, pe.item_id, fill)
        _label(ws, row, 2, eap_item.cpu_id if eap_item else pe.item_id, fill)
        _money(ws, row, 3, pe.custo_fixo_equipe, fill)
        _money(ws, row, 4, pe.custo_venda_cpu, fill)
        _money(ws, row, 5, pe.delta, fill, bold=True, color=_ST_CLR.get(st, _C["preto"]))
        _label(ws, row, 6, st, fill, bold=True, color=_ST_CLR.get(st, _C["preto"]), align="center")
        row += 1
    bpe = canonical.resultado.balanco_ponto_equilibrio if canonical.resultado else None
    if bpe:
        row += 2
        _section_row(ws, row, "BALANÇO DE PONTO DE EQUILÍBRIO", NCOLS)
        row += 1
        metricas = [
            ("HC Operacional (direto)", bpe.total_hc_operacional, False),
            ("HC Indireto", bpe.total_hc_indireto, False),
            ("HC Total", bpe.total_hc, False),
            ("Custo Fixo Mensal", bpe.custo_fixo_mensal, True),
            ("Custo Fixo Total", bpe.custo_fixo_total, True),
            ("Nº Equipes Diretas", bpe.numero_equipes_diretas, False),
        ]
        for label, val, is_money_val in metricas:
            fm = _fill(_C["cinza_lt"])
            _label(ws, row, 1, label, fm, bold=True)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            if is_money_val:
                _money(ws, row, 3, val, fm, bold=True)
            else:
                cell = ws.cell(row=row, column=3, value=val)
                cell.fill = fm
                cell.border = _border()
                cell.font = _font(bold=True)
                cell.alignment = _align("right")
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=NCOLS)
            row += 1


