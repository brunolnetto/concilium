"""Excel sheets — resultado (parte extraída de resultado.py, grupo _indiretos)."""

from __future__ import annotations

from openpyxl import Workbook

from concilium.domain.models.canonical import CanonicalModel

from ...styles import (
    _C,
    _align,
    _border,
    _fill,
    _font,
    _freeze,
    _header_row,
    _label,
    _money,
    _section_row,
    _set_col_widths,
    _title_row,
)


def _rir_build_lookups(canonical) -> tuple:
    """Build category lookup dicts and total_sem for render_indiretos_res."""
    inds_by_cat = {ind.categoria: ind for ind in canonical.indiretos}
    bdi_by_cat = {ibd.categoria: ibd for ibd in canonical.indiretos_com_bdi}
    total_sem = sum(ind.custo_total for ind in canonical.indiretos) or 1.0
    return inds_by_cat, bdi_by_cat, total_sem


def _rir_render_total_row(ws, row: int, subtotal_rows_sem: list, subtotal_rows_com: list, fill_tot) -> None:
    """Render totals row in INDIRETOS.RES."""
    f_sem = ("=" + "+".join(f"C{r}" for r in subtotal_rows_sem)) if subtotal_rows_sem else "0"
    f_com = ("=" + "+".join(f"E{r}" for r in subtotal_rows_com)) if subtotal_rows_com else "0"
    _money(ws, row, 3, f_sem, fill_tot, bold=True, color=_C["white"])
    ws.cell(row=row, column=4).fill = fill_tot
    ws.cell(row=row, column=4).border = _border()
    _money(ws, row, 5, f_com, fill_tot, bold=True, color=_C["white"])
    ws.cell(row=row, column=6).fill = fill_tot
    ws.cell(row=row, column=6).border = _border()


def _rir_render_cat_row(ws, row: int, cat: str, ind, ibd, fill, total_sem: float) -> int:
    """Render one indirect-cost category row. Returns next row."""
    custo_sem = ind.custo_total
    bdi_val = ibd.bdi_aplicado if ibd else 0.0
    custo_com = ibd.custo_com_bdi if ibd else custo_sem
    pct = custo_sem / total_sem * 100
    _label(ws, row, 1, cat, fill, bold=True, color=_C["nav_dk"])
    _label(ws, row, 2, ind.descricao, fill)
    _money(ws, row, 3, custo_sem, fill)
    cell = ws.cell(row=row, column=4, value=round(bdi_val * 100, 2))
    cell.fill = fill
    cell.border = _border()
    cell.number_format = '0.00"%"'
    cell.alignment = _align("right")
    _money(ws, row, 5, custo_com, fill)
    cell = ws.cell(row=row, column=6, value=round(pct, 2))
    cell.fill = fill
    cell.border = _border()
    cell.number_format = '0.00"%"'
    cell.alignment = _align("right")
    return row + 1


def _rir_render_kpis(ws, row: int, res, ncols: int) -> int:
    """Render KPI indicators section in INDIRETOS.RES. Returns next row."""
    cd = res.custo.direto_equipes
    ci = res.custo.total - cd
    pvt = res.venda.total_com_bdi
    kpis = [
        ("TOTAL DIRETO (catálogo s/BDI)", cd, False),
        ("TOTAL INDIRETO s/BDI", ci, False),
        ("RAZÃO CI / CD", ci / cd * 100 if cd > 0 else 0.0, True),
        ("PREÇO DE VENDA TOTAL", pvt, False),
        ("CI / PVT", ci / pvt * 100 if pvt > 0 else 0.0, True),
    ]
    for label, val, is_pct in kpis:
        fm = _fill(_C["cinza_lt"])
        _label(ws, row, 1, label, fm, bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        if is_pct:
            cell = ws.cell(row=row, column=3, value=round(val, 2))
            cell.fill = fm
            cell.border = _border()
            cell.number_format = '0.00"%"'
            cell.alignment = _align("right")
        else:
            _money(ws, row, 3, val, fm, bold=True)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=ncols)
        row += 1
    return row


def render_indiretos_res(wb: Workbook, canonical: CanonicalModel) -> None:
    """INDIRETOS.RES — Custos indiretos resumidos por categoria + KPIs de composição."""
    ws = wb.create_sheet("INDIRETOS.RES")
    NCOLS = 6
    _set_col_widths(ws, [14, 36, 18, 10, 18, 14])
    _freeze(ws, "A3")
    row = 1
    _title_row(ws, row, "INDIRETOS RESUMIDO — CATEGORIAS + INDICADORES", NCOLS)
    row += 1
    _header_row(ws, row, ["Categoria", "Descrição", "Custo s/BDI (R$)", "BDI (%)", "Custo c/BDI (R$)", "% do Total"])
    row += 1
    ws.freeze_panes = f"A{row}"
    ORDEM_CAT = ["ADM", "TOP", "LAB", "QSMS", "MOB"]
    inds_by_cat, bdi_by_cat, total_sem = _rir_build_lookups(canonical)
    fill_colors = [_fill(_C["white"]), _fill(_C["row_alt"])]
    subtotal_rows_sem: list[int] = []
    subtotal_rows_com: list[int] = []
    for i, cat in enumerate(ORDEM_CAT):
        ind = inds_by_cat.get(cat)  # type: ignore[call-overload]
        if not ind:
            continue
        ibd = bdi_by_cat.get(cat)  # type: ignore[call-overload]
        fill = fill_colors[i % 2]
        row = _rir_render_cat_row(ws, row, cat, ind, ibd, fill, total_sem)
        subtotal_rows_sem.append(row - 1)
        subtotal_rows_com.append(row - 1)
    row += 1
    fill_tot = _fill(_C["nav_dk"])
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell = ws.cell(row=row, column=1, value="TOTAL INDIRETO")
    cell.fill = fill_tot
    cell.font = _font(bold=True, color=_C["white"], size=11)
    cell.alignment = _align("right")
    cell.border = _border()
    _rir_render_total_row(ws, row, subtotal_rows_sem, subtotal_rows_com, fill_tot)
    row += 3
    _section_row(ws, row, "INDICADORES DE COMPOSIÇÃO DE CUSTO", NCOLS)
    row += 1
    res = canonical.resultado
    if res:
        row = _rir_render_kpis(ws, row, res, NCOLS)


