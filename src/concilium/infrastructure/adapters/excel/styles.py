from __future__ import annotations

from functools import cache
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_C = {
    "nav_dk": "1A3A5C",  # azul marinho escuro — cabeçalhos principais
    "nav_md": "2B6CB0",  # azul médio — subcabeçalhos
    "orange": "C05C18",  # laranja queimado — totais, destaques
    "orange_lt": "F5E6D8",  # laranja claro — linhas de total
    "row_alt": "F0F5FA",  # azul quase branco — linha alternada
    "white": "FFFFFF",
    "verde": "276749",
    "verde_lt": "DCEEE5",
    "vermelho": "9B1C1C",
    "verm_lt": "FADDDD",
    "amarelo": "92400E",
    "amar_lt": "FEF3C7",
    "cinza_bd": "CBD5E0",  # borda
    "cinza_lt": "F7FAFC",  # cabeçalho de seção leve
    "preto": "1A202C",
}

MESES_PT = {
    1: "Jan",
    2: "Fev",
    3: "Mar",
    4: "Abr",
    5: "Mai",
    6: "Jun",
    7: "Jul",
    8: "Ago",
    9: "Set",
    10: "Out",
    11: "Nov",
    12: "Dez",
}


@cache
def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


@cache
def _font(bold: bool = False, color: str = "1A202C", size: int = 10, italic: bool = False) -> Font:
    return Font(bold=bold, color=color, size=size, italic=italic, name="Calibri")


@cache
def _border(style: str = "thin") -> Border:
    s = Side(style=style, color=_C["cinza_bd"])
    return Border(left=s, right=s, top=s, bottom=s)


@cache
def _align(h: str = "left", v: str = "center", wrap: bool = False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _money(
    ws: Any,
    row: int,
    col: int,
    value: float | str,
    fill: PatternFill | None = None,
    bold: bool = False,
    color: str | None = None,
) -> None:
    is_formula = isinstance(value, str) and value.startswith("=")
    cell_value = value if is_formula else round(float(value), 2)
    cell = ws.cell(row=row, column=col, value=cell_value)
    cell.number_format = "R$ #,##0.00"
    cell.alignment = _align("right")
    cell.border = _border()
    if fill:
        cell.fill = fill
    cell.font = _font(bold=bold, color=color or _C["preto"])


def _label(
    ws: Any,
    row: int,
    col: int,
    value: Any,
    fill: PatternFill | None = None,
    bold: bool = False,
    color: str | None = None,
    align: str = "left",
    wrap: bool = False,
) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.alignment = _align(align, wrap=wrap)
    cell.border = _border()
    if fill:
        cell.fill = fill
    cell.font = _font(bold=bold, color=color or _C["preto"])


def _header_row(ws: Any, row: int, labels: list[str], dark: bool = True) -> None:
    """Linha de cabeçalho colorida."""
    bg = _C["nav_dk"] if dark else _C["nav_md"]
    for col, text in enumerate(labels, 1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill = _fill(bg)
        cell.font = _font(bold=True, color=_C["white"], size=10)
        cell.alignment = _align("center", wrap=True)
        cell.border = _border()


def _title_row(ws: Any, row: int, title: str, n_cols: int, color: str = "nav_dk") -> None:
    """Linha de título mesclada."""
    ws.cell(row=row, column=1, value=title).value = title
    if n_cols > 1:
        ws.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=n_cols,
        )
    cell = ws.cell(row=row, column=1)
    cell.fill = _fill(_C[color])
    cell.font = _font(bold=True, color=_C["white"], size=12)
    cell.alignment = _align("center")
    cell.border = _border()


def _section_row(ws: Any, row: int, title: str, n_cols: int) -> None:
    """Sub-cabeçalho de seção."""
    if n_cols > 1:
        ws.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=n_cols,
        )
    cell = ws.cell(row=row, column=1)
    cell.value = title
    cell.fill = _fill(_C["cinza_lt"])
    cell.font = _font(bold=True, color=_C["nav_dk"], size=10)
    cell.alignment = _align("left")
    cell.border = _border()


def _set_col_widths(ws: Any, widths: list[int | float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _auto_fit_columns(ws: Any, min_width: float = 8.0, max_width: float = 42.0, padding: float = 2.0) -> None:
    """Ajusta largura das colunas com base no maior valor textual visível."""
    for column_cells in ws.columns:
        max_len = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is None:
                continue
            value = cell.value
            if isinstance(value, (int, float)):
                text = f"{value:,.2f}"
            else:
                text = str(value)
            max_len = max(max_len, len(text))
        if max_len > 0:
            ws.column_dimensions[column_letter].width = min(max_width, max(min_width, max_len + padding))


def _freeze(ws: Any, cell: str) -> None:
    ws.freeze_panes = cell
