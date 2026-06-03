from pathlib import Path

import pandas as pd

from concilium.infrastructure.exporters.catalog_excel import CatalogExporterExcel


class PriceSurveyExporter(CatalogExporterExcel):
    """Gera planilhas formatadas para pesquisa de mercado e atualização de preços."""

    def export(self, output_path: Path):
        insumos_raw = self.catalogo.get("insumos", [])

        # Categoriza para abas diferentes
        data_map = {"MATERIAL": [], "MAO_OBRA": [], "EQUIPAMENTO": []}  # type: ignore[var-annotated]

        for i in insumos_raw:
            tipo = i.get("tipo")
            if tipo not in data_map:
                continue

            row = {
                "ID": i.get("codigo"),
                "Categoria": i.get("categoria"),
                "Descrição": i.get("descricao"),
                "Unidade": i.get("unidade"),
                "Preço Atual (Benchmark)": self._get_preco_regional(i),
                "Novo Preço (Cotação)": "",  # Campo para preenchimento
                "Fornecedor": "",  # Campo para preenchimento
                "Data Cotação": "",  # Campo para preenchimento
                "Observações": "",  # Campo para preenchimento
            }
            data_map[tipo].append(row)

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for tipo, items in data_map.items():
                if not items:
                    continue
                df = pd.DataFrame(items)
                sheet_name = tipo.replace("_", " ").title()
                df.to_excel(writer, index=False, sheet_name=sheet_name)

            self._aplicar_estilo_survey(writer)

    def _aplicar_estilo_survey(self, writer):
        # Reutiliza a lógica de estilo base, mas com toques específicos para formulário
        self._aplicar_estilo(writer)

        from openpyxl.styles import Font, PatternFill

        # Cor de destaque para campos que DEVEM ser preenchidos (Amarelo Mendes Lobo)
        idv = self.catalogo.get("identidade_visual", {})
        highlight_color = idv.get("cores", {}).get("destaque", "FFC000")
        fill_input = PatternFill(start_color=highlight_color, end_color=highlight_color, fill_type="solid")
        font_bold = Font(bold=True)

        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]
            # Colunas F, G, H, I são os campos de entrada (Novo Preço, Fornecedor, etc)
            # Vamos destacar o cabeçalho desses campos
            for col_letter in ["F", "G", "H", "I"]:
                cell = ws[f"{col_letter}1"]
                cell.fill = fill_input
                cell.font = font_bold
