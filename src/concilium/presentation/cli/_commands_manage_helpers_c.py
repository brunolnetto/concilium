"""Helper functions for commands_manage (C)."""

# ruff: noqa: F401
from __future__ import annotations

import json
from datetime import UTC, date, datetime
from pathlib import Path

import typer
from rich import print as rprint
from rich.console import Console
from rich.table import Table

from concilium.presentation.cli.shared import resolve_project


def _cmr_load_bundle(contexto_id: str, mod: object) -> tuple[dict, object, list, list]:
    """Loads catalog bundle. Returns (bundle, bundle_path, cat_insumos, cat_producoes)."""

    bundle_path = mod.ROOT / f"config/catalogos/{contexto_id}/.catalog_bundle.json"
    if not bundle_path.exists():
        rprint(f"[red]Bundle não encontrado:[/red] {bundle_path}")
        raise typer.Exit(1)
    bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
    cat_insumos = bundle.get("insumos", [])
    cat_producoes = bundle.get("producoes", []) or bundle.get("produtividades", [])
    return bundle, bundle_path, cat_insumos, cat_producoes


def _cmr_print_inner_low(all_low: list, console: object) -> None:
    if not all_low:
        return
    console.print(f"\n[yellow]INNER_LOW — {len(all_low)} colisões (importados como RFI-*):[/yellow]")
    for item, etype in all_low[:15]:
        cat_e = item["catalog"]
        rfi_e = item["rfi"]
        console.print(
            f"  [{item['score']:.2f}] {etype} "
            f'cat={cat_e.get("codigo", "")[:20]} "{cat_e.get("descricao", "")[:25]}" '
            f'| rfi={rfi_e.get("codigo", "")[:20]} "{rfi_e.get("descricao", "")[:25]}"'
        )
    if len(all_low) > 15:
        console.print(f"  ... e mais {len(all_low) - 15}")


def _cmr_apply_changes(
    mod: object,
    bundle: dict,
    bundle_path: object,
    cat_insumos: list,
    cat_producoes: list,
    rfi_insumos: list,
    rfi_producoes: list,
    report_ins: object,
    report_prod: object,
    verify_pu: bool,
    console: object,
) -> None:
    new_ins = mod._apply_merge(cat_insumos, rfi_insumos, report_ins)
    new_prod = mod._apply_merge(cat_producoes, rfi_producoes, report_prod, id_key="id")
    if "producoes" in bundle:
        bundle["producoes"] = new_prod
    else:
        bundle["produtividades"] = new_prod
    bundle["insumos"] = new_ins

    bundle_path.write_text(json.dumps(bundle, indent=2, ensure_ascii=False), encoding="utf-8")
    console.print(f"\n[green]Bundle atualizado:[/green] {bundle_path}")
    console.print(f"  Insumos:   {len(new_ins)} (era {len(cat_insumos)})")
    console.print(f"  Producoes: {len(new_prod)} (era {len(cat_producoes)})")
    if verify_pu:
        from concilium.application.catalog.pu_calculator import ParametrosCalculo, calcular_catalogo

        resultado = calcular_catalogo(new_prod, new_ins, ParametrosCalculo())
        faltantes = resultado["referencias_faltantes"]
        if faltantes:
            console.print(f"\n[red]Referências faltantes após merge:[/red] {len(faltantes)}")
            for r in faltantes[:10]:
                console.print(f"  {r['producao_id']}: {r['faltantes']}")
        else:
            console.print(
                f"\n[green]Integridade referencial OK[/green] — {resultado['com_composicao']} produções verificadas"
            )


def _cmr_write_report(
    report_json: str,
    report_ins: object,
    report_prod: object,
    mode: str,
    contexto_id: str,
    console: object,
) -> None:
    from pathlib import Path

    report_data = {
        "mode": mode,
        "contexto_id": contexto_id,
        "insumos": {
            "inner_high": len(report_ins.inner_high),
            "inner_high_com_delta": sum(1 for i in report_ins.inner_high if i["delta"]),
            "inner_low": len(report_ins.inner_low),
            "right_only": len(report_ins.right_only),
            "left_only": len(report_ins.left_only),
            "deltas": [
                {"codigo": i["catalog"].get("codigo"), "delta": i["delta"]} for i in report_ins.inner_high if i["delta"]
            ][:50],
            "colisoes": [
                {"cat": i["catalog"].get("codigo"), "rfi": i["rfi"].get("codigo"), "score": i["score"]}
                for i in report_ins.inner_low
            ],
            "novos": [
                {"codigo": i["rfi"].get("codigo"), "descricao": i["rfi"].get("descricao")}
                for i in report_ins.right_only
            ],
        },
        "producoes": {
            "inner_high": len(report_prod.inner_high),
            "inner_low": len(report_prod.inner_low),
            "right_only": len(report_prod.right_only),
            "left_only": len(report_prod.left_only),
            "colisoes": [
                {"cat": i["catalog"].get("id"), "rfi": i["rfi"].get("id"), "score": i["score"]}
                for i in report_prod.inner_low
            ],
            "novos": [
                {"id": i["rfi"].get("id"), "descricao": i["rfi"].get("descricao")} for i in report_prod.right_only
            ],
        },
    }
    Path(report_json).write_text(json.dumps(report_data, indent=2, ensure_ascii=False), encoding="utf-8")
    console.print(f"\n[dim]Relatório JSON salvo em: {report_json}[/dim]")


def _cmr_add_table_row(table, label: str, rep) -> None:
    n_delta = 0
    for i in rep.inner_high:
        if i["delta"]:
            n_delta += 1
    table.add_row(
        label,
        str(len(rep.inner_high)),
        str(n_delta),
        str(len(rep.inner_low)),
        str(len(rep.right_only)),
        str(len(rep.left_only)),
    )


def _cmr_build_all_low(report_ins, report_prod) -> list:
    result = []
    for i in report_ins.inner_low:
        result.append((i, "INSUMO"))
    for i in report_prod.inner_low:
        result.append((i, "PROD"))
    return result


def _cmr_load_openpyxl(mod):
    import typer
    from rich import print as rprint

    try:
        import openpyxl
    except ImportError:
        rprint("[red]openpyxl não encontrado.[/red] Execute: uv add openpyxl")
        raise typer.Exit(2) from None
    xlsm = mod.XLSM
    if not xlsm.exists():
        rprint(f"[red]Planilha RFI não encontrada:[/red] {xlsm}")
        raise typer.Exit(2)
    return openpyxl, xlsm


def _impl_catalog_merge_rfi(
    contexto_id: str = typer.Argument(..., help="ID do contexto do catálogo."),
    apply: bool = typer.Option(False, "--apply", help="Aplica as mudanças no bundle (default: dry-run)."),
    verify_pu: bool = typer.Option(False, "--verify-pu", help="Verifica integridade referencial após merge."),
    report_json: str = typer.Option(None, "--report-json", help="Salva relatório detalhado em JSON."),
):
    """Compara e merge catálogo com PLANILHA PADRÃO RFI.

    Operações (como JOIN):
      LEFT_ONLY  : só no catálogo                   -> mantém
      INNER_HIGH : match confiante (>=0.85)          -> atualiza preço/dados
      INNER_LOW  : match incerto (0.50-0.85)         -> importa como RFI-* (sem sobrescrever)
      RIGHT_ONLY : só na planilha                    -> adiciona ao catálogo
    """
    import importlib.util

    console = Console()
    tools_dir = Path(__file__).resolve().parents[4].parent / "tools"
    merge_script = tools_dir / "merge_rfi.py"

    if not merge_script.exists():
        rprint(f"[red]Script não encontrado: {merge_script}[/red]")
        raise typer.Exit(2)

    spec = importlib.util.spec_from_file_location("merge_rfi", merge_script)
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)  # type: ignore[union-attr]

    openpyxl, xlsm = _cmr_load_openpyxl(mod)

    bundle, bundle_path, cat_insumos, cat_producoes = _cmr_load_bundle(contexto_id, mod)

    mode = "APPLY" if apply else "DRY-RUN"
    console.print(f"\n[bold]Merge Catálogo <-> RFI [{mode}][/bold] — contexto: {contexto_id}")
    console.print(f"  Catálogo: {len(cat_insumos)} insumos, {len(cat_producoes)} producoes")

    wb = openpyxl.load_workbook(xlsm, data_only=True, keep_vba=False)
    rfi = mod.extract_rfi(wb)
    wb.close()
    rfi_insumos = rfi["insumos"]
    rfi_producoes = rfi["producoes"]
    console.print(f"  Planilha: {len(rfi_insumos)} insumos, {len(rfi_producoes)} producoes")

    report_ins, _ = mod._merge_entities(cat_insumos, rfi_insumos)
    report_prod, _ = mod._merge_entities(cat_producoes, rfi_producoes, id_key="id")

    table = Table(title="Merge Report", show_header=True, header_style="bold cyan")
    table.add_column("Entidade")
    table.add_column("INNER_HIGH", justify="right")
    table.add_column("INNER_HIGH Delta", justify="right")
    table.add_column("INNER_LOW", justify="right")
    table.add_column("RIGHT_ONLY", justify="right")
    table.add_column("LEFT_ONLY", justify="right")
    for label, rep in [("INSUMOS", report_ins), ("PRODUCOES", report_prod)]:
        _cmr_add_table_row(table, label, rep)
    console.print(table)

    all_low = _cmr_build_all_low(report_ins, report_prod)
    _cmr_print_inner_low(all_low, console)

    if apply:
        _cmr_apply_changes(
            mod,
            bundle,
            bundle_path,
            cat_insumos,
            cat_producoes,
            rfi_insumos,
            rfi_producoes,
            report_ins,
            report_prod,
            verify_pu,
            console,
        )
    if report_json:
        _cmr_write_report(report_json, report_ins, report_prod, mode, contexto_id, console)
    if not apply:
        console.print("\n[dim]Dry-run — use --apply para aplicar o merge[/dim]")


def _ccp_write_bundle(bundle: dict, bundle_path, insumos: list, producoes: list, console) -> None:

    from concilium.application.catalog.plausibility_checker import verificar_plausibilidade

    if "producoes" in bundle:
        bundle["producoes"] = producoes
    else:
        bundle["produtividades"] = producoes
    bundle["insumos"] = insumos
    bundle_path.write_text(json.dumps(bundle, indent=2, ensure_ascii=False), encoding="utf-8")
    console.print(f"\n[green]Bundle atualizado:[/green] {len(insumos)} insumos, {len(producoes)} producoes")
    pos = verificar_plausibilidade(insumos, producoes)
    r = pos["resumo"]
    status = "[green]OK[/green]" if r["erros"] == 0 else "[yellow]ATENÇÃO[/yellow]"
    console.print(f"  Pós-curação: {status} — {r['erros']} erros | {r['avisos']} avisos | {r['infos']} infos")


def _ccp_print_fixed(fixed: list, apply: bool, console) -> None:
    console.print("\n[green]Correções:[/green]")
    for msg in fixed[:30]:
        prefix = "✓" if apply else "→"
        console.print(f"  {prefix} {msg}")
    if len(fixed) > 30:
        console.print(f"  ... e mais {len(fixed) - 30}")


def _impl_catalog_cure_plausibility(
    contexto_id: str = typer.Option(..., "--contexto", "-c", help="ID do contexto"),
    apply: bool = typer.Option(False, "--apply", help="Aplica correções (padrão: dry-run)"),
):
    """Curação cirúrgica de anomalias de plausibilidade do catálogo.

    Detecta e corrige (quando possível automaticamente):
      - Encargos armazenados como alpha em vez de alpha-1
      - Preço abaixo do salário mínimo com referência RFI disponível
      - Duplicatas P#### vs MO-NOME (remove duplicata, mantém código canônico)
      - Datas desatualizadas após merge RFI
      - Equipamento sem valor_locacao_mensal (deriva automaticamente)

    Itens não-automatizáveis são reportados para revisão manual.
    """
    from concilium.application.catalog.cure_plausibility import ROOT as CURE_ROOT
    from concilium.application.catalog.cure_plausibility import (
        CureReport,
        cure_duplicatas,
        cure_erros_fisicos,
        cure_locacao_mensal,
        cure_stale_rfi,
        report_manual,
    )

    console = Console()

    bundle_path = CURE_ROOT / f"config/catalogos/{contexto_id}/.catalog_bundle.json"
    if not bundle_path.exists():
        rprint(f"[red]Bundle não encontrado: {bundle_path}[/red]")
        raise typer.Exit(1)

    bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
    insumos = bundle.get("insumos", [])
    producoes = bundle.get("producoes", []) or bundle.get("produtividades", [])

    mode = "APPLY" if apply else "DRY-RUN"
    console.print(f"\n[bold]Curação de Plausibilidade [{mode}][/bold] — contexto: {contexto_id}")

    report = CureReport()
    insumos = cure_erros_fisicos(insumos, report, apply)
    insumos, producoes = cure_duplicatas(insumos, producoes, report, apply)
    insumos = cure_stale_rfi(insumos, report, apply)
    insumos = cure_locacao_mensal(insumos, report, apply)
    report_manual(insumos, producoes, report)

    table = Table(show_header=True, header_style="bold cyan")
    table.add_column("Categoria")
    table.add_column("Quantidade", justify="right")
    table.add_row("Correções automáticas", str(len(report.fixed)))
    table.add_row("Intervenção manual", str(len(report.manual)))
    console.print(table)

    if report.fixed:
        _ccp_print_fixed(report.fixed, apply, console)

    if report.manual:
        console.print("\n[yellow]Requer revisão manual:[/yellow]")
        for msg in report.manual:
            console.print(f"  ⚠ {msg}")

    if apply:
        _ccp_write_bundle(bundle, bundle_path, insumos, producoes, console)
    else:
        console.print("\n[dim]Dry-run — use --apply para aplicar[/dim]")


def _impl_structure_list(projeto: str):
    """Lista a árvore de escopos e planos do projeto."""
    project = resolve_project(projeto)
    console = Console()

    table = Table(title=f"Estrutura: {projeto}")
    table.add_column("Escopo", style="cyan")
    table.add_column("Planos (Status)", style="yellow")

    for eid in project.list_escopos():
        es = project.escopo(eid)
        plans_str = "\n".join([f"- {p} ({es.plan(p).status})" for p in es.list_plans()])
        table.add_row(eid, plans_str)

    console.print(table)


def _impl_plan_init(
    projeto: str,
    nome: str = typer.Option("Plano Principal", "--nome", "-n"),
    escopo: str = typer.Option("default", "--escopo", "-e"),
):
    """Cria novo plano de trabalho (shortcut)."""
    project = resolve_project(projeto)
    es = project.escopo(escopo)
    ps = es.create_plan(nome=nome)
    rprint(f"[green][OK][/green] Plano [bold]{ps.plano_id}[/bold] criado no escopo [bold]{escopo}[/bold].")
