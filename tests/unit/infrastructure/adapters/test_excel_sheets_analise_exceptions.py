"""Cobre os ramos de exceção silenciosa em render_curva_abc/render_s_curve (ADR-058)."""

from __future__ import annotations

from openpyxl import Workbook

from concilium.infrastructure.adapters.excel.sheets.analise import render_curva_abc, render_s_curve


class _Empty:
    """Objeto sem qualquer atributo esperado pelos builders (dispara AttributeError)."""


def test_render_curva_abc_swallows_builder_exception():
    wb = Workbook()
    canonical = _Empty()

    # canonical.curva_abc_insumo é None (getattr default) e CurvaABCBuilder().build()
    # lança AttributeError ao acessar canonical_in.equipes_diretas -> retorno silencioso.
    render_curva_abc(wb, canonical)

    assert "ANALISE.CURVA_ABC" not in wb.sheetnames


def test_render_s_curve_swallows_builder_exception():
    wb = Workbook()
    canonical = _Empty()

    # canonical.s_curve é None e SCurveBuilder().build() lança AttributeError ao
    # acessar canonical_in.cronograma -> retorno silencioso.
    render_s_curve(wb, canonical)

    assert "ANALISE.S_CURVE" not in wb.sheetnames
