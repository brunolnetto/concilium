"""Coverage sweep 11 — files below 90%:

- _indiretos/_skill_helpers_a.py (88%) — missing: 38, 122-124, 132, 137, 170, 202-203, 216-217, 242->exit
- mixins/catalog.py (88%) — missing: 49->53, 52, 85-110, 123-124, 167->162, 215-216, 230-231, 237-238,
                                       248->251, 255->251, 259->255, 264-265, 280
- recursos.py (89%) — missing: 77, 271-295
"""

from __future__ import annotations

import importlib.util
from pathlib import Path
from unittest.mock import MagicMock

_ROOT = Path(__file__).resolve().parents[4]

def _load_ind_helpers_a():
    path = (
        _ROOT
        / "src/concilium/application/skills/manifests"
        / "06-resultado-financeiro/_indiretos/_skill_helpers/_dimensionamento.py"
    )
    spec = importlib.util.spec_from_file_location("ind_helpers_a_sw11", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ── recursos.py (excel sheets) ────────────────────────────────────────────────


class TestRenderCalendarioOperacional:
    def test_item_not_dict_nor_model_uses_empty_row_data(self):
        """Line 77: item is neither model_dump-able nor dict → row_data = {}."""
        from openpyxl import Workbook

        from concilium.infrastructure.adapters.excel.sheets.recursos import render_calendario_operacional

        wb = Workbook()
        canonical = MagicMock()
        # a plain string is neither a model nor a dict
        canonical.calendario_operacional = ["plain_string"]
        canonical.cronograma = None

        render_calendario_operacional(wb, canonical)
        ws = wb["CALENDARIO.OP"]
        # semana = str({}.get("semana_iso", "")) == ""
        assert ws.cell(row=5, column=1).value == ""


class TestRenderEquipeIndireto:
    def _make_entry(self, with_equipamentos=True):
        from types import SimpleNamespace

        funcoes = [{"funcao_id": "ENG", "num_pessoas": 2, "custo_mensal_unitario": 5000.0, "custo_total": 10000.0}]
        equipamentos = (
            [{"equip_id": "EQ-1", "num_unidades": 1, "custo_mensal": 2000.0, "custo_total": 2000.0}]
            if with_equipamentos
            else []
        )
        return SimpleNamespace(
            categoria="ADM",
            custo_total=12000.0,
            custo_mao_obra=10000.0,
            custo_equipamentos=2000.0,
            custo_combustivel=0.0,
            funcoes=funcoes,
            equipamentos=equipamentos,
        )

    def _make_canonical(self, entry):
        canonical = MagicMock()
        canonical.indiretos = [entry]
        return canonical

    def test_equipamentos_loop_renders_rows(self):
        """Lines 271-295: equipamentos loop in render_equipe_indireto."""
        from openpyxl import Workbook

        from concilium.infrastructure.adapters.excel.sheets.recursos import render_equipe_indireto

        wb = Workbook()
        canonical = self._make_canonical(self._make_entry(with_equipamentos=True))
        render_equipe_indireto(wb, canonical, "ADM")
        ws = wb["EQUIPE.INDIRETO.ADM"]
        all_vals = [ws.cell(row=r, column=1).value for r in range(1, 20)]
        assert "EQUIPAMENTO" in all_vals

    def test_no_equipamentos_renders_only_funcoes(self):
        """Equipamentos loop skipped when empty; has_rows=True from funcoes."""
        from openpyxl import Workbook

        from concilium.infrastructure.adapters.excel.sheets.recursos import render_equipe_indireto

        wb = Workbook()
        canonical = self._make_canonical(self._make_entry(with_equipamentos=False))
        render_equipe_indireto(wb, canonical, "ADM")
        ws = wb["EQUIPE.INDIRETO.ADM"]
        all_vals = [ws.cell(row=r, column=1).value for r in range(1, 20)]
        assert "MÃO DE OBRA" in all_vals
        assert "EQUIPAMENTO" not in all_vals

    def test_entry_not_found_shows_message(self):
        """Lines 218-223: categoria not found → message displayed."""
        from openpyxl import Workbook

        from concilium.infrastructure.adapters.excel.sheets.recursos import render_equipe_indireto

        wb = Workbook()
        canonical = MagicMock()
        canonical.indiretos = []
        render_equipe_indireto(wb, canonical, "LAB")
        ws = wb["EQUIPE.INDIRETO.LAB"]
        assert any("não encontrada" in str(ws.cell(row=r, column=1).value or "") for r in range(1, 5))

    def test_no_funcoes_no_equipamentos_shows_fallback(self):
        """Lines 297-303: has_rows=False → fallback cell shown."""
        from types import SimpleNamespace

        from openpyxl import Workbook

        from concilium.infrastructure.adapters.excel.sheets.recursos import render_equipe_indireto

        wb = Workbook()
        entry = SimpleNamespace(
            categoria="TOP",
            custo_total=5000.0,
            custo_mao_obra=0.0,
            custo_equipamentos=0.0,
            custo_combustivel=0.0,
            funcoes=[],
            equipamentos=[],
        )
        canonical = MagicMock()
        canonical.indiretos = [entry]
        render_equipe_indireto(wb, canonical, "TOP")
        ws = wb["EQUIPE.INDIRETO.TOP"]
        all_vals = [str(ws.cell(row=r, column=1).value or "") for r in range(1, 20)]
        assert any("Sem equipes" in v for v in all_vals)

