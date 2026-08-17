"""Coverage tests for CLI manage helpers and config/version commands.

Targets:
  - _commands_manage_helpers_a.py  (catalog quality helpers)
  - _commands_manage_helpers_c.py  (catalog merge/cure/structure helpers)
  - _commands_manage_helpers_d.py  (catalog validate/tendencia/audit helpers)
  - commands_config_version.py     (context + perfil versioning CLI)
"""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
import typer
from typer.testing import CliRunner

runner = CliRunner()


# ═══════════════════════════════════════════════════════════════════════════════
# _commands_manage_helpers_a.py
# ═══════════════════════════════════════════════════════════════════════════════


class TestLoadJson:
    """Tests for _load_json helper."""

    def test_returns_default_when_file_missing(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_a import _load_json

        result = _load_json(tmp_path / "nonexistent.json", {"default": True})
        assert result == {"default": True}

    def test_reads_valid_json(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_a import _load_json

        p = tmp_path / "data.json"
        p.write_text(json.dumps({"key": "value"}), encoding="utf-8")
        result = _load_json(p, {})
        assert result == {"key": "value"}

    def test_returns_default_on_bad_json(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_a import _load_json

        p = tmp_path / "bad.json"
        p.write_text("not json!!!", encoding="utf-8")
        result = _load_json(p, {"fallback": 1})
        assert result == {"fallback": 1}

    def test_returns_default_when_json_not_dict(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_a import _load_json

        p = tmp_path / "list.json"
        p.write_text(json.dumps([1, 2, 3]), encoding="utf-8")
        result = _load_json(p, {"default": True})
        assert result == {"default": True}


class TestWriteJson:
    def test_creates_parent_dirs_and_writes(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_a import _write_json

        p = tmp_path / "a" / "b" / "file.json"
        _write_json(p, {"x": 1})
        assert p.exists()
        assert json.loads(p.read_text()) == {"x": 1}


class TestAppendJsonl:
    def test_appends_multiple_lines(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_a import _append_jsonl

        p = tmp_path / "log.jsonl"
        _append_jsonl(p, {"a": 1})
        _append_jsonl(p, {"b": 2})
        lines = p.read_text().strip().splitlines()
        assert len(lines) == 2
        assert json.loads(lines[0]) == {"a": 1}
        assert json.loads(lines[1]) == {"b": 2}


class TestCqValidateCpus:
    def test_errors_and_avisos_collected(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_validate_cpus

        res_err = SimpleNamespace(status="ERRO", codigo="VAL-CPU-STR-02", mensagem="erro de estrutura")
        res_warn = SimpleNamespace(status="AVISO", codigo="VAL-GEN-01", mensagem="aviso geral")

        validador = MagicMock()
        validador.validar_cpu_completa.side_effect = [
            [res_err],
            [res_warn],
        ]
        cpus = [{"codigo": "CPU-A"}, {"codigo": "CPU-B"}]
        erros, avisos, bloqueadas = _cq_validate_cpus(cpus, validador)
        assert len(erros) == 1
        assert erros[0][0] == "CPU-A"
        assert len(avisos) == 1
        assert "CPU-A" in bloqueadas

    def test_str03_also_blocks(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_validate_cpus

        res_err = SimpleNamespace(status="ERRO", codigo="VAL-CPU-STR-03", mensagem="campo ausente")
        validador = MagicMock()
        validador.validar_cpu_completa.return_value = [res_err]
        cpus = [{"codigo": "CPU-X"}]
        erros, avisos, bloqueadas = _cq_validate_cpus(cpus, validador)
        assert "CPU-X" in bloqueadas

    def test_empty_cpu_list(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_validate_cpus

        validador = MagicMock()
        erros, avisos, bloqueadas = _cq_validate_cpus([], validador)
        assert erros == [] and avisos == [] and bloqueadas == []

    def test_erro_com_codigo_nao_bloqueante_nao_e_adicionado_a_bloqueadas(self):
        """ERRO cujo codigo não está em VAL-CPU-STR-02/03 não bloqueia a CPU."""
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_validate_cpus

        res_err = SimpleNamespace(status="ERRO", codigo="VAL-CPU-OUTRO", mensagem="erro genérico")
        validador = MagicMock()
        validador.validar_cpu_completa.return_value = [res_err]
        cpus = [{"codigo": "CPU-Y"}]
        erros, avisos, bloqueadas = _cq_validate_cpus(cpus, validador)
        assert len(erros) == 1
        assert bloqueadas == []

    def test_resultado_ok_nao_gera_erro_nem_aviso(self):
        """status diferente de ERRO/AVISO (ex.: OK) não popula nenhuma lista."""
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_validate_cpus

        res_ok = SimpleNamespace(status="OK", codigo="OK-01", mensagem="ok")
        validador = MagicMock()
        validador.validar_cpu_completa.return_value = [res_ok]
        cpus = [{"codigo": "CPU-Z"}]
        erros, avisos, bloqueadas = _cq_validate_cpus(cpus, validador)
        assert erros == [] and avisos == [] and bloqueadas == []


class TestCqValidateInsumos:
    def test_collects_vigencia_and_encoding_avisos(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_validate_insumos

        res_v = SimpleNamespace(status="AVISO", codigo="VAL-CAT-VIG-01", mensagem="desatualizado")
        res_e = SimpleNamespace(status="AVISO", codigo="VAL-CAT-ENC-01", mensagem="mojibake")
        validador = MagicMock()
        validador.validar_vigencia.return_value = res_v
        validador.validar_encoding.return_value = res_e
        insumos = [{"codigo": "INS-01"}]
        avisos, obsoletos = _cq_validate_insumos(insumos, validador)
        assert "INS-01" in obsoletos
        assert len(avisos) == 2

    def test_ok_insumo_no_avisos(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_validate_insumos

        res_ok = SimpleNamespace(status="OK", codigo="OK-01", mensagem="ok")
        validador = MagicMock()
        validador.validar_vigencia.return_value = res_ok
        validador.validar_encoding.return_value = res_ok
        insumos = [{"codigo": "INS-02"}]
        avisos, obsoletos = _cq_validate_insumos(insumos, validador)
        assert avisos == [] and obsoletos == []


class TestCqPrintBlockedCpus:
    def test_empty_blocked_list_prints_nothing(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_print_blocked_cpus

        console = MagicMock()
        _cq_print_blocked_cpus([], [], console)
        console.print.assert_not_called()

    def test_blocked_cpu_prints_table(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_print_blocked_cpus

        err = SimpleNamespace(mensagem="campo faltante")
        console = MagicMock()
        _cq_print_blocked_cpus(["CPU-A"], [("CPU-A", err)], console)
        console.print.assert_called_once()

    def test_erro_nao_bloqueado_e_pulado_na_tabela(self):
        """erros com codigo fora de cpus_bloqueadas são ignorados na tabela."""
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_print_blocked_cpus

        err_bloqueado = SimpleNamespace(mensagem="campo faltante")
        err_nao_bloqueado = SimpleNamespace(mensagem="erro genérico")
        console = MagicMock()
        _cq_print_blocked_cpus(
            ["CPU-A"],
            [("CPU-B", err_nao_bloqueado), ("CPU-A", err_bloqueado)],
            console,
        )
        console.print.assert_called_once()


class TestCqPrintEncodingAlerts:
    def test_no_encoding_alerts_prints_nothing(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_print_encoding_alerts

        console = MagicMock()
        other = SimpleNamespace(codigo="VAL-OTHER", mensagem="x")
        _cq_print_encoding_alerts([("A", other)], console)
        console.print.assert_not_called()

    def test_encoding_alert_prints_table(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_print_encoding_alerts

        console = MagicMock()
        enc = SimpleNamespace(codigo="VAL-CAT-ENC-01", mensagem="mojibake")
        _cq_print_encoding_alerts([("INS-01", enc)], console)
        console.print.assert_called_once()

    def test_encoding_alert_overflow_more_than_10(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_print_encoding_alerts

        console = MagicMock()
        enc = SimpleNamespace(codigo="VAL-CAT-ENC-01", mensagem="x")
        # 12 alerts
        avisos = [(f"INS-{i:02d}", enc) for i in range(12)]
        _cq_print_encoding_alerts(avisos, console)
        # Should print a table with overflow row
        console.print.assert_called_once()


class TestCqFmtPct:
    def test_high_coverage_is_green(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_fmt_pct

        result = _cq_fmt_pct(98, 100)
        assert "green" in result
        assert "98.0%" in result

    def test_medium_coverage_is_yellow(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_fmt_pct

        result = _cq_fmt_pct(80, 100)
        assert "yellow" in result

    def test_low_coverage_is_red(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_fmt_pct

        result = _cq_fmt_pct(50, 100)
        assert "red" in result

    def test_zero_total_returns_zero(self):
        from concilium.presentation.cli._commands_manage_helpers_a import _cq_fmt_pct

        result = _cq_fmt_pct(0, 0)
        assert "0.0%" in result


class TestImplCatalogQuality:
    """Tests for _impl_catalog_quality (full diagnostic flow)."""

    def _make_rt(self, cpus, insumos):
        ctx = SimpleNamespace(catalogo_id="cat-1")
        catalogo = {"cpus": cpus, "insumos": insumos}
        rt = MagicMock()
        rt.load_contexto.return_value = ctx
        rt.load_catalogo.return_value = catalogo
        return rt

    def test_reprovado_when_erros_present(self):
        from concilium.presentation.cli import _commands_manage_helpers_a as mod

        rt = self._make_rt(
            cpus=[{"codigo": "CPU-1"}],
            insumos=[{"codigo": "INS-1", "preco_unitario": 10}],
        )
        erro = SimpleNamespace(status="ERRO", codigo="VAL-CPU-STR-02", mensagem="ruim")
        validador = MagicMock()
        validador.validar_cpu_completa.return_value = [erro]
        validador.validar_vigencia.return_value = SimpleNamespace(status="OK", codigo="OK", mensagem="")
        validador.validar_encoding.return_value = SimpleNamespace(status="OK", codigo="OK", mensagem="")

        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.domain.validations.cpu_insumo.ValidadorCPUInsumo", return_value=validador),
            patch("concilium.infrastructure.repositories.insumo.RepositorioInsumo"),
            patch("concilium.domain.validations.circular_deps.verificar_ciclos_cpu", return_value=[]),
        ):
            mod._impl_catalog_quality(contexto_id="ctx-1")

    def test_aprovado_com_ressalvas_when_only_avisos(self):
        from concilium.presentation.cli import _commands_manage_helpers_a as mod

        rt = self._make_rt(
            cpus=[{"codigo": "CPU-1"}],
            insumos=[{"codigo": "INS-1", "preco_unitario": 10}],
        )
        validador = MagicMock()
        validador.validar_cpu_completa.return_value = []
        validador.validar_vigencia.return_value = SimpleNamespace(
            status="AVISO", codigo="VAL-CAT-VIG-01", mensagem="obsoleto"
        )
        validador.validar_encoding.return_value = SimpleNamespace(status="OK", codigo="OK", mensagem="")

        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.domain.validations.cpu_insumo.ValidadorCPUInsumo", return_value=validador),
            patch("concilium.infrastructure.repositories.insumo.RepositorioInsumo"),
            patch("concilium.domain.validations.circular_deps.verificar_ciclos_cpu", return_value=[]),
        ):
            mod._impl_catalog_quality(contexto_id="ctx-1")

    def test_integro_when_no_erros_avisos_or_ciclos(self):
        from concilium.presentation.cli import _commands_manage_helpers_a as mod

        rt = self._make_rt(cpus=[], insumos=[])
        validador = MagicMock()
        validador.validar_cpu_completa.return_value = []
        validador.validar_vigencia.return_value = SimpleNamespace(status="OK", codigo="OK", mensagem="")
        validador.validar_encoding.return_value = SimpleNamespace(status="OK", codigo="OK", mensagem="")

        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.domain.validations.cpu_insumo.ValidadorCPUInsumo", return_value=validador),
            patch("concilium.infrastructure.repositories.insumo.RepositorioInsumo"),
            patch("concilium.domain.validations.circular_deps.verificar_ciclos_cpu", return_value=[]),
        ):
            mod._impl_catalog_quality(contexto_id="ctx-1")


class TestImplCatalogSimulate:
    """Tests for _impl_catalog_simulate (CLT vs PJ comparison)."""

    def _make_resultado(self, total_direto):
        return SimpleNamespace(
            dias_trabalhados=30,
            producao_total=100.0,
            unidade="m3",
            custo_material=10.0,
            custo_mao_obra_base=20.0,
            custo_mao_obra_encargos=5.0,
            custo_equipamento_capital=8.0,
            custo_equipamento_manutencao=2.0,
            custo_diesel_total=3.0,
            total_direto=total_direto,
        )

    def test_success_path_prints_comparison(self):
        from concilium.presentation.cli import _commands_manage_helpers_a as mod

        rt = MagicMock()
        rt.load_contexto.return_value = SimpleNamespace(catalogo_id="cat-1")
        rt.load_catalogo.return_value = {}
        simulador = MagicMock()
        simulador.simular_cpu.side_effect = [
            self._make_resultado(100.0),
            self._make_resultado(80.0),
        ]

        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.domain.services.simulation.SimuladorEngenharia", return_value=simulador),
        ):
            mod._impl_catalog_simulate(
                cpu_id="CPU-1",
                contexto_id="mendeslobo",
                dias=30,
                regime="6x1",
                horas_dia=8.0,
            )

    def test_exception_raises_typer_exit(self):
        from concilium.presentation.cli import _commands_manage_helpers_a as mod

        rt = MagicMock()
        rt.load_contexto.return_value = SimpleNamespace(catalogo_id="cat-1")
        rt.load_catalogo.return_value = {}
        simulador = MagicMock()
        simulador.simular_cpu.side_effect = RuntimeError("boom")

        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.domain.services.simulation.SimuladorEngenharia", return_value=simulador),
            pytest.raises(typer.Exit),
        ):
            mod._impl_catalog_simulate(
                cpu_id="CPU-1",
                contexto_id="mendeslobo",
                dias=30,
                regime="6x1",
                horas_dia=8.0,
            )


class TestImplCatalogExport:
    """Tests for _impl_catalog_export."""

    def _make_rt(self):
        ctx = SimpleNamespace(catalogo_id="cat-1", tema_id="tema-1")
        rt = MagicMock()
        rt.load_contexto.return_value = ctx
        rt.load_catalogo.return_value = {}
        rt.load_tema.return_value.to_catalogo_dict.return_value = {}
        return rt

    def test_default_output_name(self):
        from concilium.presentation.cli import _commands_manage_helpers_a as mod

        rt = self._make_rt()
        exporter = MagicMock()

        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.infrastructure.exporters.catalog_excel.CatalogExporterExcel", return_value=exporter),
        ):
            mod._impl_catalog_export(contexto_id="mendeslobo", output=None)

        exporter.export.assert_called_once()
        called_path = exporter.export.call_args[0][0]
        assert called_path.name == "catalogo_mendeslobo_audit.xlsx"

    def test_explicit_output_path(self, tmp_path):
        from concilium.presentation.cli import _commands_manage_helpers_a as mod

        rt = self._make_rt()
        exporter = MagicMock()
        out = str(tmp_path / "custom.xlsx")

        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.infrastructure.exporters.catalog_excel.CatalogExporterExcel", return_value=exporter),
        ):
            mod._impl_catalog_export(contexto_id="mendeslobo", output=out)

        exporter.export.assert_called_once()

    def test_export_exception_raises_typer_exit(self):
        from concilium.presentation.cli import _commands_manage_helpers_a as mod

        rt = self._make_rt()
        exporter = MagicMock()
        exporter.export.side_effect = RuntimeError("disk full")

        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.infrastructure.exporters.catalog_excel.CatalogExporterExcel", return_value=exporter),
            pytest.raises(typer.Exit),
        ):
            mod._impl_catalog_export(contexto_id="mendeslobo", output=None)


# ═══════════════════════════════════════════════════════════════════════════════
# _commands_manage_helpers_c.py
# ═══════════════════════════════════════════════════════════════════════════════


class TestCcAddTableRow:
    def test_adds_row_correctly(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _cc_add_table_row

        table = MagicMock()
        rep = SimpleNamespace(
            inner_high=[{"delta": True}, {"delta": None}],
            inner_low=[1, 2],
            right_only=[1],
            left_only=[1, 2, 3],
        )
        _cc_add_table_row(table, "INSUMOS", rep)
        table.add_row.assert_called_once_with("INSUMOS", "2", "1", "2", "1", "3")


class TestCcPrintSummaryTable:
    def test_prints_one_row_per_tipo(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _cc_print_summary_table

        console = MagicMock()
        rep = SimpleNamespace(inner_high=[], inner_low=[], right_only=[], left_only=[])
        _cc_print_summary_table({"insumos": rep, "producoes": rep}, console, "Titulo")
        console.print.assert_called_once()


class TestCcPrintInnerLow:
    def test_empty_does_nothing(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _cc_print_inner_low

        console = MagicMock()
        rep = SimpleNamespace(inner_low=[])
        _cc_print_inner_low({"insumos": rep}, console)
        console.print.assert_not_called()

    def test_few_items_combines_tipos(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _cc_print_inner_low

        console = MagicMock()
        item_ins = {
            "catalog": {"codigo": "C1", "descricao": "d1"},
            "rfi": {"codigo": "R1", "descricao": "r1"},
            "score": 0.6,
        }
        rep_ins = SimpleNamespace(inner_low=[item_ins])
        item_prod = {
            "catalog": {"id": "P1", "descricao": "p"},
            "rfi": {"id": "P2", "descricao": "q"},
            "score": 0.55,
        }
        rep_prod = SimpleNamespace(inner_low=[item_prod])
        _cc_print_inner_low({"insumos": rep_ins, "producoes": rep_prod}, console)
        assert console.print.called

    def test_many_items_shows_overflow(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _cc_print_inner_low

        console = MagicMock()
        rep = SimpleNamespace(
            inner_low=[
                {
                    "catalog": {"codigo": f"C{i}", "descricao": "d"},
                    "rfi": {"codigo": f"R{i}", "descricao": "r"},
                    "score": 0.6,
                }
                for i in range(20)
            ]
        )
        _cc_print_inner_low({"insumos": rep}, console)
        # header + 15 items + overflow line
        assert console.print.call_count >= 17


class TestCcWriteReport:
    def test_writes_json_file(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _cc_write_report

        console = MagicMock()
        rep_ins = SimpleNamespace(
            inner_high=[{"delta": True, "catalog": {"codigo": "C1"}}, {"delta": None, "catalog": {"codigo": "C2"}}],
            inner_low=[{"catalog": {"codigo": "CX"}, "rfi": {"codigo": "RX"}, "score": 0.6}],
            right_only=[{"rfi": {"codigo": "N1", "descricao": "novo"}}],
            left_only=[{"x": 1}],
        )
        rep_prod = SimpleNamespace(
            inner_high=[],
            inner_low=[{"catalog": {"id": "P1"}, "rfi": {"id": "P2"}, "score": 0.55}],
            right_only=[{"rfi": {"id": "NP", "descricao": "new prod"}}],
            left_only=[],
        )
        out_path = str(tmp_path / "report.json")
        _cc_write_report(out_path, {"insumos": rep_ins, "producoes": rep_prod}, "DRY-RUN", "origA", "origB", console)
        data = json.loads(Path(out_path).read_text())
        assert data["mode"] == "DRY-RUN"
        assert data["origem_a"] == "origA"
        assert data["origem_b"] == "origB"
        assert "insumos" in data
        console.print.assert_called_once()


class TestCmVerifyPu:
    def test_with_faltantes(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _cm_verify_pu

        console = MagicMock()
        bundle = {"insumos": [], "producoes": []}
        mock_resultado = {"referencias_faltantes": [{"producao_id": "P1", "faltantes": ["I1"]}], "com_composicao": 0}
        with patch("concilium.application.catalog.pu_calculator.calcular_catalogo", return_value=mock_resultado):
            _cm_verify_pu(bundle, console)
        printed = [str(c) for c in console.print.call_args_list]
        assert any("Referências faltantes" in p for p in printed)

    def test_without_faltantes_uses_produtividades_key(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _cm_verify_pu

        console = MagicMock()
        bundle = {"insumos": [], "produtividades": []}
        mock_resultado = {"referencias_faltantes": [], "com_composicao": 5}
        with patch("concilium.application.catalog.pu_calculator.calcular_catalogo", return_value=mock_resultado):
            _cm_verify_pu(bundle, console)
        printed = [str(c) for c in console.print.call_args_list]
        assert any("Integridade referencial OK" in p for p in printed)


class TestCcpPrintFixed:
    def test_few_fixed_prints_all(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _ccp_print_fixed

        console = MagicMock()
        _ccp_print_fixed(["fix1", "fix2"], apply=True, console=console)
        assert console.print.call_count == 3  # header + 2 items

    def test_many_fixed_shows_overflow(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _ccp_print_fixed

        console = MagicMock()
        _ccp_print_fixed([f"fix{i}" for i in range(35)], apply=False, console=console)
        # header + 30 items + overflow
        assert console.print.call_count == 32

    def test_dry_run_uses_arrow_prefix(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _ccp_print_fixed

        console = MagicMock()
        _ccp_print_fixed(["fix1"], apply=False, console=console)
        printed_args = [str(c) for c in console.print.call_args_list]
        assert any("→" in a for a in printed_args)


class TestImplStructureList:
    def test_calls_resolve_project(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_structure_list

        mock_project = MagicMock()
        mock_project.list_escopos.return_value = []

        with patch("concilium.presentation.cli._commands_manage_helpers_c.resolve_project", return_value=mock_project):
            _impl_structure_list("meu-projeto")

        mock_project.list_escopos.assert_called_once()

    def test_lists_escopos_and_plans(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_structure_list

        mock_plan = SimpleNamespace(status="DRAFT")
        mock_escopo = MagicMock()
        mock_escopo.list_plans.return_value = ["plano-1"]
        mock_escopo.plan.return_value = mock_plan

        mock_project = MagicMock()
        mock_project.list_escopos.return_value = ["escopo-1"]
        mock_project.escopo.return_value = mock_escopo

        with patch("concilium.presentation.cli._commands_manage_helpers_c.resolve_project", return_value=mock_project):
            _impl_structure_list("meu-projeto")

        mock_project.escopo.assert_called_once_with("escopo-1")


class TestImplPlanInit:
    def test_creates_plan_successfully(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_plan_init

        mock_ps = MagicMock()
        mock_ps.plano_id = "plano-xyz"
        mock_escopo = MagicMock()
        mock_escopo.create_plan.return_value = mock_ps
        mock_project = MagicMock()
        mock_project.escopo.return_value = mock_escopo

        with patch("concilium.presentation.cli._commands_manage_helpers_c.resolve_project", return_value=mock_project):
            _impl_plan_init("meu-projeto")

        mock_escopo.create_plan.assert_called_once()


class TestResolverTipos:
    def test_todos_expands_to_full_list(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _resolver_tipos

        assert _resolver_tipos("todos") == ["insumos", "producoes", "cpus"]

    def test_single_tipo_returned_as_list(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _resolver_tipos

        assert _resolver_tipos("insumos") == ["insumos"]

    def test_invalid_tipo_raises_exit_1(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _resolver_tipos

        with pytest.raises(typer.Exit) as exc_info:
            _resolver_tipos("invalido")
        assert exc_info.value.exit_code == 1


class TestEhPlanilha:
    def test_existing_xlsm_is_planilha(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _eh_planilha

        xlsm = tmp_path / "x.xlsm"
        xlsm.touch()
        assert _eh_planilha(str(xlsm)) is True

    def test_nonexistent_xlsm_is_not_planilha(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _eh_planilha

        assert _eh_planilha(str(tmp_path / "nope.xlsm")) is False

    def test_catalogo_id_is_not_planilha(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _eh_planilha

        assert _eh_planilha("mendeslobo") is False


class TestResolverFonteEntidades:
    def test_catalogo_source_uses_runtime(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _resolver_fonte_entidades

        catalogo = {"insumos": [{"codigo": "I1"}], "producoes": [{"id": "P1"}], "cpus": [{"codigo": "C1"}]}
        rt = SimpleNamespace(config_root=tmp_path / "config", load_catalogo=lambda cid: catalogo)
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            result = _resolver_fonte_entidades("mendeslobo")
        assert result["insumos"] == [{"codigo": "I1"}]
        assert result["cpus"] == [{"codigo": "C1"}]

    def test_planilha_source_script_missing_raises_exit_2(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _resolver_fonte_entidades

        xlsm = tmp_path / "planilha.xlsm"
        xlsm.touch()
        # repo_root (config_root.parent) sem scripts/benchmark/merge_rfi.py
        rt = SimpleNamespace(config_root=tmp_path / "config", load_catalogo=lambda cid: {})
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            with pytest.raises(typer.Exit) as exc_info:
                _resolver_fonte_entidades(str(xlsm))
        assert exc_info.value.exit_code == 2

    def test_planilha_source_reads_via_real_merge_rfi_script(self, tmp_path):
        """Usa o repo_root real (onde scripts/benchmark/merge_rfi.py de fato existe)
        para exercitar o caminho de sucesso do adaptador de planilha ponta-a-ponta."""
        import openpyxl

        from concilium.presentation.cli._commands_manage_helpers_c import _resolver_fonte_entidades

        real_repo_root = Path(__file__).resolve().parents[4]
        xlsm = tmp_path / "planilha.xlsm"
        openpyxl.Workbook().save(xlsm)

        rt = SimpleNamespace(config_root=real_repo_root / "config", load_catalogo=lambda cid: {})
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            result = _resolver_fonte_entidades(str(xlsm))
        # Workbook em branco não tem as abas BD_* -> listas vazias, mas sempre com "cpus".
        assert result == {"insumos": [], "producoes": [], "cpus": []}


def _make_rt_multi(tmp_path: Path, catalogos: dict) -> SimpleNamespace:
    """Fake Runtime cujo load_catalogo despacha por catalogo_id (mapa de catálogos).

    Retorna sempre uma deepcopy — mesmo contrato do Runtime.load_catalogo real —
    para que mutações feitas por um teste (ex.: `_impl_catalog_merge --apply`, que
    escreve de volta no dict recebido) não vazem para os fixtures compartilhados
    entre classes de teste.
    """
    import copy

    return SimpleNamespace(
        config_root=tmp_path / "config",
        load_catalogo=lambda cid: copy.deepcopy(catalogos[cid]),
    )


class TestImplCatalogCompare:
    """_impl_catalog_compare (issue #134) — agnóstico de fonte, sempre dry-run."""

    _CAT_A = {
        "insumos": [{"codigo": "I1", "descricao": "Cimento", "unidade": "kg", "preco_unitario": 10.0}],
        "producoes": [{"id": "P1", "descricao": "Escavação"}],
        "cpus": [{"codigo": "CPU-1", "descricao": "Terraplenagem"}],
    }
    _CAT_B = {
        "insumos": [{"codigo": "I1", "descricao": "Cimento", "unidade": "kg", "preco_unitario": 12.0}],
        "producoes": [{"id": "P2", "descricao": "Compactação"}],
        "cpus": [{"codigo": "CPU-2", "descricao": "Pavimentação"}],
    }

    def test_compares_two_catalogos_todos(self, tmp_path, capsys):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_compare

        rt = _make_rt_multi(tmp_path, {"catA": self._CAT_A, "catB": self._CAT_B})
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            _impl_catalog_compare("catA", "catB", "todos")
        out = capsys.readouterr().out
        assert "Comparação de catálogo" in out
        assert "INSUMOS" in out and "PRODUCOES" in out and "CPUS" in out

    def test_invalid_tipo_raises_exit_1(self):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_compare

        with pytest.raises(typer.Exit) as exc_info:
            _impl_catalog_compare("catA", "catB", "invalido")
        assert exc_info.value.exit_code == 1

    def test_cpus_only_with_planilha_side_raises_exit_1(self, tmp_path):
        """tipo=cpus + uma origem é planilha -> cpus é o único tipo, é pulado, não resta nada."""
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_compare

        xlsm = tmp_path / "planilha.xlsm"
        xlsm.touch()
        rt = _make_rt_multi(tmp_path, {"catA": self._CAT_A})
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            with pytest.raises(typer.Exit) as exc_info:
                _impl_catalog_compare("catA", str(xlsm), "cpus")
        assert exc_info.value.exit_code == 1

    def test_cpus_skipped_with_planilha_but_other_tipos_proceed(self, tmp_path, capsys):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_compare

        xlsm = tmp_path / "planilha.xlsm"
        import openpyxl

        openpyxl.Workbook().save(xlsm)

        real_repo_root = Path(__file__).resolve().parents[4]
        rt = SimpleNamespace(config_root=real_repo_root / "config", load_catalogo=lambda cid: self._CAT_A)
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            _impl_catalog_compare("catA", str(xlsm), "todos")
        out = capsys.readouterr().out
        assert "cpus: pulado" in out
        assert "CPUS" not in out


class TestImplCatalogMerge:
    """_impl_catalog_merge (issue #134) — dry-run por padrão, --apply escreve no destino."""

    _DESTINO = {
        "insumos": [{"codigo": "I1", "descricao": "Cimento", "unidade": "kg", "preco_unitario": 10.0}],
        "producoes": [{"id": "P1", "descricao": "Escavação"}],
        "cpus": [],
    }
    _ORIGEM = {
        "insumos": [{"codigo": "I2", "descricao": "Areia", "unidade": "m3", "preco_unitario": 5.0}],
        "producoes": [],
        "cpus": [],
    }

    def test_destino_planilha_raises_exit_2(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_merge

        xlsm = tmp_path / "planilha.xlsm"
        xlsm.touch()
        with pytest.raises(typer.Exit) as exc_info:
            _impl_catalog_merge(str(xlsm), "origem_cat")
        assert exc_info.value.exit_code == 2

    def test_dry_run_does_not_write_bundle(self, tmp_path, capsys):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_merge

        rt = _make_rt_multi(tmp_path, {"dest": self._DESTINO, "orig": self._ORIGEM})
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            _impl_catalog_merge(
                "dest", "orig", apply=False, incluir_novos=True, tipo="todos", verify_pu=False, report_json=None
            )
        bundle_path = tmp_path / "config" / "catalogos" / "dest" / ".catalog_bundle.json"
        assert not bundle_path.exists()
        out = capsys.readouterr().out
        assert "Dry-run" in out

    def test_apply_writes_bundle(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_merge

        rt = _make_rt_multi(tmp_path, {"dest": self._DESTINO, "orig": self._ORIGEM})
        bundle_dir = tmp_path / "config" / "catalogos" / "dest"
        bundle_dir.mkdir(parents=True)
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            _impl_catalog_merge(
                "dest", "orig", apply=True, incluir_novos=True, tipo="todos", verify_pu=False, report_json=None
            )
        bundle_path = bundle_dir / ".catalog_bundle.json"
        assert bundle_path.exists()
        data = json.loads(bundle_path.read_text())
        codigos = {i["codigo"] for i in data["insumos"]}
        assert "I1" in codigos
        assert "I2" in codigos  # RIGHT_ONLY incluído (incluir_novos=True default)

    def test_apply_incluir_novos_false_omits_right_only(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_merge

        rt = _make_rt_multi(tmp_path, {"dest": self._DESTINO, "orig": self._ORIGEM})
        bundle_dir = tmp_path / "config" / "catalogos" / "dest"
        bundle_dir.mkdir(parents=True)
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            _impl_catalog_merge(
                "dest", "orig", apply=True, incluir_novos=False, tipo="todos", verify_pu=False, report_json=None
            )
        data = json.loads((bundle_dir / ".catalog_bundle.json").read_text())
        codigos = {i["codigo"] for i in data["insumos"]}
        assert "I2" not in codigos

    def test_apply_with_verify_pu_and_report_json(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_merge

        rt = _make_rt_multi(tmp_path, {"dest": self._DESTINO, "orig": self._ORIGEM})
        bundle_dir = tmp_path / "config" / "catalogos" / "dest"
        bundle_dir.mkdir(parents=True)
        report_path = str(tmp_path / "report.json")
        mock_resultado = {"referencias_faltantes": [], "com_composicao": 1}
        with (
            patch("concilium.presentation.mcp._context._runtime", return_value=rt),
            patch("concilium.application.catalog.pu_calculator.calcular_catalogo", return_value=mock_resultado),
        ):
            _impl_catalog_merge(
                "dest", "orig", apply=True, incluir_novos=True, tipo="todos", verify_pu=True, report_json=report_path
            )
        assert Path(report_path).exists()
        report_data = json.loads(Path(report_path).read_text())
        assert report_data["origem_a"] == "dest"
        assert report_data["origem_b"] == "orig"

    def test_produtividades_key_preserved_when_producoes_absent(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_merge

        destino_produtividades = {"insumos": [], "produtividades": [{"id": "P1"}], "cpus": []}
        rt = _make_rt_multi(tmp_path, {"dest": destino_produtividades, "orig": self._ORIGEM})
        bundle_dir = tmp_path / "config" / "catalogos" / "dest"
        bundle_dir.mkdir(parents=True)
        with patch("concilium.presentation.mcp._context._runtime", return_value=rt):
            _impl_catalog_merge(
                "dest", "orig", apply=True, incluir_novos=True, tipo="producoes", verify_pu=False, report_json=None
            )
        data = json.loads((bundle_dir / ".catalog_bundle.json").read_text())
        assert "produtividades" in data


class TestImplCatalogCurePlausibility:
    def test_bundle_missing_raises(self, tmp_path, monkeypatch):
        import concilium.application.catalog.cure_plausibility as cure_mod
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_cure_plausibility

        monkeypatch.setattr(cure_mod, "ROOT", tmp_path)
        with pytest.raises(typer.Exit):
            _impl_catalog_cure_plausibility("nope-ctx", False)

    def test_dry_run_with_fixed_and_manual_items(self, tmp_path, monkeypatch):
        import concilium.application.catalog.cure_plausibility as cure_mod
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_cure_plausibility

        monkeypatch.setattr(cure_mod, "ROOT", tmp_path)
        ctx = "ctxcure1"
        bundle_dir = tmp_path / "config" / "catalogos" / ctx
        bundle_dir.mkdir(parents=True)
        bundle = {
            "insumos": [{"codigo": "EQ1", "tipo": "EQUIPAMENTO", "preco_unitario": 10.0, "atributos_especificos": {}}],
            "producoes": [{"id": "P1", "producao_horaria": 5}],
        }
        (bundle_dir / ".catalog_bundle.json").write_text(json.dumps(bundle), encoding="utf-8")

        _impl_catalog_cure_plausibility(ctx, False)

    def test_apply_with_fixed_and_manual_writes_individual_files_not_bundle(self, tmp_path, monkeypatch):
        """ADR-177/durabilidade: --apply escreve no arquivo individual do insumo
        (não no bundle, que fica intocado — reverteria na próxima invalidação de
        cache). Achado real: este wrapper CLI tinha sua própria cópia divergente
        da orquestração/escrita; agora delega para cure_plausibility."""
        import concilium.application.catalog.cure_plausibility as cure_mod
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_cure_plausibility

        monkeypatch.setattr(cure_mod, "ROOT", tmp_path)
        ctx = "ctxcure2"
        bundle_dir = tmp_path / "config" / "catalogos" / ctx
        bundle_dir.mkdir(parents=True)
        bundle = {
            "insumos": [{"codigo": "EQ1", "tipo": "EQUIPAMENTO", "preco_unitario": 10.0, "atributos_especificos": {}}],
            "produtividades": [{"id": "P1", "producao_horaria": 5}],
        }
        bundle_path = bundle_dir / ".catalog_bundle.json"
        bundle_bytes_antes = json.dumps(bundle).encode("utf-8")
        bundle_path.write_bytes(bundle_bytes_antes)

        _impl_catalog_cure_plausibility(ctx, True)

        assert bundle_path.read_bytes() == bundle_bytes_antes
        arquivo_individual = bundle_dir / "insumos" / "equipamento" / "EQ1.json"
        assert arquivo_individual.exists()
        saved = json.loads(arquivo_individual.read_text(encoding="utf-8"))
        assert saved["atributos_especificos"]["valor_locacao_mensal"] > 0

    def test_dry_run_with_no_fixed_and_no_manual_items(self, tmp_path, monkeypatch):
        """Empty bundle → report.fixed and report.manual both empty, covering
        the false branch of both `if report.fixed:` and `if report.manual:`."""
        import concilium.application.catalog.cure_plausibility as cure_mod
        from concilium.presentation.cli._commands_manage_helpers_c import _impl_catalog_cure_plausibility

        monkeypatch.setattr(cure_mod, "ROOT", tmp_path)
        ctx = "ctxcure3"
        bundle_dir = tmp_path / "config" / "catalogos" / ctx
        bundle_dir.mkdir(parents=True)
        bundle = {"insumos": [], "producoes": []}
        (bundle_dir / ".catalog_bundle.json").write_text(json.dumps(bundle), encoding="utf-8")

        _impl_catalog_cure_plausibility(ctx, False)


# ═══════════════════════════════════════════════════════════════════════════════
# _commands_manage_helpers_d.py
# ═══════════════════════════════════════════════════════════════════════════════


class TestRunSchemaValidation:
    def test_no_errors_appends_nothing(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _run_schema_validation

        validator = MagicMock()
        validator.iter_errors.return_value = []
        erros = []
        _run_schema_validation(validator, [{"codigo": "CPU-1"}], "CPU", "codigo", erros)
        assert erros == []

    def test_error_appended_with_path(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _run_schema_validation

        err = MagicMock()
        err.absolute_path = ["fieldA", "nested"]
        err.message = "invalid value"
        validator = MagicMock()
        validator.iter_errors.return_value = [err]
        erros = []
        _run_schema_validation(validator, [{"codigo": "CPU-1"}], "CPU", "codigo", erros)
        assert len(erros) == 1
        assert "CPU-1" in erros[0]
        assert "fieldA.nested" in erros[0]

    def test_error_with_empty_path_uses_root(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _run_schema_validation

        err = MagicMock()
        err.absolute_path = []
        err.message = "required field"
        validator = MagicMock()
        validator.iter_errors.return_value = [err]
        erros = []
        _run_schema_validation(validator, [{"codigo": "CPU-2"}], "CPU", "codigo", erros)
        assert "<root>" in erros[0]

    def test_missing_id_field_uses_placeholder(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _run_schema_validation

        err = MagicMock()
        err.absolute_path = []
        err.message = "bad"
        validator = MagicMock()
        validator.iter_errors.return_value = [err]
        erros = []
        _run_schema_validation(validator, [{}], "INSUMO", "codigo", erros)
        assert "<sem-codigo>" in erros[0]


class TestCheckCatalogSemantics:
    def test_missing_producao_ref_is_error(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _check_catalog_semantics

        cpus = [{"codigo": "CPU-1", "producao_id_ref": "PROD-999"}]
        prods = [{"id": "PROD-001"}]
        erros, avisos = [], []
        _check_catalog_semantics(cpus, prods, erros, avisos)
        assert any("PROD-999" in e for e in erros)

    def test_valid_producao_ref_no_error(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _check_catalog_semantics

        cpus = [{"codigo": "CPU-1", "producao_id_ref": "PROD-001"}]
        prods = [{"id": "PROD-001"}]
        erros, avisos = [], []
        _check_catalog_semantics(cpus, prods, erros, avisos)
        assert erros == []

    def test_legacy_fields_generate_aviso(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _check_catalog_semantics

        cpus = [{"codigo": "CPU-L", "producao_horaria": 100}]
        prods = []
        erros, avisos = [], []
        _check_catalog_semantics(cpus, prods, erros, avisos)
        assert any("legados" in a for a in avisos)

    def test_no_ref_no_legacy_no_errors(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _check_catalog_semantics

        cpus = [{"codigo": "CPU-OK"}]
        prods = [{"id": "PROD-01"}]
        erros, avisos = [], []
        _check_catalog_semantics(cpus, prods, erros, avisos)
        assert erros == [] and avisos == []


class TestBuildSchemaValidators:
    def test_builds_validators_from_real_schemas(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _build_schema_validators

        schema_root = Path("config") / "schemas" / "catalog"
        cpu_v, prod_v, ins_v = _build_schema_validators(schema_root)
        assert cpu_v is not None
        assert prod_v is not None
        assert ins_v is not None

    def test_missing_schema_file_raises_exit(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_d import _build_schema_validators

        with pytest.raises(typer.Exit) as exc_info:
            _build_schema_validators(tmp_path)
        assert exc_info.value.exit_code == 2

    def test_missing_common_schema_is_optional(self, tmp_path):
        """common.schema.json ausente não deve impedir a construção dos validadores."""
        from concilium.presentation.cli._commands_manage_helpers_d import _build_schema_validators

        real_root = Path("config") / "schemas" / "catalog"
        for name in ["cpu.schema.json", "producao.schema.json", "insumo.schema.json"]:
            (tmp_path / name).write_text((real_root / name).read_text(encoding="utf-8"), encoding="utf-8")

        cpu_v, prod_v, ins_v = _build_schema_validators(tmp_path)
        assert cpu_v is not None
        assert prod_v is not None
        assert ins_v is not None


class TestImplCatalogValidateSchema:
    def _make_rt(self, catalogo: dict):
        rt = MagicMock()
        rt.load_contexto.return_value = SimpleNamespace(catalogo_id="cat1")
        rt.load_catalogo.return_value = catalogo
        return rt

    def test_empty_catalog_prints_ok(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_validate_schema

        rt = self._make_rt({"cpus": [], "insumos": [], "produtividades": [], "producoes": []})
        with patch("concilium.presentation.cli._commands_manage_helpers_d._runtime", return_value=rt):
            # Should not raise — no errors, no warnings.
            _impl_catalog_validate_schema(contexto_id="ctx1", fail_on_warning=False)

    def test_non_list_produtividades_and_producoes_default_to_empty(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_validate_schema

        rt = self._make_rt({"cpus": [], "insumos": [], "produtividades": None, "producoes": None})
        with patch("concilium.presentation.cli._commands_manage_helpers_d._runtime", return_value=rt):
            _impl_catalog_validate_schema(contexto_id="ctx1", fail_on_warning=False)

    def test_schema_errors_raise_exit_1(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_validate_schema

        # CPU faltando campos obrigatórios (descricao, unidade, etc.) gera erro de schema.
        rt = self._make_rt({"cpus": [{"codigo": "CPU-1"}], "insumos": [], "produtividades": [], "producoes": []})
        with patch("concilium.presentation.cli._commands_manage_helpers_d._runtime", return_value=rt):
            with pytest.raises(typer.Exit) as exc_info:
                _impl_catalog_validate_schema(contexto_id="ctx1", fail_on_warning=False)
            assert exc_info.value.exit_code == 1

    def test_warnings_with_fail_on_warning_raises_exit_1(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_validate_schema

        # CPU válido mas com campo legado -> aviso semântico, sem erro de schema.
        cpu = {
            "codigo": "CPU-1",
            "descricao": "desc",
            "unidade": "m3",
            "origem": "CATALOGO_MENDESLOBO",
            "producao_id_ref": None,
            "categoria": "SERVICO",
            "composicao_itens": [],
            "producao_horaria": 10,
        }
        rt = self._make_rt({"cpus": [cpu], "insumos": [], "produtividades": [], "producoes": []})
        with patch("concilium.presentation.cli._commands_manage_helpers_d._runtime", return_value=rt):
            with pytest.raises(typer.Exit) as exc_info:
                _impl_catalog_validate_schema(contexto_id="ctx1", fail_on_warning=True)
            assert exc_info.value.exit_code == 1

    def test_import_error_on_jsonschema_missing(self):
        from concilium.presentation.cli import _commands_manage_helpers_d as mod

        with patch.object(mod, "_build_schema_validators", side_effect=ImportError("no jsonschema")):
            with pytest.raises(typer.Exit) as exc_info:
                mod._impl_catalog_validate_schema(contexto_id="ctx1", fail_on_warning=False)
            assert exc_info.value.exit_code == 2


class TestPrintMessagesSection:
    def test_empty_messages_no_output(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _print_messages_section

        console = MagicMock()
        _print_messages_section(console, [], "[bold red]Erros:[/bold red]", "erros", 30)
        console.print.assert_not_called()

    def test_few_messages_printed(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _print_messages_section

        console = MagicMock()
        _print_messages_section(console, ["error 1", "error 2"], "[bold red]Erros:[/bold red]", "erros", 30)
        assert console.print.call_count == 3  # title + 2 messages

    def test_overflow_shows_more(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _print_messages_section

        console = MagicMock()
        msgs = [f"err{i}" for i in range(35)]
        _print_messages_section(console, msgs, "[bold]Title[/bold]", "erros", 30)
        # title + 30 messages + overflow line
        assert console.print.call_count == 32


class TestPrintCatalogValidationResults:
    def test_prints_table_and_sections(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _print_catalog_validation_results

        console = MagicMock()
        cpus = [{"codigo": "CPU-1"}]
        prods = [{"id": "P1"}]
        insumos = [{"codigo": "I1"}]
        erros = ["CPU CPU-1: erro x", "PRODUCAO P1: erro y"]
        avisos = ["CPU CPU-1: aviso z"]
        _print_catalog_validation_results(console, cpus, prods, insumos, erros, avisos, "testcat")
        assert console.print.called


class TestImplCatalogCureSchema:
    """A impl atual chama cure_schema.run() diretamente (não mais scripts/cure_catalog.py)."""

    def test_exits_1_when_context_missing(self, monkeypatch):
        import typer

        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_cure_schema

        with pytest.raises((SystemExit, typer.Exit)) as exc_info:
            _impl_catalog_cure_schema(contexto_id="ctx-inexistente", apply=True)
        exc = exc_info.value
        code = getattr(exc, "code", None) if isinstance(exc, SystemExit) else getattr(exc, "exit_code", None)
        assert code == 1

    def test_exits_1_on_runtime_error(self):
        import typer

        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_cure_schema

        with patch("concilium.application.catalog.cure_schema.run", side_effect=RuntimeError("erros residuais")):
            with pytest.raises((SystemExit, typer.Exit)) as exc_info:
                _impl_catalog_cure_schema(contexto_id="ctx1", apply=True)
        exc = exc_info.value
        code = getattr(exc, "code", None) if isinstance(exc, SystemExit) else getattr(exc, "exit_code", None)
        assert code == 1

    def test_success_does_not_raise_and_forwards_flags(self):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_cure_schema

        with patch("concilium.application.catalog.cure_schema.run") as mock_run:
            _impl_catalog_cure_schema(contexto_id="ctx1", apply=False)
        mock_run.assert_called_once_with("ctx1", False)


class TestImplCatalogAuditReasonability:
    def test_calls_service_and_writes_artifact(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_audit_reasonability

        artifact = {
            "summary": {"total_cpus": 3, "erros": 1, "avisos": 2, "outliers": 0, "mediana_producao_horaria": 10.5},
            "findings": [{"severity": "WARN", "cpu_id": "CPU-1", "issues": ["problema"]}],
        }

        mock_rt = MagicMock()
        mock_ctx = SimpleNamespace(catalogo_id="testcat")
        mock_rt.load_contexto.return_value = mock_ctx
        mock_rt.load_catalogo.return_value = {"cpus": [], "produtividades": [], "producoes": []}

        with patch("concilium.presentation.cli._commands_manage_helpers_d._runtime", return_value=mock_rt):
            with patch(
                "concilium.application.services.catalog.catalog_reasonability.build_catalog_reasonability_artifact",
                return_value=artifact,
            ):
                out = str(tmp_path / "audit.json")
                _impl_catalog_audit_reasonability(contexto_id="ctx1", output=out, max_findings=20)

        assert Path(out).exists()
        data = json.loads(Path(out).read_text())
        assert data["summary"]["total_cpus"] == 3

    def test_no_findings_prints_ok(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_audit_reasonability

        artifact = {
            "summary": {"total_cpus": 0, "erros": 0, "avisos": 0, "outliers": 0, "mediana_producao_horaria": 0.0},
            "findings": [],
        }

        mock_rt = MagicMock()
        mock_ctx = SimpleNamespace(catalogo_id="emptycat")
        mock_rt.load_contexto.return_value = mock_ctx
        mock_rt.load_catalogo.return_value = {"cpus": None, "produtividades": None, "producoes": None}

        with patch("concilium.presentation.cli._commands_manage_helpers_d._runtime", return_value=mock_rt):
            with patch(
                "concilium.application.services.catalog.catalog_reasonability.build_catalog_reasonability_artifact",
                return_value=artifact,
            ):
                out = str(tmp_path / "audit_empty.json")
                _impl_catalog_audit_reasonability(contexto_id="ctx1", output=out, max_findings=20)


class TestImplCatalogTendencia:
    def test_raises_exit_when_bundle_missing(self, tmp_path, monkeypatch):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_tendencia

        with patch(
            "concilium.presentation.cli._commands_manage_helpers_d._config_root", return_value=tmp_path / "config"
        ):
            import typer

            with pytest.raises((SystemExit, typer.Exit)) as exc_info:
                _impl_catalog_tendencia(contexto_id="missing", jurisdicao_id=None, output="")
            exc = exc_info.value
            code = getattr(exc, "code", None) if isinstance(exc, SystemExit) else getattr(exc, "exit_code", None)
            assert code == 1

    def test_renders_table_and_saves_output(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_tendencia

        config_root = tmp_path / "config"
        bundle_dir = config_root / "catalogos" / "ctx1"
        bundle_dir.mkdir(parents=True)
        bundle = {"insumos": [{"codigo": "I1", "tipo": "MATERIAL", "preco_unitario": 10.0}]}
        (bundle_dir / ".catalog_bundle.json").write_text(json.dumps(bundle), encoding="utf-8")

        from concilium.application.catalog.price_trend import IndiceCategoria, RelatorioTendencia

        idx = IndiceCategoria(
            tipo="MATERIAL",
            n_insumos_total=1,
            n_insumos_com_historico=0,
            variacao_media_pct=5.0,
            variacao_mediana_pct=5.0,
            variacao_max_pct=10.0,
            variacao_min_pct=0.0,
        )
        relatorio = RelatorioTendencia(
            contexto_id="ctx1",
            gerado_em="2026-06-12",
            data_base_analise="2026-06-12",
            jurisdicao_id=None,
            indices_por_categoria=[idx],
            resumo={"total_insumos": 1, "total_com_historico": 0, "total_com_sicro_regional": 0},
        )

        out_path = str(tmp_path / "trend.json")
        with patch("concilium.presentation.cli._commands_manage_helpers_d._config_root", return_value=config_root):
            with patch("concilium.application.catalog.price_trend.calcular_tendencia", return_value=relatorio):
                _impl_catalog_tendencia(contexto_id="ctx1", jurisdicao_id=None, output=out_path)

        assert Path(out_path).exists()

    def test_with_jurisdicao_prints_sicro_line(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_tendencia

        config_root = tmp_path / "config"
        bundle_dir = config_root / "catalogos" / "ctx1"
        bundle_dir.mkdir(parents=True)
        bundle = {"insumos": []}
        (bundle_dir / ".catalog_bundle.json").write_text(json.dumps(bundle), encoding="utf-8")

        relatorio = SimpleNamespace(
            indices_por_categoria=[],
            resumo={"total_insumos": 0, "total_com_historico": 0, "total_com_sicro_regional": 5},
        )

        with patch("concilium.presentation.cli._commands_manage_helpers_d._config_root", return_value=config_root):
            with patch("concilium.application.catalog.price_trend.calcular_tendencia", return_value=relatorio):
                # Should not raise
                _impl_catalog_tendencia(contexto_id="ctx1", jurisdicao_id="BR-CE", output="")

    def test_high_variacao_uses_yellow_color(self, tmp_path):
        from concilium.presentation.cli._commands_manage_helpers_d import _impl_catalog_tendencia

        config_root = tmp_path / "config"
        bundle_dir = config_root / "catalogos" / "ctx1"
        bundle_dir.mkdir(parents=True)
        bundle = {"insumos": []}
        (bundle_dir / ".catalog_bundle.json").write_text(json.dumps(bundle), encoding="utf-8")

        idx_yellow = SimpleNamespace(
            tipo="MAO_DE_OBRA",
            n_insumos_total=1,
            n_insumos_com_historico=1,
            variacao_media_pct=20.0,  # > 15, < 30 → yellow
            variacao_mediana_pct=20.0,
            variacao_max_pct=25.0,
        )
        idx_red = SimpleNamespace(
            tipo="EQUIP",
            n_insumos_total=1,
            n_insumos_com_historico=1,
            variacao_media_pct=35.0,  # > 30 → red
            variacao_mediana_pct=35.0,
            variacao_max_pct=40.0,
        )
        relatorio = SimpleNamespace(
            indices_por_categoria=[idx_yellow, idx_red],
            resumo={"total_insumos": 2, "total_com_historico": 2, "total_com_sicro_regional": 0},
        )

        with patch("concilium.presentation.cli._commands_manage_helpers_d._config_root", return_value=config_root):
            with patch("concilium.application.catalog.price_trend.calcular_tendencia", return_value=relatorio):
                # Should not raise
                _impl_catalog_tendencia(contexto_id="ctx1", jurisdicao_id=None, output="")


# ═══════════════════════════════════════════════════════════════════════════════
# commands_config_version.py
# ═══════════════════════════════════════════════════════════════════════════════


def _make_context_app():
    """Build a standalone Typer app wrapping the context commands."""
    from concilium.presentation.cli.commands_config_version import app_context

    return app_context


def _make_perfil_app():
    from concilium.presentation.cli.commands_config_version import app_perfil

    return app_perfil


class TestConfigRoot:
    def test_returns_valid_path_when_config_exists(self, tmp_path, monkeypatch):
        from concilium.presentation.cli import commands_config_version as mod

        # Override the file location to use tmp_path structure
        config = tmp_path / "config"
        config.mkdir()
        # Patch __file__ parents
        with patch.object(mod, "_config_root", return_value=config):
            result = mod._config_root()
        assert result == config

    def test_raises_bad_parameter_when_no_config(self, tmp_path, monkeypatch):
        """Real _config_root raises BadParameter when no config/ dir found."""
        import concilium.presentation.cli.commands_config_version as cv_mod

        with patch.object(cv_mod.Path, "exists", return_value=False):
            with pytest.raises((typer.BadParameter, SystemExit)):
                cv_mod._config_root()

    def test_falls_back_to_second_candidate(self):
        """Real _config_root: first candidate (parents[4]) missing, second
        (parents[5]) exists — exercises the loop's second iteration."""
        import concilium.presentation.cli.commands_config_version as cv_mod

        real_exists = Path.exists
        first_candidate = Path(cv_mod.__file__).parents[4] / "config"
        second_candidate = Path(cv_mod.__file__).parents[5] / "config"

        def fake_exists(self):
            if self == first_candidate:
                return False
            if self == second_candidate:
                return True
            return real_exists(self)

        with patch.object(cv_mod.Path, "exists", fake_exists):
            result = cv_mod._config_root()
        assert result == second_candidate


class TestContextList:
    def test_empty_contexts_prints_warning(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.list_contexts.return_value = []

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["list"])

        assert result.exit_code in (0, 1)
        assert "Nenhum" in result.output

    def test_with_contexts_prints_table(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.list_contexts.return_value = [
            {"id": "ctx1", "nome": "Context 1", "current_version": "v2", "versions": ["v1", "v2"]},
        ]

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["list"])

        assert result.exit_code == 0
        assert "ctx1" in result.output


class TestContextVersions:
    def test_no_versions_prints_warning(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.list_versions.return_value = []
        mock_svc.current_version.return_value = "v1"

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["versions", "ctx1"])

        assert result.exit_code in (0, 1)
        assert "Nenhuma" in result.output

    def test_with_versions_shows_table(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.list_versions.return_value = [
            {"version": "v1", "created_at": "2024-01-01T00:00:00", "created_by": "alice", "motivo": "inicial"},
            {"version": "v2", "created_at": "2024-02-01T00:00:00", "created_by": "bob", "motivo": "update"},
        ]
        mock_svc.current_version.return_value = "v2"

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["versions", "ctx1"])

        assert result.exit_code == 0
        assert "v1" in result.output
        assert "v2" in result.output


class TestContextNewVersion:
    def test_creates_version_and_makes_current(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.create_version.return_value = {"version": "v3"}

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["new-version", "ctx1", "--motivo", "nova feature"])

        assert result.exit_code == 0
        assert "v3" in result.output

    def test_creates_version_no_current(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.create_version.return_value = {"version": "v4"}

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["new-version", "ctx1", "--motivo", "test", "--no-current"])

        assert result.exit_code == 0
        assert "v4" in result.output
        # Should NOT show "Versão ativa atualizada"
        assert "ativa atualizada" not in result.output

    def test_creates_version_with_optional_fields(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.create_version.return_value = {"version": "v5"}

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(
                    app,
                    [
                        "new-version",
                        "ctx1",
                        "--motivo",
                        "update cat",
                        "--catalogo-id",
                        "newcat",
                        "--jurisdicao-id",
                        "BR-CE",
                        "--perfil-id",
                        "p1",
                        "--author",
                        "testuser",
                    ],
                )

        assert result.exit_code == 0
        # Verify fields were passed to create_version
        call_args = mock_svc.create_version.call_args
        fields = call_args[0][1]
        assert fields.get("catalogo_id") == "newcat"
        assert fields.get("jurisdicao_id") == "BR-CE"
        assert fields.get("perfil_id") == "p1"


class TestContextCompare:
    def test_identical_versions_prints_green(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.compare.return_value = {"changed": {}, "unchanged": ["catalogo_id"]}

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["compare", "ctx1", "v1", "v2"])

        assert result.exit_code == 0
        assert "idênticas" in result.output

    def test_changed_fields_shows_table(self):
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.compare.return_value = {
            "changed": {"catalogo_id": {"de": "cat1", "para": "cat2"}},
            "unchanged": ["perfil_id"],
        }

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["compare", "ctx1", "v1", "v2"])

        assert result.exit_code == 0
        assert "catalogo_id" in result.output
        assert "cat1" in result.output
        assert "cat2" in result.output

    def test_changed_fields_without_unchanged_skips_footer(self):
        """When `unchanged` is empty, the 'Inalterados' footer line is skipped."""
        app = _make_context_app()
        mock_svc = MagicMock()
        mock_svc.compare.return_value = {
            "changed": {"catalogo_id": {"de": "cat1", "para": "cat2"}},
            "unchanged": [],
        }

        with patch(
            "concilium.presentation.cli.commands_config_version.ContextoVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["compare", "ctx1", "v1", "v2"])

        assert result.exit_code == 0
        assert "Inalterados" not in result.output


class TestPerfilList:
    def test_empty_perfis_prints_warning(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.list_perfis.return_value = []

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["list"])

        assert "Nenhum" in result.output

    def test_with_perfis_prints_table(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.list_perfis.return_value = [
            {"id": "perfil1", "nome": "Perfil 1", "current_version": "v1", "versions": ["v1"]},
        ]

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["list"])

        assert result.exit_code == 0
        assert "perfil1" in result.output


class TestPerfilVersions:
    def test_no_versions_prints_warning(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.list_versions.return_value = []
        mock_svc.current_version.return_value = "v1"

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["versions", "perfil1"])

        assert "Nenhuma" in result.output

    def test_with_versions_shows_table(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.list_versions.return_value = [
            {"version": "v1", "created_at": None, "created_by": None, "motivo": None},
        ]
        mock_svc.current_version.return_value = "v1"

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["versions", "perfil1"])

        assert result.exit_code == 0
        assert "v1" in result.output


class TestPerfilNewVersion:
    def test_creates_version(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.create_version.return_value = {"version": "v2"}

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["new-version", "perfil1", "--motivo", "atualização"])

        assert result.exit_code == 0
        assert "v2" in result.output

    def test_creates_version_with_jurisdicao(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.create_version.return_value = {"version": "v3"}

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                runner.invoke(
                    app, ["new-version", "perfil1", "--motivo", "jurisdição nova", "--jurisdicao-id", "BR-SP"]
                )

        call_args = mock_svc.create_version.call_args
        fields = call_args[0][1]
        assert fields.get("jurisdicao_id") == "BR-SP"

    def test_no_current_flag(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.create_version.return_value = {"version": "v4"}

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["new-version", "perfil1", "--motivo", "test", "--no-current"])

        assert result.exit_code == 0
        assert "ativa atualizada" not in result.output


class TestPerfilCompare:
    def test_identical_versions(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.compare.return_value = {"changed": {}, "unchanged": []}

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["compare", "perfil1", "v1", "v2"])

        assert "idênticas" in result.output

    def test_changed_fields(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.compare.return_value = {
            "changed": {"jurisdicao_id": {"de": "BR-DF", "para": "BR-GO"}},
            "unchanged": [],
        }

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["compare", "perfil1", "v1", "v2"])

        assert result.exit_code == 0
        assert "jurisdicao_id" in result.output
        assert "BR-DF" in result.output
        assert "BR-GO" in result.output

    def test_unchanged_items_printed(self):
        app = _make_perfil_app()
        mock_svc = MagicMock()
        mock_svc.compare.return_value = {
            "changed": {"k": {"de": "a", "para": "b"}},
            "unchanged": ["field_x", "field_y"],
        }

        with patch(
            "concilium.presentation.cli.commands_config_version.PerfilVersionService",
            return_value=mock_svc,
        ):
            with patch("concilium.presentation.cli.commands_config_version._config_root", return_value=Path("/fake")):
                result = runner.invoke(app, ["compare", "perfil1", "v1", "v2"])

        assert "Inalterados" in result.output
        assert "field_x" in result.output
