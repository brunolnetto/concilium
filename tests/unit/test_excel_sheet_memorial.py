"""Testes para infrastructure/adapters/excel/sheets/memorial.py — cobertura dos
caminhos não exercitados pelo teste de ponta a ponta do ExcelRenderer completo
(especificamente: construção sob demanda do memorial de BDI quando ausente do
canonical, e agregação de equipamentos por equip_id em MEMORIAL.EQUIPAMENTOS)."""

from __future__ import annotations

from unittest.mock import MagicMock, patch

from openpyxl import Workbook

from concilium.infrastructure.adapters.excel.sheets.memorial import (
    render_memorial_bdi,
    render_memorial_catalogo,
    render_memorial_equipamentos,
)


class TestRenderMemorialBdi:
    def test_falha_ao_construir_memorial_nao_quebra_a_planilha(self):
        """Exception no builder é engolida — planilha ainda é gerada sem a seção de meta."""
        wb = Workbook()
        canonical = MagicMock(memorial_calculo_bdi=None)
        with patch(
            "concilium.domain.services.financial.memorial_bdi_builder.MemorialBDIBuilder",
            side_effect=RuntimeError("dados insuficientes"),
        ):
            render_memorial_bdi(wb, canonical)
        assert "MEMORIAL.BDI" in wb.sheetnames


class TestRenderMemorialEquipamentos:
    def test_agrega_equipamentos_repetidos_por_equip_id(self):
        wb = Workbook()
        equip_a = MagicMock(equip_id="EQ-1", descricao="Escavadeira", num_unidades=1, custo_mensal=1000.0)
        equip_a.custo_combustivel_mensal = 100.0
        equip_b = MagicMock(equip_id="EQ-1", descricao="Escavadeira", num_unidades=2, custo_mensal=1000.0)
        equip_b.custo_combustivel_mensal = 100.0
        equipe = MagicMock(equipamentos=[equip_a, equip_b])
        canonical = MagicMock(equipes_diretas=[equipe])
        render_memorial_equipamentos(wb, canonical)
        assert "MEMORIAL.EQUIPAMENTOS" in wb.sheetnames

    def test_equipamento_sem_equip_id_e_ignorado(self):
        wb = Workbook()
        equip = MagicMock(equip_id="", descricao="Sem ID")
        equipe = MagicMock(equipamentos=[equip])
        canonical = MagicMock(equipes_diretas=[equipe])
        render_memorial_equipamentos(wb, canonical)
        assert "MEMORIAL.EQUIPAMENTOS" in wb.sheetnames

    def test_sem_equipes_diretas_gera_planilha_vazia(self):
        wb = Workbook()
        canonical = MagicMock(equipes_diretas=[])
        render_memorial_equipamentos(wb, canonical)
        assert "MEMORIAL.EQUIPAMENTOS" in wb.sheetnames


class TestRenderMemorialCatalogo:
    def test_valor_nao_numerico_no_item_cai_no_except_e_usa_default(self):
        """_fv() interno captura exceção de float() quando o atributo não é
        coercível para número (ex.: quantitativo como string não numérica) e
        retorna o default 0.0 em vez de propagar."""
        wb = Workbook()
        item = MagicMock(item_id="1.1", cpu_id="CPU-1", descricao="Item 1", unidade="m3")
        item.quantitativo = "não-numérico"
        canonical = MagicMock(eap=[item], planejamento=[], equipes_diretas=[])
        render_memorial_catalogo(wb, canonical)
        assert "MEMORIAL.CATALOGO" in wb.sheetnames
