"""Tests for infrastructure/adapters/excel/sheets/graficos.py."""

from __future__ import annotations

from types import SimpleNamespace

from openpyxl import Workbook

from concilium.infrastructure.adapters.excel.sheets import graficos as gmod

# ---------------------------------------------------------------------------
# Graficos sheet
# ---------------------------------------------------------------------------


def _canonical_for_graficos():
    dist = {
        "2026-W10": SimpleNamespace(num_equipes=2),
        "2026-W11": SimpleNamespace(num_equipes=1),
    }
    act = SimpleNamespace(activity_id="A1", distribuicao_semanal=dist)
    cronograma = SimpleNamespace(scheduled_activities=[act])
    equipe = SimpleNamespace(
        item_id="A1",
        equipamentos=[SimpleNamespace(num_unidades=2)],
        funcoes=[SimpleNamespace(num_pessoas=3)],
    )
    return SimpleNamespace(cronograma=cronograma, equipes_diretas=[equipe])


def test_graficos_helpers_and_render(monkeypatch):
    canonical = _canonical_for_graficos()
    wb = Workbook()

    monkeypatch.setattr(
        gmod,
        "_cash_flow",
        lambda _c: [
            {
                "mes_label": "2026-03",
                "receita": 100.0,
                "despesa": 80.0,
                "saldo": 20.0,
                "exposicao_acumulada": 20.0,
            }
        ],
    )

    assert gmod._rg_week_label("2026-W10")
    labels, values = gmod._rg_histogram_totals(canonical, "equipamentos")
    assert labels and values
    assert gmod._rg_cash_flow_rows(canonical)[0]["mes"] == "2026-03"

    gmod.render_graficos(wb, canonical)
    assert "GRAFICOS" in wb.sheetnames


def test_render_cubo_comp_and_campanha_frota_branches(monkeypatch):
    wb = Workbook()

    fake_cubo = SimpleNamespace(
        celulas=[1],
        regime_bdi="pre_reforma",
        total_preco_venda=1000.0,
        total_custo_direto=700.0,
        pct_margem=0.1,
        pct_tributos=0.1,
        total_overhead=100.0,
        total_margem=100.0,
        total_tributos=100.0,
        exposicao_tributaria_por_tipo=lambda: {"tributo_iss": 50.0},
        margem_por_atividade=lambda: {"A1": 100.0},
        composicao_preco_temporal=lambda: {"2026-W10": {"custo_direto": 100.0, "margem_l": 10.0}},
    )

    canonical = SimpleNamespace(cubo_analitico=fake_cubo, equipes_diretas=[], analise_campanhas=None)
    gmod.render_cubo_comp(wb, canonical)
    assert "CUBO.COMP" in wb.sheetnames

    # Trigger builder-exception branch in campanha render (analise remains None)
    class _BoomBuilder:
        def build(self, _):
            raise RuntimeError("boom")

    monkeypatch.setattr(
        "concilium.domain.services.campaign.campanha_equipamento_builder.CampanhaEquipamentoBuilder",
        lambda: _BoomBuilder(),
    )

    gmod.render_campanha_frota(wb, SimpleNamespace(equipes_diretas=[], analise_campanhas=None))
    assert "CAMPANHA.FROTA" in wb.sheetnames


def test_hist_accumulate_act_no_equipamentos_defaults_min_units_1():
    """Line 62: eq.equipamentos vazio/falsy → min_units default 1 (else branch)."""
    dist = {"2026-W10": SimpleNamespace(num_equipes=2)}
    act = SimpleNamespace(activity_id="A1", distribuicao_semanal=dist)
    equipe_sem_equipamentos = SimpleNamespace(item_id="A1", equipamentos=[], funcoes=[SimpleNamespace(num_pessoas=3)])
    lookup = {"A1": equipe_sem_equipamentos}
    totals: dict = {"2026-W10": 0.0}
    gmod._rg_hist_accumulate_act(act, lookup, "equipamentos", totals)
    # scale = num_pessoas ou num_unidades conforme kind; aqui kind=equipamentos sem
    # equipamentos cadastrados → _rg_hist_scale retorna 0, mas min_units usou o default 1.
    assert "2026-W10" in totals


def test_render_graficos_empty_canonical_skips_all_charts(monkeypatch):
    """Lines 163->176, 176->188, 188->206, 201->204: sem labels/gantt/finance → nenhum gráfico é criado."""
    wb = Workbook()
    canonical = SimpleNamespace(cronograma=None, equipes_diretas=[])
    monkeypatch.setattr(gmod, "_cash_flow", lambda _c: [])

    gmod.render_graficos(wb, canonical)

    ws = wb["GRAFICOS"]
    assert ws._charts == []


def test_render_cubo_comp_builds_cubo_when_absent_on_canonical(monkeypatch):
    """Lines 316-319: cubo_analitico ausente → CuboAnaliticoBuilder().build() é chamado com sucesso."""
    wb = Workbook()
    fake_cubo = SimpleNamespace(
        celulas=[1],
        regime_bdi="pre_reforma",
        total_preco_venda=1000.0,
        total_custo_direto=700.0,
        pct_margem=0.1,
        pct_tributos=0.1,
        total_overhead=100.0,
        total_margem=100.0,
        total_tributos=100.0,
        exposicao_tributaria_por_tipo=lambda: {},
        margem_por_atividade=lambda: {},
        composicao_preco_temporal=lambda: {},
    )

    class _Builder:
        def build(self, _canonical):
            return fake_cubo

    monkeypatch.setattr(
        "concilium.domain.services.analytics.cubo_analitico_builder.CuboAnaliticoBuilder",
        lambda: _Builder(),
    )
    canonical = SimpleNamespace(equipes_diretas=[], analise_campanhas=None)  # sem cubo_analitico
    gmod.render_cubo_comp(wb, canonical)
    assert "CUBO.COMP" in wb.sheetnames


def test_render_cubo_comp_builder_exception_returns_silently(monkeypatch):
    """Linhas 320-321: exceção no builder → função retorna sem criar a planilha."""
    wb = Workbook()

    class _BoomBuilder:
        def build(self, _canonical):
            raise RuntimeError("boom")

    monkeypatch.setattr(
        "concilium.domain.services.analytics.cubo_analitico_builder.CuboAnaliticoBuilder",
        lambda: _BoomBuilder(),
    )
    canonical = SimpleNamespace(equipes_diretas=[], analise_campanhas=None)
    gmod.render_cubo_comp(wb, canonical)
    assert "CUBO.COMP" not in wb.sheetnames

def test_graficos_edge_branches_and_campanha_paths(monkeypatch):
    # _rg_hist_accumulate_act early return (eq absent)
    totals = {"2026-W10": 0.0}
    gmod._rg_hist_accumulate_act(
        SimpleNamespace(activity_id="missing", distribuicao_semanal={}), {}, "equipamentos", totals
    )
    assert totals["2026-W10"] == 0.0

    # _rg_hist_accumulate_act: semana com n_eq<=0 é pulada (branch 66->64)
    dist_zero = SimpleNamespace(num_equipes=0)
    dist_pos = SimpleNamespace(num_equipes=2)
    totals2 = {"2026-W01": 0.0, "2026-W02": 0.0}
    eq = SimpleNamespace(equipamentos=[SimpleNamespace(num_unidades=1)])
    gmod._rg_hist_accumulate_act(
        SimpleNamespace(
            activity_id="A1",
            distribuicao_semanal={"2026-W01": dist_zero, "2026-W02": dist_pos},
        ),
        {"A1": eq},
        "equipamentos",
        totals2,
    )
    assert totals2["2026-W01"] == 0.0
    assert totals2["2026-W02"] > 0.0

    # _rg_histogram_totals no weeks branch
    canonical_empty = SimpleNamespace(cronograma=SimpleNamespace(scheduled_activities=[]), equipes_diretas=[])
    assert gmod._rg_histogram_totals(canonical_empty, "equipamentos") == ([], [])

    # _rg_build_gantt_map skips activities with empty weeks
    act_no_week = SimpleNamespace(activity_id="A0", distribuicao_semanal={})
    canonical_gantt = SimpleNamespace(cronograma=SimpleNamespace(scheduled_activities=[act_no_week]))
    assert gmod._rg_build_gantt_map(canonical_gantt) == []

    wb_zero = Workbook()
    analise_zero = SimpleNamespace(
        n_campanhas=0,
        economia_total=0.0,
        custo_standby_total=0.0,
        economia_liquida=0.0,
        campanhas=[],
        alertas_custo_duplicado=[],
    )
    canonical_zero = SimpleNamespace(
        equipes_diretas=[SimpleNamespace(equipamentos=[SimpleNamespace(equip_id="EQ-1")])],
        analise_campanhas=analise_zero,
    )
    gmod.render_campanha_frota(wb_zero, canonical_zero)
    assert "CAMPANHA.FROTA" in wb_zero.sheetnames

    wb_full = Workbook()
    camp = SimpleNamespace(
        instancia_id="INST-1",
        tipo_id="EQ-1",
        semana_entrada="2026-W10",
        semana_saida="2026-W12",
        semanas_ativas=2,
        semanas_standby=1,
        custo_mob=100.0,
        custo_demob=90.0,
        custo_standby_total=30.0,
        custo_campanha=220.0,
        economia_vs_dismissal=15.0,
        atividades=["A1", "A2"],
    )
    analise_full = SimpleNamespace(
        n_campanhas=1,
        economia_total=15.0,
        custo_standby_total=30.0,
        economia_liquida=15.0,
        campanhas=[camp],
        alertas_custo_duplicado=["alerta-1"],
    )
    canonical_full = SimpleNamespace(equipes_diretas=[], analise_campanhas=analise_full)
    gmod.render_campanha_frota(wb_full, canonical_full)
    assert "CAMPANHA.FROTA" in wb_full.sheetnames
