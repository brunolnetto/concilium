"""Tests for concilium.infrastructure.adapters.excel.styles — branch coverage gaps.

Coverage targets:
  _money / _label: fill=None branch (no fill applied)
  _title_row / _section_row: n_cols <= 1 branch (no merge_cells call)
"""

from __future__ import annotations

from openpyxl import Workbook

from concilium.infrastructure.adapters.excel.styles import (
    _label,
    _money,
    _section_row,
    _title_row,
)


def _make_ws():
    wb = Workbook()
    return wb.active


class TestMoney:
    def test_without_fill_leaves_default_fill(self):
        ws = _make_ws()
        _money(ws, 1, 1, 123.456)
        cell = ws.cell(row=1, column=1)
        assert cell.value == 123.46
        # PatternFill default (no explicit fill applied) is 'none' fill type.
        assert cell.fill.fill_type is None

    def test_with_fill_applies_fill(self):
        from concilium.infrastructure.adapters.excel.styles import _fill

        ws = _make_ws()
        fill = _fill("FF0000")
        _money(ws, 1, 1, 10.0, fill=fill)
        cell = ws.cell(row=1, column=1)
        assert cell.fill.fill_type == "solid"


class TestLabel:
    def test_without_fill_leaves_default_fill(self):
        ws = _make_ws()
        _label(ws, 1, 1, "Texto")
        cell = ws.cell(row=1, column=1)
        assert cell.value == "Texto"
        assert cell.fill.fill_type is None

    def test_with_fill_applies_fill(self):
        from concilium.infrastructure.adapters.excel.styles import _fill

        ws = _make_ws()
        fill = _fill("00FF00")
        _label(ws, 1, 1, "Texto", fill=fill)
        cell = ws.cell(row=1, column=1)
        assert cell.fill.fill_type == "solid"


class TestTitleRow:
    def test_single_column_skips_merge(self):
        ws = _make_ws()
        _title_row(ws, 1, "Titulo", n_cols=1)
        assert len(ws.merged_cells.ranges) == 0
        assert ws.cell(row=1, column=1).value == "Titulo"

    def test_multi_column_merges(self):
        ws = _make_ws()
        _title_row(ws, 1, "Titulo", n_cols=3)
        assert len(ws.merged_cells.ranges) == 1


class TestSectionRow:
    def test_single_column_skips_merge(self):
        ws = _make_ws()
        _section_row(ws, 1, "Secao", n_cols=1)
        assert len(ws.merged_cells.ranges) == 0

    def test_multi_column_merges(self):
        ws = _make_ws()
        _section_row(ws, 1, "Secao", n_cols=3)
        assert len(ws.merged_cells.ranges) == 1
