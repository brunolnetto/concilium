"""Greatest-gap coverage batch 1.

Prioritizes lowest-coverage modules first:
- infrastructure/adapters/excel/sheets/graficos.py
- infrastructure/adapters/blob_storage/sharepoint.py
- infrastructure/catalog/query_engine.py
"""

from __future__ import annotations

import json
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from concilium.infrastructure.adapters.blob_storage.sharepoint import SharePointBlobStorage, SharePointCredential
from concilium.infrastructure.adapters.excel.sheets import graficos as gmod
from concilium.infrastructure.catalog.query_engine import CatalogQueryEngine

# ---------------------------------------------------------------------------
# SharePoint adapter
# ---------------------------------------------------------------------------


class _Resp:
    def __init__(self, status_code=200, content=b"", headers=None, data=None):
        self.status_code = status_code
        self.content = content
        self.headers = headers or {}
        self._data = data or {}

    def json(self):
        return self._data

    def raise_for_status(self):
        if self.status_code >= 400 and self.status_code not in (404, 412):
            raise RuntimeError(f"http {self.status_code}")


@pytest.fixture
def sp_cred() -> SharePointCredential:
    cred = SharePointCredential(tenant_id="t", client_id="c", client_secret="s")
    cred._app = MagicMock()
    cred._app.acquire_token_for_client.return_value = {"access_token": "tok"}
    return cred


def test_sharepoint_from_env_missing_drive_id_raises(monkeypatch):
    class _SP:
        drive_id = None
        root_path = "root"

    class _Storage:
        sharepoint = _SP()

    monkeypatch.setattr("concilium.infrastructure.adapters.blob_storage.sharepoint.get_storage", lambda: _Storage())
    with pytest.raises(OSError):
        SharePointBlobStorage.from_env()


def test_sharepoint_cached_etag_refresh_path(tmp_path, monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred, cache_dir=tmp_path)
    storage._cache.put("a.json", b"old", "etag0", "lm0")

    # First GET (item metadata) returns 200; second GET (content) returns 200 with new payload.
    httpx = SimpleNamespace(
        get=MagicMock(
            side_effect=[
                _Resp(status_code=200, headers={"ETag": "etag1", "Last-Modified": "lm1"}),
                _Resp(status_code=200, content=b"new"),
            ]
        )
    )
    monkeypatch.setattr(storage, "_client", lambda: httpx)

    data = storage.ler_blob("a.json")
    assert data == b"new"


def test_sharepoint_cached_exception_offline_returns_stale(tmp_path, monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred, cache_dir=tmp_path, offline=True)
    storage._cache.put("b.json", b"stale", "etag0", "lm0")
    httpx = SimpleNamespace(get=MagicMock(side_effect=RuntimeError("network down")))
    monkeypatch.setattr(storage, "_client", lambda: httpx)

    assert storage.ler_blob("b.json") == b"stale"


def test_sharepoint_direct_offline_no_cache_raises(tmp_path, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred, cache_dir=tmp_path, offline=True)
    with pytest.raises(FileNotFoundError, match="offline"):
        storage._ler_blob_direct("missing.json", httpx=SimpleNamespace())


def test_sharepoint_direct_request_exception_wrapped(monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred)
    httpx = SimpleNamespace(get=MagicMock(side_effect=RuntimeError("boom")))
    with pytest.raises(FileNotFoundError, match="Erro ao acessar SharePoint"):
        storage._ler_blob_direct("x.json", httpx=httpx)


def test_sharepoint_gravar_blob_large_upload_and_cache_invalidate(tmp_path, monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred, cache_dir=tmp_path)
    storage._cache.put("big.bin", b"old", "e", "lm")

    called = {"upload": False, "invalidated": False}

    def _upload(_httpx, path, content):
        called["upload"] = path == "big.bin" and len(content) > 0

    def _invalidate(path):
        called["invalidated"] = path == "big.bin"

    monkeypatch.setattr(storage, "_gravar_upload_session", _upload)
    monkeypatch.setattr(storage._cache, "invalidate", _invalidate)
    monkeypatch.setattr(storage, "_client", lambda: SimpleNamespace())

    storage.gravar_blob("big.bin", b"x" * (4 * 1024 * 1024))
    assert called["upload"] is True
    assert called["invalidated"] is True


def test_sharepoint_upload_session_chunk_loop(monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred)
    post_resp = _Resp(status_code=200, data={"uploadUrl": "https://upload"})
    put_resp = _Resp(status_code=200)
    httpx = SimpleNamespace(post=MagicMock(return_value=post_resp), put=MagicMock(return_value=put_resp))

    data = b"y" * (5 * 1024 * 1024 + 10)
    storage._gravar_upload_session(httpx=httpx, path="file.bin", content=data)
    # 2 chunks expected: one full 5MB + remainder
    assert httpx.put.call_count == 2


def test_sharepoint_listar_blobs_offline_exception_returns_empty(tmp_path, monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred, cache_dir=tmp_path, offline=True)
    httpx = SimpleNamespace(get=MagicMock(side_effect=RuntimeError("offline")))
    monkeypatch.setattr(storage, "_client", lambda: httpx)

    assert storage.listar_blobs("abc") == []


def test_sharepoint_listar_blobs_404_returns_empty(monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred)
    httpx = SimpleNamespace(get=MagicMock(return_value=_Resp(status_code=404)))
    monkeypatch.setattr(storage, "_client", lambda: httpx)

    assert storage.listar_blobs("abc") == []


def test_sharepoint_listar_blobs_recurses_into_subfolder(monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred)
    r1 = _Resp(status_code=200, data={"value": [{"name": "sub", "folder": {}}, {"name": "f1.json", "file": {}}]})
    r2 = _Resp(status_code=200, data={"value": [{"name": "f2.json", "file": {}}]})
    httpx = SimpleNamespace(get=MagicMock(side_effect=[r1, r2]))
    monkeypatch.setattr(storage, "_client", lambda: httpx)

    blobs = storage.listar_blobs("prefix")
    assert "prefix/f1.json" in blobs
    assert "prefix/sub/f2.json" in blobs


def test_sharepoint_blob_existe_exception_false(monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred)
    httpx = SimpleNamespace(get=MagicMock(side_effect=RuntimeError("x")))
    monkeypatch.setattr(storage, "_client", lambda: httpx)

    assert storage.blob_existe("any") is False


def test_sharepoint_aprovar_versao_404_raises(monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred)
    httpx = SimpleNamespace(post=MagicMock(return_value=_Resp(status_code=404)))
    monkeypatch.setattr(storage, "_client", lambda: httpx)

    with pytest.raises(FileNotFoundError):
        storage.aprovar_versao("x")


def test_sharepoint_gerar_url_temporaria_non_200_and_exception(monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred)

    httpx_non200 = SimpleNamespace(post=MagicMock(return_value=_Resp(status_code=403)))
    monkeypatch.setattr(storage, "_client", lambda: httpx_non200)
    assert storage.gerar_url_temporaria("x") is None

    httpx_exc = SimpleNamespace(post=MagicMock(side_effect=RuntimeError("blocked")))
    monkeypatch.setattr(storage, "_client", lambda: httpx_exc)
    assert storage.gerar_url_temporaria("x") is None


def test_sharepoint_health_check_non_200_error_code(monkeypatch, sp_cred):
    storage = SharePointBlobStorage("d", "root", sp_cred)
    httpx = SimpleNamespace(get=MagicMock(return_value=_Resp(status_code=503)))
    monkeypatch.setattr(storage, "_client", lambda: httpx)

    result = storage.health_check()
    assert result["status"] == "error"
    assert result["code"] == 503


# ---------------------------------------------------------------------------
# Graficos sheet
# ---------------------------------------------------------------------------


def _canonical_for_graficos():
    dist = {
        "2026-W10": SimpleNamespace(num_equipes=2),
        "2026-W11": SimpleNamespace(num_equipes=1),
    }
    act = SimpleNamespace(activity_id="A1", distribuicao_semanal=dist)
    cronograma = SimpleNamespace(scheduled_activities=[act])
    equipe = SimpleNamespace(
        item_id="A1",
        equipamentos=[SimpleNamespace(num_unidades=2)],
        funcoes=[SimpleNamespace(num_pessoas=3)],
    )
    return SimpleNamespace(cronograma=cronograma, equipes_diretas=[equipe])


def test_graficos_helpers_and_render(monkeypatch):
    canonical = _canonical_for_graficos()
    wb = Workbook()

    monkeypatch.setattr(
        gmod,
        "_cash_flow",
        lambda _c: [
            {
                "mes_label": "2026-03",
                "receita": 100.0,
                "despesa": 80.0,
                "saldo": 20.0,
                "exposicao_acumulada": 20.0,
            }
        ],
    )

    assert gmod._rg_week_label("2026-W10")
    labels, values = gmod._rg_histogram_totals(canonical, "equipamentos")
    assert labels and values
    assert gmod._rg_cash_flow_rows(canonical)[0]["mes"] == "2026-03"

    gmod.render_graficos(wb, canonical)
    assert "GRAFICOS" in wb.sheetnames


def test_render_cubo_comp_and_campanha_frota_branches(monkeypatch):
    wb = Workbook()

    fake_cubo = SimpleNamespace(
        celulas=[1],
        regime_bdi="pre_reforma",
        total_preco_venda=1000.0,
        total_custo_direto=700.0,
        pct_margem=0.1,
        pct_tributos=0.1,
        total_overhead=100.0,
        total_margem=100.0,
        total_tributos=100.0,
        exposicao_tributaria_por_tipo=lambda: {"tributo_iss": 50.0},
        margem_por_atividade=lambda: {"A1": 100.0},
        composicao_preco_temporal=lambda: {"2026-W10": {"custo_direto": 100.0, "margem_l": 10.0}},
    )

    canonical = SimpleNamespace(cubo_analitico=fake_cubo, equipes_diretas=[], analise_campanhas=None)
    gmod.render_cubo_comp(wb, canonical)
    assert "CUBO.COMP" in wb.sheetnames

    # Trigger builder-exception branch in campanha render (analise remains None)
    class _BoomBuilder:
        def build(self, _):
            raise RuntimeError("boom")

    monkeypatch.setattr(
        "concilium.domain.services.campaign.campanha_equipamento_builder.CampanhaEquipamentoBuilder",
        lambda: _BoomBuilder(),
    )

    gmod.render_campanha_frota(wb, SimpleNamespace(equipes_diretas=[], analise_campanhas=None))
    assert "CAMPANHA.FROTA" in wb.sheetnames


# ---------------------------------------------------------------------------
# Query engine internals
# ---------------------------------------------------------------------------


class _FakeResult:
    def __init__(self, rows):
        self._rows = rows

    def fetchall(self):
        return self._rows


class _FakeConn:
    def __init__(self):
        self.executed = []
        self.executed_many = []

    def execute(self, sql, *args):
        self.executed.append((sql, args))
        if "SELECT raw_json FROM cpus" in sql:
            cpu = {
                "codigo": "CPU-1",
                "composicao_itens": [
                    {"insumo_id": "I1", "coeficiente": 2.0, "categoria_custo": "mat", "unidade": "kg"}
                ],
            }
            return _FakeResult([(json.dumps(cpu),)])
        if "SELECT raw_json FROM produtividades" in sql:
            prod = {
                "id": "P1",
                "composicao_padrao_equipe": {
                    "funcoes": [{"funcao_id": "MO-1", "num_pessoas": 2, "papel": "direto"}],
                    "equipamentos": [{"equip_id": "EQ-1", "num_unidades": 1, "consumo_litros_hora": 3.0}],
                    "materiais": [{"insumo_id": "MAT-1", "consumo_por_hora": 4.0, "unidade": "kg"}],
                },
            }
            return _FakeResult([(json.dumps(prod),)])
        return _FakeResult([])

    def executemany(self, sql, rows):
        self.executed_many.append((sql, list(rows)))


class _DuckConn:
    def __init__(self, rows):
        self.rows = rows
        self.closed = False

    def execute(self, _sql):
        return _FakeResult(self.rows)

    def close(self):
        self.closed = True


def test_query_engine_build_junction_tables_internal(tmp_path):
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()
    engine._build_junction_tables(conn, cpus_rows=[], prods_rows=[])

    assert any("CREATE TABLE composicao_itens" in sql for sql, _ in conn.executed)
    assert any("INSERT INTO composicao_itens" in sql for sql, _ in conn.executed_many)
    assert any("INSERT INTO equipe_funcoes" in sql for sql, _ in conn.executed_many)


def test_query_engine_build_methods_tolerate_invalid_json(tmp_path):
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()

    cpus_dir = tmp_path / "cpus"
    cpus_dir.mkdir(parents=True)
    (cpus_dir / "bad.json").write_text("{invalid")
    (cpus_dir / "list.json").write_text("[]")
    engine._build_cpus(conn)  # exception + non-dict branches

    ins_dir = tmp_path / "insumos"
    ins_dir.mkdir(exist_ok=True)
    (ins_dir / "materiais.json").write_text("{invalid")
    engine._build_insumos(conn)

    (tmp_path / "produtividades.json").write_text("{invalid")
    engine._build_produtividades(conn)

    (tmp_path / "funcoes.json").write_text("{invalid")
    engine._build_funcoes(conn)

    assert any("CREATE TABLE cpus" in sql for sql, _ in conn.executed)
    assert any("CREATE TABLE insumos" in sql for sql, _ in conn.executed)
    assert any("CREATE TABLE produtividades" in sql for sql, _ in conn.executed)
    assert any("CREATE TABLE funcoes" in sql for sql, _ in conn.executed)


def test_query_engine_load_table_branches(monkeypatch, tmp_path):
    engine = CatalogQueryEngine(tmp_path)

    # extra_col branch
    dc1 = _DuckConn(rows=[("tipo-a", json.dumps({"id": "x"}))])

    class _Duck1:
        @staticmethod
        def connect(_path, read_only=True):
            return dc1

    monkeypatch.setitem(__import__("sys").modules, "duckdb", _Duck1)
    rows = engine._load_table("insumos", extra_col="tipo")
    assert rows[0]["tipo"] == "tipo-a"
    assert dc1.closed is True

    # non-extra_col branch
    dc2 = _DuckConn(rows=[(json.dumps({"codigo": "CPU-1"}),)])

    class _Duck2:
        @staticmethod
        def connect(_path, read_only=True):
            return dc2

    monkeypatch.setitem(__import__("sys").modules, "duckdb", _Duck2)
    rows2 = engine._load_table("cpus")
    assert rows2[0]["codigo"] == "CPU-1"
    assert dc2.closed is True
