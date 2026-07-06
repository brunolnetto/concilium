"""Greatest-gap coverage batch 2.

Targets remaining high-impact branches after batch 1, prioritizing:
- infrastructure/catalog/query_engine.py
- infrastructure/adapters/blob_storage/sharepoint.py
- infrastructure/adapters/excel/sheets/graficos.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from concilium.infrastructure.adapters.blob_storage.sharepoint import SharePointBlobStorage, SharePointCredential
from concilium.infrastructure.adapters.excel.sheets import graficos as gmod
from concilium.infrastructure.catalog.query_engine import CatalogQueryEngine, CatalogSession


class _FakeResult:
    def __init__(self, rows):
        self._rows = rows

    def fetchall(self):
        return self._rows


class _FakeConn:
    def __init__(self):
        self.executed: list[str] = []
        self.executed_many: list[tuple[str, list[tuple]]] = []
        self.closed = False
        self.checkpoint_called = False

    def execute(self, sql, *args):
        self.executed.append(str(sql))
        return _FakeResult([])

    def executemany(self, sql, rows):
        self.executed_many.append((str(sql), list(rows)))

    def checkpoint(self):
        self.checkpoint_called = True

    def close(self):
        self.closed = True


def _sp_cred() -> SharePointCredential:
    cred = SharePointCredential(tenant_id="t", client_id="c", client_secret="s")
    cred._app = MagicMock()
    cred._app.acquire_token_for_client.return_value = {"access_token": "tok"}
    return cred


class _Resp:
    def __init__(self, status_code=200, content=b"", headers=None, data=None):
        self.status_code = status_code
        self.content = content
        self.headers = headers or {}
        self._data = data or {}

    def json(self):
        return self._data

    def raise_for_status(self):
        if self.status_code >= 400 and self.status_code not in (404, 412):
            raise RuntimeError(f"http {self.status_code}")


def test_query_engine_rebuild_handles_stale_db_and_writes_timestamp(tmp_path, monkeypatch):
    engine = CatalogQueryEngine(tmp_path)
    # Create directory at DB path so unlink raises OSError and exercises the safe fallback.
    engine._db_path.mkdir(parents=True, exist_ok=True)

    conn = _FakeConn()

    class _Duckdb:
        @staticmethod
        def connect(_path):
            return conn

    monkeypatch.setitem(sys.modules, "duckdb", _Duckdb)
    monkeypatch.setattr(engine, "_build_cpus", lambda c: c.execute("CREATE CPUS"))
    monkeypatch.setattr(engine, "_build_insumos", lambda c: c.execute("CREATE INSUMOS"))
    monkeypatch.setattr(engine, "_build_produtividades", lambda c: c.execute("CREATE PRODS"))
    monkeypatch.setattr(engine, "_build_funcoes", lambda c: c.execute("CREATE FUNCOES"))

    engine._rebuild()

    assert conn.checkpoint_called is True
    assert conn.closed is True
    assert engine._ts_path.exists()
    payload = json.loads(engine._ts_path.read_text(encoding="utf-8"))
    assert payload["schema_ver"] == CatalogQueryEngine.SCHEMA_VER
    assert "rebuild_ms" in payload


def test_query_engine_build_methods_cover_layout_variants(tmp_path):
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()

    # cpus.json fallback layout
    (tmp_path / "cpus.json").write_text(
        json.dumps({"cpus": [{"codigo": "CPU-1", "descricao": "d", "categoria": "c", "unidade": "u"}]}),
        encoding="utf-8",
    )

    # insumos/<tipo>/<ID>.json layout with subdirectory
    tipo_dir = tmp_path / "insumos" / "equipamento"
    tipo_dir.mkdir(parents=True)
    (tipo_dir / "EQ-1.json").write_text(json.dumps({"id": "EQ-1", "descricao": "Escavadeira"}), encoding="utf-8")

    # produtividades from alternate folder name "producoes"
    prods_dir = tmp_path / "producoes"
    prods_dir.mkdir(parents=True)
    (prods_dir / "PROD-1.json").write_text(json.dumps({"id": "PROD-1", "cpu_id": "CPU-1"}), encoding="utf-8")

    # funcoes from funcoes.json
    (tmp_path / "funcoes.json").write_text(
        json.dumps({"funcoes": [{"id": "FN-1", "descricao": "Op"}]}), encoding="utf-8"
    )

    engine._build_cpus(conn)
    engine._build_insumos(conn)
    engine._build_produtividades(conn)
    engine._build_funcoes(conn)

    inserts = "\n".join(sql for sql, _ in conn.executed_many)
    assert "INSERT INTO cpus" in inserts
    assert "INSERT INTO insumos" in inserts
    assert "INSERT INTO produtividades" in inserts
    assert "INSERT INTO funcoes" in inserts


def test_query_engine_wrapper_load_methods(tmp_path, monkeypatch):
    engine = CatalogQueryEngine(tmp_path)

    def _fake_load(table, extra_col=None):
        if table == "cpus":
            return [{"codigo": "CPU-1"}]
        if table == "insumos":
            return [{"tipo": "equipamento", "id": "EQ-1"}, {"id": "GEN-1"}]
        if table == "produtividades":
            return [{"id": "PROD-1"}]
        return [{"id": "FN-1"}]

    monkeypatch.setattr(engine, "_load_table", _fake_load)

    assert engine.load_cpus()[0]["codigo"] == "CPU-1"
    grouped = engine.load_insumos()
    assert grouped["equipamento"][0]["id"] == "EQ-1"
    assert grouped["outros"][0]["id"] == "GEN-1"
    assert engine.load_produtividades()[0]["id"] == "PROD-1"
    assert engine.load_funcoes()[0]["id"] == "FN-1"


def test_query_engine_is_ready_returns_false_when_cache_file_missing(tmp_path, monkeypatch):
    from concilium.infrastructure.catalog import query_engine

    engine = CatalogQueryEngine(tmp_path)
    monkeypatch.setattr(engine, "_is_stale", lambda: False)
    monkeypatch.setattr(query_engine, "_AVAILABLE", True)
    assert engine.is_ready() is False


def test_query_engine_is_stale_ignores_oserror_in_json_stat(tmp_path, monkeypatch):
    engine = CatalogQueryEngine(tmp_path)
    engine._db_path.touch()
    engine._ts_path.write_text(
        json.dumps({"schema_ver": CatalogQueryEngine.SCHEMA_VER, "mtime": 1e20}),
        encoding="utf-8",
    )
    noisy_json = tmp_path / "broken.json"
    noisy_json.write_text("{}", encoding="utf-8")

    original_stat = Path.stat

    def _stat(self, *args, **kwargs):
        if self == noisy_json:
            raise OSError("stat fail")
        return original_stat(self, *args, **kwargs)

    monkeypatch.setattr(Path, "stat", _stat)
    assert engine._is_stale() is False


def test_query_engine_build_cpus_invalid_entries_and_bad_payload(tmp_path):
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()

    cpus_dir = tmp_path / "cpus"
    cpus_dir.mkdir(parents=True)
    # Valid JSON but not dict -> no append branch.
    (cpus_dir / "arr.json").write_text("[]", encoding="utf-8")
    # Invalid JSON -> exception branch.
    (cpus_dir / "bad.json").write_text("{", encoding="utf-8")

    engine._build_cpus(conn)
    assert any("CREATE TABLE cpus" in sql for sql in conn.executed)

    # Remove dir to exercise cpus.json fallback exception path.
    for f in cpus_dir.glob("*.json"):
        f.unlink()
    cpus_dir.rmdir()
    (tmp_path / "cpus.json").write_text("{", encoding="utf-8")
    engine._build_cpus(conn)
    assert any("CREATE INDEX idx_cpus_codigo" in sql for sql in conn.executed)


def test_query_engine_build_insumos_edge_paths(tmp_path):
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()

    # Layout A with one non-directory entry and one bad json file.
    ins = tmp_path / "insumos"
    ins.mkdir(parents=True)
    (ins / "README.txt").write_text("x", encoding="utf-8")
    tipo = ins / "equipamento"
    tipo.mkdir()
    (tipo / "bad.json").write_text("{", encoding="utf-8")
    engine._build_insumos(conn)

    # Layout B with mixed list entries (non-dict + dict).
    for f in tipo.glob("*.json"):
        f.unlink()
    tipo.rmdir()
    (ins / "equipamento.json").write_text(json.dumps({"insumos": ["bad", {"id": "EQ-1"}]}), encoding="utf-8")
    engine._build_insumos(conn)

    assert any("CREATE TABLE insumos" in sql for sql in conn.executed)


def test_query_engine_build_produtividades_edge_paths(tmp_path):
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()

    prod_dir = tmp_path / "produtividades"
    prod_dir.mkdir(parents=True)
    (prod_dir / "bad.json").write_text("{", encoding="utf-8")
    engine._build_produtividades(conn)

    # Fallback single-file path with mixed entries.
    (tmp_path / "produtividades.json").write_text(
        json.dumps({"produtividades": ["bad", {"id": "P1", "cpu_id": "CPU-1"}]}),
        encoding="utf-8",
    )
    # No rows from directory branch after removing the folder.
    for f in prod_dir.glob("*.json"):
        f.unlink()
    prod_dir.rmdir()
    engine._build_produtividades(conn)

    assert any("CREATE TABLE produtividades" in sql for sql in conn.executed)


def test_query_engine_build_cpus_neither_dir_nor_file_exists(tmp_path):
    """Neither cpus/ nor cpus.json present: rows stays empty, both branches of
    the if/elif are skipped entirely (line 257->267 arc)."""
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()
    engine._build_cpus(conn)
    assert not any("INSERT INTO cpus" in sql for sql, _ in conn.executed_many)
    assert any("CREATE TABLE cpus" in sql for sql in conn.executed)


def test_query_engine_build_cpus_file_fallback_multiple_entries(tmp_path):
    """cpus.json fallback list with 2+ entries exercises the for-loop back-edge."""
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()
    (tmp_path / "cpus.json").write_text(
        json.dumps(
            {
                "cpus": [
                    {"codigo": "CPU-1", "descricao": "d1", "categoria": "c1", "unidade": "u1"},
                    {"codigo": "CPU-2", "descricao": "d2", "categoria": "c2", "unidade": "u2"},
                ]
            }
        ),
        encoding="utf-8",
    )
    engine._build_cpus(conn)
    ((_, rows),) = [pair for pair in conn.executed_many if "INSERT INTO cpus" in pair[0]]
    assert len(rows) == 2


def test_query_engine_build_insumos_multiple_files_in_same_tipo_dir(tmp_path):
    """Two JSON files in the same insumos/<tipo>/ dir exercise the glob loop back-edge."""
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()
    tipo_dir = tmp_path / "insumos" / "equipamento"
    tipo_dir.mkdir(parents=True)
    (tipo_dir / "EQ-1.json").write_text(json.dumps({"id": "EQ-1", "descricao": "A"}), encoding="utf-8")
    (tipo_dir / "EQ-2.json").write_text(json.dumps({"id": "EQ-2", "descricao": "B"}), encoding="utf-8")
    engine._build_insumos(conn)
    ((_, rows),) = [pair for pair in conn.executed_many if "INSERT INTO insumos" in pair[0]]
    assert len(rows) == 2


def test_query_engine_build_produtividades_multiple_files_in_dir(tmp_path):
    """Two JSON files in produtividades/ exercise the glob loop back-edge."""
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()
    prod_dir = tmp_path / "produtividades"
    prod_dir.mkdir(parents=True)
    (prod_dir / "P1.json").write_text(json.dumps({"id": "P1", "cpu_id": "CPU-1"}), encoding="utf-8")
    (prod_dir / "P2.json").write_text(json.dumps({"id": "P2", "cpu_id": "CPU-2"}), encoding="utf-8")
    engine._build_produtividades(conn)
    ((_, rows),) = [pair for pair in conn.executed_many if "INSERT INTO produtividades" in pair[0]]
    assert len(rows) == 2


def test_query_engine_build_funcoes_file_multiple_entries(tmp_path):
    """funcoes.json list with 2+ entries exercises the for-loop back-edge."""
    engine = CatalogQueryEngine(tmp_path)
    conn = _FakeConn()
    (tmp_path / "funcoes.json").write_text(
        json.dumps({"funcoes": [{"id": "FN-1", "descricao": "A"}, {"id": "FN-2", "descricao": "B"}]}),
        encoding="utf-8",
    )
    engine._build_funcoes(conn)
    ((_, rows),) = [pair for pair in conn.executed_many if "INSERT INTO funcoes" in pair[0]]
    assert len(rows) == 2


def test_catalog_session_load_material_id_fallback_branch():
    catalogo = {
        "cpus": [
            {
                "codigo": "CPU-1",
                "descricao": "Esc",
                "categoria": "terraplenagem",
                "unidade": "m3",
                "composicao_itens": [],
            }
        ],
        "insumos": {"material": [{"id": "MAT-1", "preco_unitario": 5.0}]},
        "produtividades": [
            {
                "id": "PROD-1",
                "cpu_id": "CPU-1",
                "composicao_padrao_equipe": {
                    "funcoes": [],
                    "equipamentos": [],
                    "materiais": [{"material_id": "MAT-1", "consumo_por_hora": 2.0, "unidade": "kg"}],
                },
            }
        ],
    }
    with CatalogSession.from_resolved(catalogo) as session:
        equipe = session.get_equipe_com_precos("PROD-1")
    assert equipe["materiais"][0]["insumo_id"] == "MAT-1"


def test_sharepoint_success_and_error_branches(tmp_path, monkeypatch):
    cred = _sp_cred()
    # Garantir que variáveis de env de SP não poluem o teste via settings cacheados
    monkeypatch.delenv("CONCILIUM__SP_DRIVE_ID", raising=False)
    monkeypatch.delenv("CONCILIUM__SP_CLIENT_SECRET", raising=False)

    class _SP:
        drive_id = "drv"
        root_path = "root"

    class _Storage:
        sharepoint = _SP()

    monkeypatch.setattr("concilium.infrastructure.adapters.blob_storage.sharepoint.get_storage", lambda: _Storage())
    monkeypatch.setattr(
        "concilium.infrastructure.adapters.blob_storage.sharepoint.SharePointCredential.from_env",
        lambda: cred,
    )

    storage = SharePointBlobStorage.from_env(cache_dir=tmp_path)
    assert isinstance(storage, SharePointBlobStorage)

    # _ler_blob_direct offline branch using stale cache
    storage._cache.put("offline.json", b"cached", "e", "lm")
    storage._offline = True
    assert storage._ler_blob_direct("offline.json", httpx=SimpleNamespace()) == b"cached"
    storage._offline = False

    # Small write path and optimistic concurrency 412 branch
    httpx_put_412 = SimpleNamespace(put=MagicMock(return_value=_Resp(status_code=412)))
    with pytest.raises(Exception):  # noqa: B017
        storage._gravar_pequeno(httpx_put_412, "x.bin", b"abc")

    # listar_blobs exception while online should raise wrapped FileNotFoundError
    httpx_list_err = SimpleNamespace(get=MagicMock(side_effect=RuntimeError("net")))
    monkeypatch.setattr(storage, "_client", lambda: httpx_list_err)
    with pytest.raises(FileNotFoundError):
        storage.listar_blobs("pref")

    # Success branches: approve, temporary URL, health_check=ok
    httpx_ok = SimpleNamespace(
        post=MagicMock(side_effect=[_Resp(status_code=200), _Resp(status_code=200, data={"link": {"webUrl": "u"}})]),
        get=MagicMock(return_value=_Resp(status_code=200, data={"name": "Drive", "driveType": "documentLibrary"})),
    )
    monkeypatch.setattr(storage, "_client", lambda: httpx_ok)
    storage.aprovar_versao("ok.txt")
    assert storage.gerar_url_temporaria("ok.txt") == "u"
    health = storage.health_check()
    assert health["status"] == "ok"
    assert health["drive_name"] == "Drive"


def test_graficos_edge_branches_and_campanha_paths(monkeypatch):
    # _rg_hist_accumulate_act early return (eq absent)
    totals = {"2026-W10": 0.0}
    gmod._rg_hist_accumulate_act(
        SimpleNamespace(activity_id="missing", distribuicao_semanal={}), {}, "equipamentos", totals
    )
    assert totals["2026-W10"] == 0.0

    # _rg_hist_accumulate_act: semana com n_eq<=0 é pulada (branch 66->64)
    dist_zero = SimpleNamespace(num_equipes=0)
    dist_pos = SimpleNamespace(num_equipes=2)
    totals2 = {"2026-W01": 0.0, "2026-W02": 0.0}
    eq = SimpleNamespace(equipamentos=[SimpleNamespace(num_unidades=1)])
    gmod._rg_hist_accumulate_act(
        SimpleNamespace(
            activity_id="A1",
            distribuicao_semanal={"2026-W01": dist_zero, "2026-W02": dist_pos},
        ),
        {"A1": eq},
        "equipamentos",
        totals2,
    )
    assert totals2["2026-W01"] == 0.0
    assert totals2["2026-W02"] > 0.0

    # _rg_histogram_totals no weeks branch
    canonical_empty = SimpleNamespace(cronograma=SimpleNamespace(scheduled_activities=[]), equipes_diretas=[])
    assert gmod._rg_histogram_totals(canonical_empty, "equipamentos") == ([], [])

    # _rg_build_gantt_map skips activities with empty weeks
    act_no_week = SimpleNamespace(activity_id="A0", distribuicao_semanal={})
    canonical_gantt = SimpleNamespace(cronograma=SimpleNamespace(scheduled_activities=[act_no_week]))
    assert gmod._rg_build_gantt_map(canonical_gantt) == []

    wb_zero = Workbook()
    analise_zero = SimpleNamespace(
        n_campanhas=0,
        economia_total=0.0,
        custo_standby_total=0.0,
        economia_liquida=0.0,
        campanhas=[],
        alertas_custo_duplicado=[],
    )
    canonical_zero = SimpleNamespace(
        equipes_diretas=[SimpleNamespace(equipamentos=[SimpleNamespace(equip_id="EQ-1")])],
        analise_campanhas=analise_zero,
    )
    gmod.render_campanha_frota(wb_zero, canonical_zero)
    assert "CAMPANHA.FROTA" in wb_zero.sheetnames

    wb_full = Workbook()
    camp = SimpleNamespace(
        instancia_id="INST-1",
        tipo_id="EQ-1",
        semana_entrada="2026-W10",
        semana_saida="2026-W12",
        semanas_ativas=2,
        semanas_standby=1,
        custo_mob=100.0,
        custo_demob=90.0,
        custo_standby_total=30.0,
        custo_campanha=220.0,
        economia_vs_dismissal=15.0,
        atividades=["A1", "A2"],
    )
    analise_full = SimpleNamespace(
        n_campanhas=1,
        economia_total=15.0,
        custo_standby_total=30.0,
        economia_liquida=15.0,
        campanhas=[camp],
        alertas_custo_duplicado=["alerta-1"],
    )
    canonical_full = SimpleNamespace(equipes_diretas=[], analise_campanhas=analise_full)
    gmod.render_campanha_frota(wb_full, canonical_full)
    assert "CAMPANHA.FROTA" in wb_full.sheetnames
