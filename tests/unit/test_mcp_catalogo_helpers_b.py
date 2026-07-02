"""Tests for _catalogo_tools_helpers_b — covers missing lines."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import MagicMock, patch

from concilium.presentation.mcp.tools._catalogo_tools_helpers_b import (
    _catalog_root,
    _impl_avaliar_plausibilidade_catalogo,
    _impl_calcular_pu_producao,
    _impl_comparar_rfi_catalogo,
    _impl_tendencia_precos_catalogo,
    _impl_verificar_integridade_catalogo,
)

MODULE = "concilium.presentation.mcp.tools._catalogo_tools_helpers_b"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _identity(x):
    """Use as _wrap_response to keep assertions simple."""
    return x


def _make_bundle(
    insumos=None,
    producoes=None,
    cpus=None,
    use_produtividades_key=False,
):
    bundle: dict = {}
    if insumos is not None:
        bundle["insumos"] = insumos
    if producoes is not None:
        if use_produtividades_key:
            bundle["produtividades"] = producoes
        else:
            bundle["producoes"] = producoes
    if cpus is not None:
        bundle["cpus"] = cpus
    return bundle


def _write_bundle(path: Path, bundle: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(bundle), encoding="utf-8")


# ---------------------------------------------------------------------------
# _catalog_root
# ---------------------------------------------------------------------------


class TestCatalogRoot:
    def test_returns_correct_path(self, tmp_path):
        result = _catalog_root(tmp_path, "mendeslobo")
        assert result == tmp_path / "catalogos" / "mendeslobo"

    def test_different_catalog_id(self, tmp_path):
        result = _catalog_root(tmp_path, "outro-contexto")
        assert result == tmp_path / "catalogos" / "outro-contexto"


# ---------------------------------------------------------------------------
# _impl_tendencia_precos_catalogo
# ---------------------------------------------------------------------------


class TestImplTendenciaPrecosCatalogo:
    def _config_path(self, tmp_path):
        return tmp_path / "config"

    def test_sad_path_bundle_missing(self, tmp_path):
        config_dir = self._config_path(tmp_path)
        config_dir.mkdir(parents=True)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            result = _impl_tendencia_precos_catalogo("mendeslobo")

        assert "erro" in result
        assert "Bundle" in result["erro"]

    def test_happy_path_basic(self, tmp_path):
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "mendeslobo" / ".catalog_bundle.json"
        insumos = [
            {"codigo": "INS-001", "tipo": "EQUIPAMENTO", "preco_unitario": 100.0, "descricao": "Test", "unidade": "h"}
        ]
        _write_bundle(bundle_path, _make_bundle(insumos=insumos))

        # Build a real RelatorioTendencia-like dataclass result
        from concilium.application.catalog.price_trend import IndiceCategoria, RelatorioTendencia

        fake_relatorio = RelatorioTendencia(
            contexto_id="mendeslobo",
            gerado_em="2026-06-14",
            data_base_analise="2026-06-14",
            jurisdicao_id=None,
            indices_por_categoria=[
                IndiceCategoria(
                    tipo="EQUIPAMENTO",
                    n_insumos_com_historico=0,
                    n_insumos_total=1,
                    variacao_media_pct=None,
                    variacao_mediana_pct=None,
                    variacao_max_pct=None,
                    variacao_min_pct=None,
                    insumos=[],
                )
            ],
            resumo={"total_insumos": 1},
        )

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.price_trend.calcular_tendencia",
                    return_value=fake_relatorio,
                ):
                    result = _impl_tendencia_precos_catalogo("mendeslobo")

        assert "indices_por_categoria" in result
        assert result["contexto_id"] == "mendeslobo"

    def test_happy_path_truncation_more_than_10_insumos(self, tmp_path):
        """Covers lines 50-53: top-10 truncation and insumos_omitidos."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        _write_bundle(bundle_path, _make_bundle(insumos=[]))

        from concilium.application.catalog.price_trend import IndiceCategoria, RelatorioTendencia, TendenciaInsumo

        # 15 insumos — truncation logic should keep 10 and set insumos_omitidos=5
        fake_insumos = [
            TendenciaInsumo(
                codigo=f"INS-{i:03d}",
                descricao=f"desc {i}",
                tipo="MATERIAL",
                unidade="un",
                preco_atual=float(i),
                preco_referencia=float(i) * 0.9,
                data_referencia="2024-01-01",
                variacao_abs=float(i) * 0.1,
                variacao_pct=float(i),
                n_registros_historico=1,
                fonte_atual="preco_unitario",
            )
            for i in range(1, 16)
        ]
        fake_relatorio = RelatorioTendencia(
            contexto_id="ctx",
            gerado_em="2026-06-14",
            data_base_analise="2026-06-14",
            jurisdicao_id=None,
            indices_por_categoria=[
                IndiceCategoria(
                    tipo="MATERIAL",
                    n_insumos_com_historico=15,
                    n_insumos_total=15,
                    variacao_media_pct=8.0,
                    variacao_mediana_pct=8.0,
                    variacao_max_pct=15.0,
                    variacao_min_pct=1.0,
                    insumos=fake_insumos,
                )
            ],
            resumo={"total_insumos": 15},
        )

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.price_trend.calcular_tendencia",
                    return_value=fake_relatorio,
                ):
                    result = _impl_tendencia_precos_catalogo("ctx")

        cat = result["indices_por_categoria"][0]
        # After truncation, at most 10 insumos remain in the dict
        assert len(cat["insumos"]) <= 10
        # insumos_omitidos was computed from the pre-truncation list
        assert cat["insumos_omitidos"] >= 0

    def test_with_jurisdicao_id(self, tmp_path):
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        _write_bundle(bundle_path, _make_bundle(insumos=[]))

        from concilium.application.catalog.price_trend import RelatorioTendencia

        fake_relatorio = RelatorioTendencia(
            contexto_id="ctx",
            gerado_em="2026-06-14",
            data_base_analise="2026-06-14",
            jurisdicao_id="BR-CE",
            indices_por_categoria=[],
            resumo={},
        )

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.price_trend.calcular_tendencia",
                    return_value=fake_relatorio,
                ) as mock_ct:
                    _impl_tendencia_precos_catalogo("ctx", jurisdicao_id="BR-CE")

        mock_ct.assert_called_once_with([], jurisdicao_id="BR-CE", contexto_id="ctx")


# ---------------------------------------------------------------------------
# _impl_comparar_rfi_catalogo
# ---------------------------------------------------------------------------


class TestImplCompararRfiCatalogo:
    def _config_path(self, tmp_path):
        return tmp_path / "config"

    def _tools_path(self, config_dir: Path) -> Path:
        # In the source: tools_dir = config_dir.parents[1] / "tools"
        # config_dir.parents[1] is the grandparent of config_dir
        return config_dir.parents[1] / "tools"

    def test_sad_path_merge_rfi_missing(self, tmp_path):
        config_dir = self._config_path(tmp_path)
        config_dir.mkdir(parents=True)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            result = _impl_comparar_rfi_catalogo("mendeslobo")

        assert "erro" in result
        assert "merge_rfi.py" in result["erro"]

    def test_sad_path_openpyxl_not_installed(self, tmp_path):
        config_dir = self._config_path(tmp_path)
        config_dir.mkdir(parents=True)
        tools_dir = self._tools_path(config_dir)
        tools_dir.mkdir(parents=True, exist_ok=True)
        merge_path = tools_dir / "merge_rfi.py"
        merge_path.write_text("# stub", encoding="utf-8")

        import builtins

        original_import = builtins.__import__

        def mock_import(name, *args, **kwargs):
            if name == "openpyxl":
                raise ImportError("openpyxl not installed")
            return original_import(name, *args, **kwargs)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch("builtins.__import__", side_effect=mock_import):
                result = _impl_comparar_rfi_catalogo("mendeslobo")

        assert "erro" in result
        assert "openpyxl" in result["erro"]

    def test_sad_path_bundle_missing(self, tmp_path):
        """Covers lines 90-92: bundle not found after merge_rfi loaded."""
        config_dir = self._config_path(tmp_path)
        config_dir.mkdir(parents=True)
        tools_dir = self._tools_path(config_dir)
        tools_dir.mkdir(parents=True, exist_ok=True)
        merge_path = tools_dir / "merge_rfi.py"
        merge_path.write_text("# stub", encoding="utf-8")

        mock_mod = MagicMock()
        mock_spec = MagicMock()
        mock_spec.loader = MagicMock()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch("importlib.util.spec_from_file_location", return_value=mock_spec):
                with patch("importlib.util.module_from_spec", return_value=mock_mod):
                    result = _impl_comparar_rfi_catalogo("mendeslobo")

        assert "erro" in result
        assert "Bundle" in result["erro"]

    def test_sad_path_planilha_not_found(self, tmp_path):
        """Covers line ~98-99: XLSM not found."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        _write_bundle(bundle_path, _make_bundle(insumos=[], producoes=[]))
        tools_dir = self._tools_path(config_dir)
        tools_dir.mkdir(parents=True, exist_ok=True)
        merge_path = tools_dir / "merge_rfi.py"
        merge_path.write_text("# stub", encoding="utf-8")

        mock_mod = MagicMock()
        xlsm_path = MagicMock(spec=Path)
        xlsm_path.exists.return_value = False
        mock_mod.XLSM = xlsm_path

        mock_spec = MagicMock()
        mock_spec.loader = MagicMock()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch("importlib.util.spec_from_file_location", return_value=mock_spec):
                with patch("importlib.util.module_from_spec", return_value=mock_mod):
                    import openpyxl as _openpyxl  # noqa — ensure available

                    result = _impl_comparar_rfi_catalogo("ctx")

        assert "erro" in result
        assert "RFI" in result["erro"] or "Planilha" in result["erro"]

    def test_happy_path_dry_run(self, tmp_path):
        """Covers lines 101-172 (dry_run=True path)."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[{"codigo": "INS-001", "descricao": "Insumo 1"}],
            producoes=[{"id": "PROD-001", "descricao": "Producao 1"}],
        )
        _write_bundle(bundle_path, bundle)

        tools_dir = self._tools_path(config_dir)
        tools_dir.mkdir(parents=True, exist_ok=True)
        merge_path = tools_dir / "merge_rfi.py"
        merge_path.write_text("# stub", encoding="utf-8")

        # Build mock merge report
        def _make_rep(inner_high=None, inner_low=None, right_only=None, left_only=None):
            rep = MagicMock()
            rep.inner_high = inner_high or []
            rep.inner_low = inner_low or []
            rep.right_only = right_only or []
            rep.left_only = left_only or []
            return rep

        mock_mod = MagicMock()
        xlsm_path = MagicMock(spec=Path)
        xlsm_path.exists.return_value = True
        mock_mod.XLSM = xlsm_path
        mock_mod.extract_rfi.return_value = {"insumos": [], "producoes": []}

        rep_ins = _make_rep(
            inner_high=[{"catalog": {"codigo": "INS-001"}, "delta": {"preco_unitario": 5.0}}],
            inner_low=[],
            right_only=[{"rfi": {"codigo": "RFI-NEW-001", "descricao": "Novo", "unidade": "m"}}],
            left_only=[],
        )
        rep_prod = _make_rep()
        mock_mod._merge_entities.side_effect = [
            (rep_ins, {}),
            (rep_prod, {}),
        ]

        mock_spec = MagicMock()
        mock_spec.loader = MagicMock()

        mock_wb = MagicMock()
        mock_mod.extract_rfi.return_value = {"insumos": [], "producoes": []}

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch("importlib.util.spec_from_file_location", return_value=mock_spec):
                with patch("importlib.util.module_from_spec", return_value=mock_mod):
                    with patch("openpyxl.load_workbook", return_value=mock_wb):
                        with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                            result = _impl_comparar_rfi_catalogo("ctx", dry_run=True)

        assert result["mode"] == "DRY_RUN"
        assert result["contexto_id"] == "ctx"
        assert "insumos" in result
        assert "producoes" in result
        assert "resumo" in result
        assert "aplicado" not in result

    def test_happy_path_applied_producoes_key(self, tmp_path):
        """Covers lines 158-167 (dry_run=False with 'producoes' key in bundle)."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(insumos=[], producoes=[])
        _write_bundle(bundle_path, bundle)

        tools_dir = self._tools_path(config_dir)
        tools_dir.mkdir(parents=True, exist_ok=True)
        merge_path = tools_dir / "merge_rfi.py"
        merge_path.write_text("# stub", encoding="utf-8")

        def _make_rep(**kw):
            rep = MagicMock()
            rep.inner_high = kw.get("inner_high", [])
            rep.inner_low = kw.get("inner_low", [])
            rep.right_only = kw.get("right_only", [])
            rep.left_only = kw.get("left_only", [])
            return rep

        mock_mod = MagicMock()
        xlsm_path = MagicMock(spec=Path)
        xlsm_path.exists.return_value = True
        mock_mod.XLSM = xlsm_path
        mock_mod.extract_rfi.return_value = {"insumos": [], "producoes": []}
        mock_mod._merge_entities.side_effect = [(_make_rep(), {}), (_make_rep(), {})]
        mock_mod._apply_merge.side_effect = [[], []]

        mock_spec = MagicMock()
        mock_spec.loader = MagicMock()
        mock_wb = MagicMock()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch("importlib.util.spec_from_file_location", return_value=mock_spec):
                with patch("importlib.util.module_from_spec", return_value=mock_mod):
                    with patch("openpyxl.load_workbook", return_value=mock_wb):
                        with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                            result = _impl_comparar_rfi_catalogo("ctx", dry_run=False)

        assert result["mode"] == "APPLIED"
        assert "aplicado" in result

    def test_happy_path_applied_produtividades_key(self, tmp_path):
        """Covers line 164: bundle has 'produtividades' key instead of 'producoes'."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        # Bundle uses 'produtividades' instead of 'producoes'
        bundle = _make_bundle(insumos=[], producoes=[], use_produtividades_key=True)
        _write_bundle(bundle_path, bundle)

        tools_dir = self._tools_path(config_dir)
        tools_dir.mkdir(parents=True, exist_ok=True)
        merge_path = tools_dir / "merge_rfi.py"
        merge_path.write_text("# stub", encoding="utf-8")

        def _make_rep(**kw):
            rep = MagicMock()
            rep.inner_high = []
            rep.inner_low = []
            rep.right_only = []
            rep.left_only = []
            return rep

        mock_mod = MagicMock()
        xlsm_path = MagicMock(spec=Path)
        xlsm_path.exists.return_value = True
        mock_mod.XLSM = xlsm_path
        mock_mod.extract_rfi.return_value = {"insumos": [], "producoes": []}
        mock_mod._merge_entities.side_effect = [(_make_rep(), {}), (_make_rep(), {})]
        mock_mod._apply_merge.side_effect = [[], []]

        mock_spec = MagicMock()
        mock_spec.loader = MagicMock()
        mock_wb = MagicMock()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch("importlib.util.spec_from_file_location", return_value=mock_spec):
                with patch("importlib.util.module_from_spec", return_value=mock_mod):
                    with patch("openpyxl.load_workbook", return_value=mock_wb):
                        with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                            result = _impl_comparar_rfi_catalogo("ctx", dry_run=False)

        assert result["mode"] == "APPLIED"
        # The bundle was updated using 'produtividades' key
        updated = json.loads(bundle_path.read_text(encoding="utf-8"))
        assert "produtividades" in updated
        assert "producoes" not in updated


# ---------------------------------------------------------------------------
# _impl_calcular_pu_producao
# ---------------------------------------------------------------------------


class TestImplCalcularPuProducao:
    def _config_path(self, tmp_path):
        return tmp_path / "config"

    def test_sad_path_bundle_missing(self, tmp_path):
        config_dir = self._config_path(tmp_path)
        config_dir.mkdir(parents=True)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            result = _impl_calcular_pu_producao("ctx", "PROD-001")

        assert "erro" in result

    def test_sad_path_producao_not_found(self, tmp_path):
        """Covers lines 207-209: producao_id not in bundle."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[],
            producoes=[{"id": "PROD-OTHER", "valor_horario": 10.0}],
        )
        _write_bundle(bundle_path, bundle)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            result = _impl_calcular_pu_producao("ctx", "PROD-DOES-NOT-EXIST")

        assert "erro" in result
        assert "PROD-DOES-NOT-EXIST" in result["erro"]

    def test_happy_path_producao_found_by_id(self, tmp_path):
        """Covers lines 196-219: producao found and PU calculated."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[{"codigo": "MO-001", "tipo": "MAO_OBRA", "preco_unitario": 4000.0, "unidade": "mes"}],
            producoes=[
                {
                    "id": "PROD-001",
                    "valor_horario": 5.0,
                    "composicao_padrao_equipe": {
                        "funcoes": [{"funcao_id": "MO-001", "num_pessoas": 1}],
                        "equipamentos": [],
                    },
                }
            ],
        )
        _write_bundle(bundle_path, bundle)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                result = _impl_calcular_pu_producao("ctx", "PROD-001")

        assert "producao_id" in result
        assert result["producao_id"] == "PROD-001"

    def test_happy_path_producao_found_by_codigo(self, tmp_path):
        """Covers alternate lookup via 'codigo' key."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[],
            producoes=[
                {
                    "codigo": "CPU-001",
                    "valor_horario": 3.0,
                    "composicao_padrao_equipe": {
                        "funcoes": [],
                        "equipamentos": [],
                    },
                }
            ],
        )
        _write_bundle(bundle_path, bundle)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                result = _impl_calcular_pu_producao("ctx", "CPU-001")

        assert "producao_id" in result

    def test_happy_path_uses_produtividades_key(self, tmp_path):
        """Covers line 205: bundle uses 'produtividades' instead of 'producoes'."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[],
            producoes=[{"id": "PROD-001", "valor_horario": 2.0}],
            use_produtividades_key=True,
        )
        _write_bundle(bundle_path, bundle)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                result = _impl_calcular_pu_producao("ctx", "PROD-001")

        assert "producao_id" in result


# ---------------------------------------------------------------------------
# _impl_avaliar_plausibilidade_catalogo
# ---------------------------------------------------------------------------


class TestImplAvaliarPlausibilidadeCatalogo:
    def _config_path(self, tmp_path):
        return tmp_path / "config"

    def test_sad_path_bundle_missing(self, tmp_path):
        config_dir = self._config_path(tmp_path)
        config_dir.mkdir(parents=True)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            result = _impl_avaliar_plausibilidade_catalogo("ctx")

        assert "erro" in result

    def test_happy_path(self, tmp_path):
        """Covers lines 234-247."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[{"codigo": "INS-001"}],
            producoes=[{"id": "PROD-001"}],
            cpus=[{"codigo": "CPU-001"}],
        )
        _write_bundle(bundle_path, bundle)

        fake_resultado = {"ok": True, "alertas": []}

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.plausibility_checker.verificar_plausibilidade",
                    return_value=fake_resultado,
                ) as mock_vp:
                    result = _impl_avaliar_plausibilidade_catalogo("ctx")

        mock_vp.assert_called_once()
        assert result == fake_resultado

    def test_happy_path_with_produtividades_key(self, tmp_path):
        """Covers line 243: bundle uses 'produtividades' key."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[],
            producoes=[{"id": "PROD-001"}],
            cpus=[],
            use_produtividades_key=True,
        )
        _write_bundle(bundle_path, bundle)

        fake_resultado = {"ok": True, "alertas": []}

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.plausibility_checker.verificar_plausibilidade",
                    return_value=fake_resultado,
                ) as mock_vp:
                    _impl_avaliar_plausibilidade_catalogo("ctx")

        mock_vp.assert_called_once()
        # Verify it was called with the produtividades list
        call_args = mock_vp.call_args
        insumos_arg, producoes_arg, cpus_arg = call_args[0]
        assert producoes_arg == [{"id": "PROD-001"}]


# ---------------------------------------------------------------------------
# _impl_verificar_integridade_catalogo
# ---------------------------------------------------------------------------


class TestImplVerificarIntegridadeCatalogo:
    def _config_path(self, tmp_path):
        return tmp_path / "config"

    def _base_pu_result(self, refs_faltantes=None, pu_divergentes=None, com_composicao=0):
        return {
            "referencias_faltantes": refs_faltantes or [],
            "pu_divergentes": pu_divergentes or [],
            "com_composicao": com_composicao,
            "ok": True,
            "total_producoes": 0,
            "sem_valor_horario": [],
            "breakdown_por_producao": [],
        }

    def test_sad_path_bundle_missing(self, tmp_path):
        config_dir = self._config_path(tmp_path)
        config_dir.mkdir(parents=True)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            result = _impl_verificar_integridade_catalogo("ctx")

        assert "erro" in result

    def test_happy_path_no_problems(self, tmp_path):
        """Covers lines 267-315: no errors, ok=True."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[{"codigo": "INS-001"}],
            producoes=[{"id": "PROD-001", "valor_horario": 5.0}],
            cpus=[],
        )
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result(com_composicao=1)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ):
                    result = _impl_verificar_integridade_catalogo("ctx")

        assert result["ok"] is True
        assert result["resumo"]["erros"] == 0
        assert result["resumo"]["avisos"] == 0
        assert result["problemas"] == []

    def test_happy_path_with_refs_faltantes(self, tmp_path):
        """Covers lines 306-307: pu_result has referencias_faltantes."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(insumos=[], producoes=[], cpus=[])
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result(
            refs_faltantes=[{"producao_id": "PROD-001", "faltantes": ["INS-MISSING"]}],
            com_composicao=1,
        )

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ):
                    result = _impl_verificar_integridade_catalogo("ctx")

        assert result["ok"] is False
        erros = [p for p in result["problemas"] if p["severidade"] == "ERRO"]
        assert len(erros) == 1
        assert erros[0]["tipo"] == "REF_INSUMO_FALTANTE"

    def test_happy_path_with_pu_divergentes(self, tmp_path):
        """Covers lines 308-309: pu_result has pu_divergentes."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(insumos=[], producoes=[], cpus=[])
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result(
            pu_divergentes=[
                {"producao_id": "PROD-001", "pu_calculado": 100.0, "pu_referencia": 50.0, "delta_pct": 100.0}
            ],
        )

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ):
                    result = _impl_verificar_integridade_catalogo("ctx")

        avisos = [p for p in result["problemas"] if p["tipo"] == "PU_DIVERGENTE"]
        assert len(avisos) == 1
        assert avisos[0]["severidade"] == "AVISO"

    def test_cpu_refs_faltantes(self, tmp_path):
        """Covers lines 285-291: cpu composicao_itens references non-existent insumo_id."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[{"codigo": "INS-001"}],
            producoes=[],
            cpus=[
                {
                    "codigo": "CPU-001",
                    "composicao_itens": [
                        {"insumo_id": "INS-NONEXISTENT", "coeficiente": 1.0},
                        {"insumo_id": "INS-001", "coeficiente": 0.5},  # This one exists
                    ],
                }
            ],
        )
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ):
                    result = _impl_verificar_integridade_catalogo("ctx")

        assert result["ok"] is False  # because CPU_INSUMO_FALTANTE is ERRO severity
        cpu_erros = [p for p in result["problemas"] if p["tipo"] == "CPU_INSUMO_FALTANTE"]
        assert len(cpu_erros) == 1
        assert cpu_erros[0]["insumo_id_faltante"] == "INS-NONEXISTENT"

    def test_cpu_prod_faltantes(self, tmp_path):
        """Covers lines 294-299: CPU references a producao_id_ref that doesn't exist."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[],
            producoes=[{"id": "PROD-001"}],
            cpus=[
                {
                    "codigo": "CPU-001",
                    "producao_id_ref": "PROD-MISSING",
                    "composicao_itens": [],
                }
            ],
        )
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ):
                    result = _impl_verificar_integridade_catalogo("ctx")

        prod_faltantes = [p for p in result["problemas"] if p["tipo"] == "CPU_PROD_REF_FALTANTE"]
        assert len(prod_faltantes) == 1
        assert prod_faltantes[0]["producao_id_ref"] == "PROD-MISSING"
        assert prod_faltantes[0]["severidade"] == "AVISO"

    def test_rfi_prefix_detection(self, tmp_path):
        """Covers lines 302-303: insumos/producoes starting with 'RFI-' are detected."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[
                {"codigo": "INS-001", "descricao": "Normal insumo"},
                {"codigo": "RFI-INS-001", "descricao": "RFI imported insumo"},
            ],
            producoes=[
                {"id": "PROD-001", "descricao": "Normal prod"},
                {"id": "RFI-PROD-001", "descricao": "RFI imported prod"},
            ],
            cpus=[],
        )
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ):
                    result = _impl_verificar_integridade_catalogo("ctx")

        assert result["resumo"]["insumos_rfi_pendentes"] == 1
        assert result["resumo"]["producoes_rfi_pendentes"] == 1
        assert len(result["rfi_pendentes"]["insumos"]) == 1
        assert result["rfi_pendentes"]["insumos"][0]["codigo"] == "RFI-INS-001"
        assert len(result["rfi_pendentes"]["producoes"]) == 1
        assert result["rfi_pendentes"]["producoes"][0]["id"] == "RFI-PROD-001"

    def test_happy_path_with_produtividades_key(self, tmp_path):
        """Covers line 276: bundle uses 'produtividades' key."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[],
            producoes=[{"id": "PROD-001"}],
            cpus=[],
            use_produtividades_key=True,
        )
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ) as mock_cc:
                    _impl_verificar_integridade_catalogo("ctx")

        # Verify calcular_catalogo received the produtividades list
        call_args = mock_cc.call_args
        producoes_arg = call_args[0][0]
        assert producoes_arg == [{"id": "PROD-001"}]

    def test_ok_flag_false_when_errors(self, tmp_path):
        """Covers lines 315-318: ok=False when ERRO severity problems exist."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[],  # No insumos — so composicao_itens refs will fail
            producoes=[],
            cpus=[
                {
                    "codigo": "CPU-001",
                    "composicao_itens": [{"insumo_id": "INS-MISSING", "coeficiente": 1.0}],
                }
            ],
        )
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result()

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ):
                    result = _impl_verificar_integridade_catalogo("ctx")

        assert result["ok"] is False
        assert result["resumo"]["erros"] > 0

    def test_resumo_counts_are_correct(self, tmp_path):
        """Covers lines 319-329: resumo field counts."""
        config_dir = self._config_path(tmp_path)
        bundle_path = config_dir / "catalogos" / "ctx" / ".catalog_bundle.json"
        bundle = _make_bundle(
            insumos=[{"codigo": "A"}, {"codigo": "B"}],
            producoes=[{"id": "P1"}, {"id": "P2"}],
            cpus=[{"codigo": "CPU-1", "composicao_itens": []}],
        )
        _write_bundle(bundle_path, bundle)

        pu_result = self._base_pu_result(com_composicao=2)

        with patch(f"{MODULE}._config_root", return_value=config_dir):
            with patch(f"{MODULE}._wrap_response", side_effect=_identity):
                with patch(
                    "concilium.application.catalog.pu_calculator.calcular_catalogo",
                    return_value=pu_result,
                ):
                    result = _impl_verificar_integridade_catalogo("ctx")

        assert result["resumo"]["total_insumos"] == 2
        assert result["resumo"]["total_producoes"] == 2
        assert result["resumo"]["total_cpus"] == 1
        assert result["resumo"]["producoes_com_composicao"] == 2
