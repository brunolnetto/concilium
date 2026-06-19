"""Tests for src/concilium/infrastructure/catalog/extraction — tier-2 coverage."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest

# ---------------------------------------------------------------------------
# __init__ importability
# ---------------------------------------------------------------------------


def test_init_imports():
    from concilium.infrastructure.catalog.extraction import (  # noqa: F401
        consolidate_reference,
        extract_cpu_references,
        extract_engine_iteration_references,
        extract_reference_data,
        iter_benchmark_reports,
    )


# ---------------------------------------------------------------------------
# catalog_refs
# ---------------------------------------------------------------------------

from concilium.infrastructure.catalog.extraction.catalog_refs import (
    _load_json,
    _safe_float,
    extract_cpu_references,
)


class TestLoadJsonCatalogRefs:
    def test_nonexistent_path_returns_empty(self, tmp_path):
        result = _load_json(tmp_path / "does_not_exist.json")
        assert result == {}

    def test_valid_json_returns_dict(self, tmp_path):
        p = tmp_path / "valid.json"
        p.write_text('{"key": "value"}', encoding="utf-8")
        assert _load_json(p) == {"key": "value"}

    def test_invalid_json_returns_empty(self, tmp_path):
        p = tmp_path / "bad.json"
        p.write_text("not json {{", encoding="utf-8")
        assert _load_json(p) == {}


class TestSafeFloat:
    def test_none_returns_none(self):
        assert _safe_float(None) is None

    def test_valid_float_string(self):
        assert _safe_float("3.14") == pytest.approx(3.14)

    def test_int_returns_float(self):
        result = _safe_float(42)
        assert result == 42.0
        assert isinstance(result, float)

    def test_invalid_string_returns_none(self):
        assert _safe_float("abc") is None

    def test_empty_string_returns_none(self):
        assert _safe_float("") is None


def _make_catalog_repo(tmp: Path):
    """Helper: returns (bench_root, cpus_dir, prods_dir) after creating them."""
    bench_root = tmp / "benchmarks" / "data"
    cpus_dir = tmp / "config" / "catalogos" / "mendeslobo" / "cpus"
    prods_dir = tmp / "config" / "catalogos" / "mendeslobo" / "produtividades"
    for d in (bench_root, cpus_dir, prods_dir):
        d.mkdir(parents=True, exist_ok=True)
    return bench_root, cpus_dir, prods_dir


class TestExtractCpuReferences:
    def test_no_referencia_json_returns_empty_cpu_referencias(self, tmp_path):
        _make_catalog_repo(tmp_path)
        result = extract_cpu_references("SLUG", root=tmp_path)
        assert result["slug"] == "SLUG"
        assert result["cpu_referencias"]["total"] == 0
        assert result["cpu_referencias"]["catalog_cpus"] == {}

    def test_eap_cpu_id_missing_in_catalog(self, tmp_path):
        bench_root, _cpus, _prods = _make_catalog_repo(tmp_path)
        slug_dir = bench_root / "PROJ"
        slug_dir.mkdir()
        ref = {"eap": [{"cpu_id": "CPU-MISSING"}]}
        (slug_dir / "referencia.json").write_text(json.dumps(ref), encoding="utf-8")

        result = extract_cpu_references("PROJ", root=tmp_path)
        assert result["cpu_referencias"]["catalog_cpus"]["CPU-MISSING"]["status"] == "NAO_ENCONTRADO_NO_CATALOGO"

    def test_cpu_present_but_no_producao_id_ref(self, tmp_path):
        bench_root, cpus_dir, _prods = _make_catalog_repo(tmp_path)
        slug_dir = bench_root / "PROJ"
        slug_dir.mkdir()
        ref = {"eap": [{"cpu_id": "CPU-001"}]}
        (slug_dir / "referencia.json").write_text(json.dumps(ref), encoding="utf-8")
        cpu_def = {"descricao": "Escavação", "unidade": "m3"}  # no producao_id_ref
        (cpus_dir / "CPU-001.json").write_text(json.dumps(cpu_def), encoding="utf-8")

        result = extract_cpu_references("PROJ", root=tmp_path)
        assert result["cpu_referencias"]["catalog_cpus"]["CPU-001"]["status"] == "ERRO_SEM_PRODUCAO_ID_REF"

    def test_cpu_with_producao_id_ref_but_productivity_missing(self, tmp_path):
        bench_root, cpus_dir, _prods = _make_catalog_repo(tmp_path)
        slug_dir = bench_root / "PROJ"
        slug_dir.mkdir()
        ref = {"eap": [{"cpu_id": "CPU-001"}]}
        (slug_dir / "referencia.json").write_text(json.dumps(ref), encoding="utf-8")
        cpu_def = {"descricao": "Escavação", "unidade": "m3", "producao_id_ref": "PROD-X"}
        (cpus_dir / "CPU-001.json").write_text(json.dumps(cpu_def), encoding="utf-8")
        # produtividades dir empty — PROD-X.json not created

        result = extract_cpu_references("PROJ", root=tmp_path)
        entry = result["cpu_referencias"]["catalog_cpus"]["CPU-001"]
        assert entry["status"] == "ERRO_PRODUTIVIDADE_NAO_ENCONTRADA"
        assert entry["producao_id_ref"] == "PROD-X"

    def test_full_happy_path(self, tmp_path):
        bench_root, cpus_dir, prods_dir = _make_catalog_repo(tmp_path)
        slug_dir = bench_root / "PROJ"
        slug_dir.mkdir()
        ref = {"eap": [{"cpu_id": "CPU-001"}]}
        (slug_dir / "referencia.json").write_text(json.dumps(ref), encoding="utf-8")
        cpu_def = {"descricao": "Escavação", "unidade": "m3", "producao_id_ref": "PROD-X"}
        (cpus_dir / "CPU-001.json").write_text(json.dumps(cpu_def), encoding="utf-8")
        prod_def = {
            "descricao": "Produção escavação",
            "valor_horario": "150.5",
            "cpu_ids": ["CPU-001"],
        }
        (prods_dir / "PROD-X.json").write_text(json.dumps(prod_def), encoding="utf-8")

        result = extract_cpu_references("PROJ", root=tmp_path)
        entry = result["cpu_referencias"]["catalog_cpus"]["CPU-001"]
        assert entry["producao_horaria"] == pytest.approx(150.5)
        assert entry["producao_id_ref"] == "PROD-X"
        assert "producoes" in entry
        assert entry["producoes"]["PROD-X"]["producao_horaria"] == pytest.approx(150.5)


# ---------------------------------------------------------------------------
# consolidation
# ---------------------------------------------------------------------------

from concilium.infrastructure.catalog.extraction.consolidation import (
    consolidate_reference,
)


def _make_consolidation_repo(tmp: Path, slug: str, projeto_nome: str):
    """Create minimal repo structure for consolidation tests."""
    bench_data = tmp / "benchmarks" / "data" / slug
    bench_data.mkdir(parents=True)
    ref = {"slug": slug, "projeto": projeto_nome, "eap": []}
    (bench_data / "referencia.json").write_text(json.dumps(ref), encoding="utf-8")

    proj_uuid = "aaaabbbb-cccc-dddd-eeee-ffffffffffff"
    proj_dir = tmp / "projetos" / proj_uuid
    v1_dir = proj_dir / "v1"
    v1_dir.mkdir(parents=True)
    meta = {"nome": projeto_nome, "criado_em": "2024-01-01T00:00:00"}
    (proj_dir / "meta.json").write_text(json.dumps(meta), encoding="utf-8")
    return proj_dir, v1_dir


class TestConsolidateReference:
    def test_no_referencia_json_returns_empty_dict(self, tmp_path):
        (tmp_path / "benchmarks" / "data").mkdir(parents=True)
        (tmp_path / "projetos").mkdir(parents=True)
        result = consolidate_reference("MISSING", root=tmp_path)
        assert result == {}

    def test_project_not_found_returns_raw_ref(self, tmp_path):
        bench_data = tmp_path / "benchmarks" / "data" / "PROJ"
        bench_data.mkdir(parents=True)
        ref = {"slug": "PROJ", "projeto": "Projeto Fantasma", "eap": []}
        (bench_data / "referencia.json").write_text(json.dumps(ref), encoding="utf-8")
        (tmp_path / "projetos").mkdir(parents=True)
        # No matching project dir

        result = consolidate_reference("PROJ", root=tmp_path)
        # Falls back to raw ref (no project found)
        assert result == ref

    def test_full_happy_path_enriches_with_equipes_and_planejamento(self, tmp_path):
        _proj_dir, v1_dir = _make_consolidation_repo(tmp_path, "PROJ", "Obra Principal")
        canonical = {
            "equipes_diretas": [
                {
                    "equipe_id": "EQ-1",
                    "item_id": "1.1",
                    "cpu_id": "CPU-001",
                    "nome": "Equipe A",
                    "funcoes": [{"funcao": "Operador"}],
                    "equipamentos": [],
                    "custo_total": 10000,
                    "custo_mao_obra_mensal": 8000,
                    "custo_equipamentos_mensal": 2000,
                    "custo_combustivel_total_mensal": 500,
                    "custo_material_mensal": 200,
                    "quantitativo": 100,
                }
            ],
            "planejamento": [
                {
                    "item_id": "1.1",
                    "descricao": "Escavação",
                    "quantitativo": 100,
                    "subtotal": {"sem_bdi": 9000, "com_bdi": 10500},
                }
            ],
        }
        (v1_dir / "canonical.json").write_text(json.dumps(canonical), encoding="utf-8")

        result = consolidate_reference("PROJ", root=tmp_path)
        assert "equipes_detalhadas" in result
        assert "planejamento_detalhado" in result
        assert result["_fonte"] == "canonical_consolidado"
        assert result["equipes_detalhadas"][0]["equipe_id"] == "EQ-1"
        assert result["planejamento_detalhado"][0]["subtotal_sem_bdi"] == 9000
        assert result["planejamento_detalhado"][0]["subtotal_com_bdi"] == 10500

    def test_empty_canonical_yields_empty_lists(self, tmp_path):
        _proj_dir, v1_dir = _make_consolidation_repo(tmp_path, "PROJ2", "Obra Vazia")
        (v1_dir / "canonical.json").write_text(json.dumps({}), encoding="utf-8")

        result = consolidate_reference("PROJ2", root=tmp_path)
        assert result["equipes_detalhadas"] == []
        assert result["planejamento_detalhado"] == []


# ---------------------------------------------------------------------------
# engine_refs
# ---------------------------------------------------------------------------

from concilium.infrastructure.catalog.extraction.engine_refs import (
    _check_project_dir,
    _normalize_project_token,
    extract_engine_iteration_references,
    iter_benchmark_reports,
)


class TestNormalizeProjectToken:
    def test_removes_calibrado(self):
        assert _normalize_project_token("Obra CALIBRADO X") == "OBRA X"

    def test_uppercases(self):
        assert _normalize_project_token("obra x") == "OBRA X"

    def test_collapses_spaces(self):
        assert _normalize_project_token("  OBRA   X  ") == "OBRA X"

    def test_empty_string(self):
        assert _normalize_project_token("") == ""

    def test_calibrado_standalone(self):
        assert _normalize_project_token("CALIBRADO") == ""

    def test_underscores_become_spaces(self):
        result = _normalize_project_token("OBRA_X_CALIBRADO")
        assert "CALIBRADO" not in result
        assert "OBRA" in result
        assert "X" in result


class TestIterBenchmarkReports:
    def test_empty_dir_returns_empty_list(self, tmp_path):
        result = iter_benchmark_reports(root=tmp_path)
        assert result == []

    def test_finds_benchmark_json_files(self, tmp_path):
        (tmp_path / "benchmark_CTC.json").write_text("{}", encoding="utf-8")
        (tmp_path / "benchmark_AGIS.json").write_text("{}", encoding="utf-8")
        result = iter_benchmark_reports(root=tmp_path)
        names = {p.name for p in result}
        assert "benchmark_CTC.json" in names
        assert "benchmark_AGIS.json" in names

    def test_skips_benchmark_template(self, tmp_path):
        (tmp_path / "benchmark_template.json").write_text("{}", encoding="utf-8")
        (tmp_path / "benchmark_REAL.json").write_text("{}", encoding="utf-8")
        result = iter_benchmark_reports(root=tmp_path)
        names = {p.name for p in result}
        assert "benchmark_template.json" not in names
        assert "benchmark_REAL.json" in names

    def test_non_matching_files_ignored(self, tmp_path):
        (tmp_path / "report_CTC.json").write_text("{}", encoding="utf-8")
        (tmp_path / "other.txt").write_text("x", encoding="utf-8")
        result = iter_benchmark_reports(root=tmp_path)
        assert result == []


class TestCheckProjectDir:
    def test_non_dir_returns_best_unchanged(self, tmp_path):
        f = tmp_path / "file.txt"
        f.write_text("x")
        result = _check_project_dir(f, "TARGET", None)
        assert result is None

    def test_dir_without_meta_returns_best_unchanged(self, tmp_path):
        d = tmp_path / "proj"
        d.mkdir()
        result = _check_project_dir(d, "TARGET", None)
        assert result is None

    def test_dir_with_matching_meta_returns_tuple(self, tmp_path):
        d = tmp_path / "proj"
        d.mkdir()
        meta = {"nome": "TARGET", "criado_em": "2024-01-01T00:00:00"}
        (d / "meta.json").write_text(json.dumps(meta), encoding="utf-8")
        result = _check_project_dir(d, "TARGET", None)
        assert result is not None
        assert result[1] == d

    def test_prefers_later_created_at(self, tmp_path):
        d1 = tmp_path / "proj1"
        d1.mkdir()
        (d1 / "meta.json").write_text(
            json.dumps({"nome": "TARGET", "criado_em": "2023-01-01T00:00:00"}), encoding="utf-8"
        )
        d2 = tmp_path / "proj2"
        d2.mkdir()
        (d2 / "meta.json").write_text(
            json.dumps({"nome": "TARGET", "criado_em": "2024-06-01T00:00:00"}), encoding="utf-8"
        )
        best = _check_project_dir(d1, "TARGET", None)
        best = _check_project_dir(d2, "TARGET", best)
        assert best[1] == d2


class TestExtractEngineIterationReferences:
    def _build_repo(self, tmp: Path, slug: str, projeto_nome: str):
        reports_dir = tmp
        proj_uuid = "test-uuid-0001"
        proj_dir = tmp / "projetos" / proj_uuid
        v1_dir = proj_dir / "v1"
        v1_dir.mkdir(parents=True)
        meta = {"nome": projeto_nome, "criado_em": "2024-01-01T00:00:00"}
        (proj_dir / "meta.json").write_text(json.dumps(meta), encoding="utf-8")
        canonical = {
            "equipes_diretas": [{"equipe_id": "EQ-1", "item_id": "1"}],
            "planejamento": [{"item_id": "1", "descricao": "Escav"}],
            "resultado": {"custo_direto": 100000},
            "indiretos": [{"categoria": "ADM", "total": 5000}],
        }
        (v1_dir / "canonical.json").write_text(json.dumps(canonical), encoding="utf-8")
        report = {"slug": slug, "projeto": projeto_nome}
        report_path = reports_dir / f"benchmark_{slug}.json"
        report_path.write_text(json.dumps(report), encoding="utf-8")
        return report_path

    def test_happy_path_returns_dict_with_slug(self, tmp_path):
        out_dir = tmp_path / "extracted"
        report_path = self._build_repo(tmp_path, "TESTSLUG", "Meu Projeto")
        result = extract_engine_iteration_references(
            report_path=report_path,
            root=tmp_path,
            out_dir=out_dir,
        )
        assert result["slug"] == "TESTSLUG"
        assert result["projeto"] == "Meu Projeto"

    def test_happy_path_writes_output_file(self, tmp_path):
        out_dir = tmp_path / "extracted"
        report_path = self._build_repo(tmp_path, "TESTSLUG", "Meu Projeto")
        extract_engine_iteration_references(
            report_path=report_path,
            root=tmp_path,
            out_dir=out_dir,
        )
        out_file = out_dir / "iteracao_referencias_testslug.json"
        assert out_file.exists()

    def test_indiretos_populated_from_canonical(self, tmp_path):
        out_dir = tmp_path / "extracted2"
        report_path = self._build_repo(tmp_path, "TESTSLUG2", "Meu Projeto 2")
        result = extract_engine_iteration_references(
            report_path=report_path,
            root=tmp_path,
            out_dir=out_dir,
        )
        assert result["indiretos"]["ADM"]["total"] == 5000


# ---------------------------------------------------------------------------
# excel
# ---------------------------------------------------------------------------

from openpyxl import Workbook

from concilium.infrastructure.catalog.extraction.excel import (
    _extract_bdi,
    _extract_eap_col_map,
    _extract_indiretos,
    _extract_resultado,
    _get_val_near_label,
    extract_reference_data,
)


class TestGetValNearLabel:
    def test_label_found_returns_adjacent_value(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "TOTAL DIRETO"
        ws["B1"] = 99999.0
        result = _get_val_near_label(ws, "TOTAL DIRETO", col_offset=1)
        assert result == 99999.0

    def test_label_not_found_returns_none(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Something else"
        result = _get_val_near_label(ws, "TOTAL DIRETO", col_offset=1)
        assert result is None

    def test_case_insensitive_match(self):
        wb = Workbook()
        ws = wb.active
        ws["A2"] = "total direto"
        ws["B2"] = 42.0
        result = _get_val_near_label(ws, "TOTAL DIRETO", col_offset=1)
        assert result == 42.0

    def test_partial_label_match(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "GRAND TOTAL DIRETO OBRAS"
        ws["B1"] = 55.0
        result = _get_val_near_label(ws, "TOTAL DIRETO", col_offset=1)
        assert result == 55.0


class TestExtractEapColMap:
    def test_header_row_with_item_returns_col_map(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "ITEM"
        ws["B1"] = "DESCRICAO"
        ws["C1"] = "VALOR"
        result = _extract_eap_col_map(ws)
        assert "ITEM" in result
        assert "DESCRICAO" in result

    def test_no_matching_headers_returns_empty_dict(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "FOO"
        ws["B1"] = "BAR"
        result = _extract_eap_col_map(ws)
        assert result == {}

    def test_codigo_triggers_col_map(self):
        wb = Workbook()
        ws = wb.active
        ws["A3"] = "CODIGO"
        ws["B3"] = "DESCRICAO"
        result = _extract_eap_col_map(ws)
        assert "CODIGO" in result


class TestExtractResultado:
    def test_no_resultado_sheet_returns_empty(self):
        wb = Workbook()
        wb.active.title = "Dados"
        result = _extract_resultado(wb, wb.sheetnames)
        assert result == {}

    def test_resultado_sheet_present_returns_dict(self):
        wb = Workbook()
        wb.active.title = "RESULTADO"
        ws = wb["RESULTADO"]
        ws["A1"] = "TOTAL DIRETO"
        ws["C1"] = 1500000.0
        result = _extract_resultado(wb, wb.sheetnames)
        assert isinstance(result, dict)
        assert "custo_direto_total" in result


class TestExtractBdi:
    def test_no_bdi_sheet_returns_empty(self):
        wb = Workbook()
        wb.active.title = "Dados"
        result = _extract_bdi(wb, wb.sheetnames)
        assert result == {}

    def test_bdi_sheet_present_returns_dict(self):
        wb = Workbook()
        wb.active.title = "BDI"
        result = _extract_bdi(wb, wb.sheetnames)
        assert isinstance(result, dict)
        assert "mao_de_obra" in result
        assert "materiais" in result


class TestExtractIndiretos:
    def test_no_indirect_sheets_returns_empty(self):
        wb = Workbook()
        wb.active.title = "Dados"
        result = _extract_indiretos(wb, wb.sheetnames)
        assert result == {}

    def test_adm_sheet_found_returns_adm_entry(self):
        wb = Workbook()
        wb.active.title = "ADM"
        ws = wb["ADM"]
        ws["A1"] = "CUSTO FIXO MENSAL"
        ws["C1"] = 20000.0
        result = _extract_indiretos(wb, wb.sheetnames)
        assert "ADM" in result
        assert result["ADM"]["custo_fixo_mensal_total"] == 20000.0


class TestExtractReferenceData:
    def _make_xlsx_bytes(self) -> bytes:
        """Build an in-memory xlsx with RESULTADO, BDI, ADM, and ORCAMENTO sheets."""
        wb = Workbook()
        # Rename default sheet
        ws_res = wb.active
        ws_res.title = "RESULTADO"
        ws_res["A1"] = "TOTAL DIRETO"
        ws_res["C1"] = 5000000.0

        ws_bdi = wb.create_sheet("BDI")
        ws_bdi["A1"] = "Total do BDI"
        ws_bdi["B1"] = 0.25

        ws_adm = wb.create_sheet("ADM")
        ws_adm["A1"] = "CUSTO FIXO MENSAL"
        ws_adm["B1"] = 10000.0

        ws_orc = wb.create_sheet("ORCAMENTO PLAN")
        ws_orc["A1"] = "ITEM"
        ws_orc["B1"] = "DESCRICAO"

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_returns_expected_top_level_keys(self, tmp_path):
        xlsx_bytes = self._make_xlsx_bytes()
        xlsx_path = tmp_path / "ref.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        result = extract_reference_data(xlsx_path)
        assert "resultado" in result
        assert "bdi" in result
        assert "indiretos_detalhamento" in result
        assert "eap" in result

    def test_resultado_populated(self, tmp_path):
        xlsx_bytes = self._make_xlsx_bytes()
        xlsx_path = tmp_path / "ref.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        result = extract_reference_data(xlsx_path)
        assert result["resultado"].get("custo_direto_total") == 5000000.0

    def test_bdi_populated(self, tmp_path):
        xlsx_bytes = self._make_xlsx_bytes()
        xlsx_path = tmp_path / "ref.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        result = extract_reference_data(xlsx_path)
        assert "mao_de_obra" in result["bdi"]
        assert result["bdi"]["mao_de_obra"]["bdi_total"] == pytest.approx(0.25)

    def test_indiretos_populated(self, tmp_path):
        xlsx_bytes = self._make_xlsx_bytes()
        xlsx_path = tmp_path / "ref.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        result = extract_reference_data(xlsx_path)
        assert "ADM" in result["indiretos_detalhamento"]

    def test_eap_is_list(self, tmp_path):
        xlsx_bytes = self._make_xlsx_bytes()
        xlsx_path = tmp_path / "ref.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        result = extract_reference_data(xlsx_path)
        assert isinstance(result["eap"], list)

    def test_string_path_accepted(self, tmp_path):
        xlsx_bytes = self._make_xlsx_bytes()
        xlsx_path = tmp_path / "ref.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        result = extract_reference_data(str(xlsx_path))
        assert "resultado" in result
