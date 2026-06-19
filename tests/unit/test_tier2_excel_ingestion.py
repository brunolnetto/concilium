"""Tests for concilium.infrastructure.catalog.excel_ingestion."""

from datetime import date
from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from concilium.infrastructure.catalog.excel_ingestion import ExcelIngestionPipeline

# ── Helpers ───────────────────────────────────────────────────────────────────


def _make_excel_bytes(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_pipeline() -> ExcelIngestionPipeline:
    storage = MagicMock()
    storage.gravar_blob.return_value = None
    return ExcelIngestionPipeline(storage)


# ── _extract_rows_from_sheet ──────────────────────────────────────────────────


def test_extract_rows_sheet_not_in_workbook():
    excel = _make_excel_bytes({"Insumos": [["codigo", "descricao"], ["INS-01", "Cimento"]]})
    pipeline = _make_pipeline()
    from openpyxl import load_workbook

    wb = load_workbook(filename=BytesIO(excel), read_only=True, data_only=True)
    result = pipeline._extract_rows_from_sheet(wb, "NaoExiste")
    wb.close()
    assert result == []


def test_extract_rows_empty_sheet():
    excel = _make_excel_bytes({"Insumos": []})
    pipeline = _make_pipeline()
    from openpyxl import load_workbook

    wb = load_workbook(filename=BytesIO(excel), read_only=True, data_only=True)
    result = pipeline._extract_rows_from_sheet(wb, "Insumos")
    wb.close()
    assert result == []


def test_extract_rows_all_none_header():
    # Sheet with only a fully-None row → should return []
    excel = _make_excel_bytes({"Insumos": [[None, None, None]]})
    pipeline = _make_pipeline()
    from openpyxl import load_workbook

    wb = load_workbook(filename=BytesIO(excel), read_only=True, data_only=True)
    result = pipeline._extract_rows_from_sheet(wb, "Insumos")
    wb.close()
    assert result == []


def test_extract_rows_normal():
    excel = _make_excel_bytes(
        {
            "Insumos": [
                ["codigo", "descricao"],
                ["INS-01", "Cimento Portland"],
                ["INS-02", "Areia"],
            ]
        }
    )
    pipeline = _make_pipeline()
    from openpyxl import load_workbook

    wb = load_workbook(filename=BytesIO(excel), read_only=True, data_only=True)
    result = pipeline._extract_rows_from_sheet(wb, "Insumos")
    wb.close()
    assert len(result) == 2
    assert result[0]["codigo"] == "INS-01"
    assert result[1]["descricao"] == "Areia"


def test_extract_rows_skips_empty_rows():
    excel = _make_excel_bytes(
        {
            "Insumos": [
                ["codigo", "descricao"],
                [None, None],  # all-None → skipped
                ["INS-01", "Cimento"],
                [None, ""],  # all None/"" → skipped
            ]
        }
    )
    pipeline = _make_pipeline()
    from openpyxl import load_workbook

    wb = load_workbook(filename=BytesIO(excel), read_only=True, data_only=True)
    result = pipeline._extract_rows_from_sheet(wb, "Insumos")
    wb.close()
    assert len(result) == 1
    assert result[0]["codigo"] == "INS-01"


# ── parse_planilha ────────────────────────────────────────────────────────────


def test_parse_planilha_returns_dict_keys():
    excel = _make_excel_bytes(
        {
            "Insumos": [["codigo"], ["INS-01"]],
            "CPUs": [["cpu_id"], ["CPU-01"]],
            "Jurisdição": [["sigla"], ["BR"]],
        }
    )
    pipeline = _make_pipeline()
    result = pipeline.parse_planilha(excel)
    assert "insumos" in result
    assert "cpus" in result
    assert "jurisdicoes" in result


def test_parse_planilha_missing_sheet():
    # xlsx without "CPUs" sheet → cpus key returns []
    excel = _make_excel_bytes(
        {
            "Insumos": [["codigo"], ["INS-01"]],
            "Jurisdição": [["sigla"], ["BR"]],
        }
    )
    pipeline = _make_pipeline()
    result = pipeline.parse_planilha(excel)
    assert result["cpus"] == []


# ── processar_catalogo ────────────────────────────────────────────────────────


def test_processar_catalogo_happy_path():
    """Pipeline runs end-to-end with mocked validators and storage."""
    # Minimal valid xlsx: Insumos + CPUs (no insumo refs so no validation errors)
    excel = _make_excel_bytes(
        {
            "Insumos": [["codigo", "descricao"], ["INS-01", "Cimento"]],
            "CPUs": [["cpu_id", "descricao"], ["CPU-01", "Escavação"]],
        }
    )
    storage = MagicMock()
    storage.gravar_blob.return_value = None
    pipeline = ExcelIngestionPipeline(storage)

    with (
        patch("concilium.infrastructure.catalog.excel_ingestion.ValidadorCPUInsumo") as mock_validator_cls,
        patch("concilium.infrastructure.catalog.excel_ingestion.validar_sem_ciclos") as mock_ciclos,
        patch("concilium.infrastructure.repositories.insumo.RepositorioInsumo") as mock_repo_cls,
    ):
        mock_validator_instance = MagicMock()
        mock_validator_instance.validar_cpu_completa.return_value = []
        mock_validator_cls.return_value = mock_validator_instance
        mock_ciclos.return_value = None

        blob_path = pipeline.processar_catalogo(
            data_base=date(2026, 1, 1),
            versao_tag="2026-01-test",
            excel_bytes=excel,
        )

    assert isinstance(blob_path, str)
    assert "2026-01-test" in blob_path
    storage.gravar_blob.assert_called_once()


def test_processar_catalogo_validation_error():
    """When validator returns ERRO validations, processar_catalogo raises ValueError."""
    excel = _make_excel_bytes(
        {
            "Insumos": [["codigo", "descricao"], ["INS-01", "Cimento"]],
            "CPUs": [["cpu_id", "descricao"], ["CPU-01", "Escavação"]],
        }
    )
    storage = MagicMock()
    pipeline = ExcelIngestionPipeline(storage)

    erro_mock = MagicMock()
    erro_mock.status = "ERRO"

    with (
        patch("concilium.infrastructure.catalog.excel_ingestion.ValidadorCPUInsumo") as mock_validator_cls,
        patch("concilium.infrastructure.catalog.excel_ingestion.validar_sem_ciclos"),
        patch("concilium.infrastructure.repositories.insumo.RepositorioInsumo"),
    ):
        mock_validator_instance = MagicMock()
        mock_validator_instance.validar_cpu_completa.return_value = [erro_mock]
        mock_validator_cls.return_value = mock_validator_instance

        with pytest.raises(ValueError, match="erros críticos"):
            pipeline.processar_catalogo(
                data_base=date(2026, 1, 1),
                versao_tag="2026-01-test",
                excel_bytes=excel,
            )
