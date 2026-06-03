"""Lê headers e amostra das abas de catálogo da planilha RFI."""
import io, sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

import openpyxl

XLSM = Path(__file__).resolve().parents[1] / "docs/product/3-reference_files/catalogos/PLANILHA PADRÃO RFI.xlsm"
wb = openpyxl.load_workbook(XLSM, read_only=True, data_only=True)

TARGET_SHEETS = [
    "BD_MÃO_DE_OBRA", "BD_EQUIPAMENTOS", "BD_MATERIAIS",
    "BD_PRODUTIVIDADES", "ORÇAMENTO", "ENCARGOS", "BDI", "PARÂMETROS",
    "SERVIÇOS TERCEIRIZADOS",
]

for name in TARGET_SHEETS:
    if name not in wb.sheetnames:
        print(f"\n[AUSENTE] {name}")
        continue
    ws = wb[name]
    print(f"\n{'='*65}")
    print(f"ABA: {name}")
    rows_data = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if all(c is None for c in row):
            continue
        rows_data.append(row)
        if len(rows_data) >= 8:
            break
    if not rows_data:
        print("  (vazia)")
        continue
    # Primeira linha não-nula como header
    header = rows_data[0]
    ncols = sum(1 for c in header if c is not None)
    print(f"  {ncols} colunas, {len(rows_data)-1} linhas de amostra")
    for ci, h in enumerate(header):
        if h is None:
            continue
        samples = []
        for row in rows_data[1:]:
            if ci < len(row) and row[ci] is not None:
                val = str(row[ci])[:50]
                samples.append(val)
        print(f"    col {ci:2d}: {str(h):<35}  | {' / '.join(samples[:4])}")

wb.close()
